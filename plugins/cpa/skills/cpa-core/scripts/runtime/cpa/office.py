"""Bounded Office package checks and preservation gates for CPA writers.

C1/C7 and template writers: hard rules 7 and 11; financial verification and visual review
remain separate. FIXTURE - confirm against her file. This is not full OOXML schema validation.
No external target is fetched and no Office application, macro or formula is executed.
"""
from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, time
import hashlib
import json
from pathlib import Path
import posixpath
import re
from urllib.parse import unquote, urlsplit
import zipfile

MAX_EXPANDED_BYTES = 512 * 1024 * 1024
MAX_MEMBERS = 10000
MAX_DIRECTORY_BYTES = 8 * 1024 * 1024
MAX_XML_DEPTH = 128
MAX_REFERENCES = 100000
MAX_ANNOTATION_BYTES = 8 * 1024 * 1024
_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
_DOCREL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_CT = "http://schemas.openxmlformats.org/package/2006/content-types"
_SHEET = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_DRAW = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
_EXTERNAL_FORMULA = re.compile(r"(?<![\w\[])\[(?:\d+|[^\]]+\.xls[xmb]?)\][^][!()]{0,256}!", re.I)
_WORKBOOK_PARTS = re.compile(
    r"(?:\[Content_Types\]\.xml|_rels/\.rels|docProps/(?:core|app)\.xml|"
    r"xl/(?:workbook|styles|sharedStrings)\.xml|xl/_rels/workbook\.xml\.rels|"
    r"xl/(?:worksheets/sheet\d+|theme/theme\d+|charts/chart\d+|drawings/drawing\d+|tables/table\d+)\.xml|"
    r"xl/(?:worksheets|drawings|charts)/_rels/[^/]+\.rels|xl/media/[^/]+|"
    r"xl/comments/comment\d+\.xml|xl/drawings/commentsDrawing\d+\.vml)\Z"
)


class OfficeSafetyError(ValueError):
    """A file cannot safely be edited or a candidate failed a pre-replacement check."""


@dataclass(frozen=True)
class PackageReport:
    """Scoped structural evidence; hashes are semantic for XML and byte-based for opaque parts."""

    kind: str
    parts: dict[str, str]
    external_links: int
    raw_parts: dict[str, str]


def _target(source: str, target: str) -> str:
    parsed = urlsplit(target)
    if parsed.scheme or parsed.netloc or parsed.query or "\\" in target:
        raise OfficeSafetyError("Invalid internal relationship target")
    value = unquote(parsed.path)
    resolved = posixpath.normpath(posixpath.join(posixpath.dirname(source), value))
    if value.startswith("/"):
        resolved = posixpath.normpath(value).lstrip("/")
    if resolved in ("", ".", "..") or resolved.startswith("../") or "\\" in resolved:
        raise OfficeSafetyError("Relationship target escapes the package")
    return resolved


def _rel_source(name: str) -> str:
    if name == "_rels/.rels":
        return ""
    parent, _, leaf = name.rpartition("/")
    if not parent.endswith("/_rels"):
        raise OfficeSafetyError("Invalid relationship part location")
    return parent[:-6] + "/" + leaf[:-5]


def _directory_limits(path):
    import struct

    with Path(path).open("rb") as stream:
        stream.seek(0, 2)
        length = stream.tell()
        stream.seek(max(0, length - 65557))
        trailer = stream.read(65557)
    index = trailer.rfind(b"PK\x05\x06")
    if index < 0 or index + 22 > len(trailer):
        raise OfficeSafetyError("Missing ZIP end record")
    _, disk, directory_disk, disk_count, count, size, offset, comment = struct.unpack(
        "<4s4H2IH", trailer[index:index + 22])
    if count > MAX_MEMBERS or size > MAX_DIRECTORY_BYTES:
        raise OfficeSafetyError("ZIP directory inspection limit exceeded")
    if disk or directory_disk or disk_count != count or size == 0xffffffff or offset == 0xffffffff:
        raise OfficeSafetyError("Multi-volume/ZIP64 directories are unsupported")
    if index + 22 + comment != len(trailer) or offset + size != length - len(trailer) + index:
        raise OfficeSafetyError("Ambiguous ZIP directory bounds")


def inspect_package(path: Path | str, *, workbook_edit: bool = False) -> PackageReport:
    """Inspect within the shared Windows long-path guard; see the bounded package policy below."""
    from cpa import fsutil

    with fsutil.guard_long_path(path):
        return _inspect_package(path, workbook_edit=workbook_edit)


def _inspect_package(path: Path | str, *, workbook_edit: bool = False) -> PackageReport:
    """Stream and check package structure without extraction, network access or whole-sheet loads.

    Workbook-edit mode additionally refuses untested feature classes. Inspection alone does not
    authorize a write. Expanded byte, member, XML depth and relationship budgets fail closed.
    """
    from lxml import etree

    parts: dict[str, str] = {}
    raw_parts: dict[str, str] = {}
    rels: dict[str, dict[str, tuple[str, str, str]]] = {}
    refs: dict[str, set[str]] = {}
    defaults: dict[str, str] = {}
    overrides: dict[str, str] = {}
    external = 0
    try:
        _directory_limits(path)
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
            names = {m.filename for m in members}
            if len(members) > MAX_MEMBERS or sum(m.file_size for m in members) > MAX_EXPANDED_BYTES:
                raise OfficeSafetyError("Office package exceeds inspection limit")
            if len(names) != len(members):
                raise OfficeSafetyError("Duplicate package member names")
            for member in members:
                name = member.filename
                if ("\\" in name or name.startswith("/") or ":" in name or
                        any(p in ("", ".", "..") for p in name.split("/")) or
                        (member.external_attr >> 16) & 0o170000 == 0o120000):
                    raise OfficeSafetyError("Unsafe package member name or symlink")
                if workbook_edit and not _WORKBOOK_PARTS.fullmatch(name):
                    raise OfficeSafetyError(f"Unsupported workbook part; write blocked: {name}")
            if "[Content_Types].xml" not in names or "_rels/.rels" not in names:
                raise OfficeSafetyError("Missing package content types or root relationships")
            if sum(m.file_size for m in members if _annotation(m.filename)) > MAX_ANNOTATION_BYTES:
                raise OfficeSafetyError("Annotation preservation limit exceeded")
            used_bytes = 0
            ref_count = 0
            for member in members:
                name = member.filename
                digest = hashlib.sha256()
                raw_digest = hashlib.sha256()
                is_xml = name.endswith((".xml", ".rels", ".vml"))
                parser = etree.XMLPullParser(events=("start", "end"), resolve_entities=False,
                                             load_dtd=False, no_network=True, huge_tree=False) if is_xml else None
                depth = 0
                source = _rel_source(name) if name.endswith(".rels") else None
                if source is not None:
                    if source and source not in names:
                        raise OfficeSafetyError("Relationship source part is missing")
                    rels[source] = {}
                with archive.open(member, "r") as stream:  # portability: ok ZIP binary member stream
                    while chunk := stream.read(65536):
                        raw_digest.update(chunk)
                        used_bytes += len(chunk)
                        if used_bytes > MAX_EXPANDED_BYTES:
                            raise OfficeSafetyError("Expanded package exceeds inspection limit")
                        if parser is None:
                            digest.update(chunk)
                            continue
                        parser.feed(chunk)
                        for event, elem in parser.read_events():
                            if not isinstance(elem.tag, str):
                                raise OfficeSafetyError("XML entities are not supported")
                            if event == "start":
                                depth += 1
                                if depth > MAX_XML_DEPTH:
                                    raise OfficeSafetyError("XML depth limit exceeded")
                                if elem.getroottree().docinfo.doctype:
                                    raise OfficeSafetyError("XML DTDs are not supported")
                                if workbook_edit and name.endswith(".vml"):
                                    if ((elem.tag.endswith("}ClientData") and elem.get("ObjectType") != "Note") or
                                            (elem.tag == "{urn:schemas-microsoft-com:vml}shape" and elem.get("type") != "#_x0000_t202") or
                                            elem.tag.endswith(("}imagedata", "}FmlaMacro"))):
                                        raise OfficeSafetyError("Only comment-note VML is supported; controls are blocked")
                                if workbook_edit and (elem.tag.endswith("}extLst") or
                                        elem.tag in {f"{{{_DRAW}}}sp", f"{{{_DRAW}}}grpSp", f"{{{_DRAW}}}cxnSp"} or
                                        (elem.tag == f"{{{_SHEET}}}f" and elem.get("t") in ("array", "dataTable"))):
                                    raise OfficeSafetyError(f"Unsupported workbook feature in {name}; write blocked")
                                for key, value in elem.attrib.items():
                                    if key.startswith("{" + _DOCREL + "}"):
                                        refs.setdefault(name, set()).add(value)
                                        ref_count += 1
                                digest.update(repr(("start", elem.tag, sorted(elem.attrib.items()))).encode("utf-8"))
                                continue
                            if (workbook_edit and elem.tag in {f"{{{_SHEET}}}f", f"{{{_SHEET}}}definedName"}
                                    and _EXTERNAL_FORMULA.search(elem.text or "")):
                                raise OfficeSafetyError("External workbook dependency in formula/name; write blocked")
                            if elem.tag == f"{{{_REL}}}Relationship":
                                if source is None:
                                    raise OfficeSafetyError("Relationship outside a relationships part")
                                rid, target, kind = (elem.get(k, "") for k in ("Id", "Target", "Type"))
                                mode = elem.get("TargetMode", "Internal")
                                if not rid or not target or not kind or rid in rels[source] or mode not in ("Internal", "External"):
                                    raise OfficeSafetyError("Invalid or duplicate relationship")
                                rels[source][rid] = (target, kind, mode)
                                ref_count += 1
                                if mode == "External":
                                    external += 1
                                    if workbook_edit and not kind.endswith("/hyperlink"):
                                        raise OfficeSafetyError("External workbook dependency; write blocked")
                            if name == "[Content_Types].xml":
                                if elem.tag == f"{{{_CT}}}Default":
                                    key = elem.get("Extension", "")
                                    if not key or key in defaults:
                                        raise OfficeSafetyError("Duplicate or invalid content type default")
                                    defaults[key] = elem.get("ContentType", "")
                                elif elem.tag == f"{{{_CT}}}Override":
                                    key = elem.get("PartName", "").lstrip("/")
                                    if not key or key in overrides:
                                        raise OfficeSafetyError("Duplicate or invalid content type override")
                                    overrides[key] = elem.get("ContentType", "")
                            if ref_count > MAX_REFERENCES:
                                raise OfficeSafetyError("Relationship inspection limit exceeded")
                            text = elem.text or ""
                            if not text.strip() and not elem.tag.endswith("}t"):
                                text = ""
                            digest.update(repr(("end", elem.tag, text)).encode("utf-8"))
                            depth -= 1
                            elem.clear()
                            while elem.getprevious() is not None:
                                del elem.getparent()[0]
                    if parser is not None:
                        root = parser.close()
                        if root.getroottree().docinfo.doctype:
                            raise OfficeSafetyError("XML DTDs are not supported")
                parts[name] = digest.hexdigest()
                raw_parts[name] = raw_digest.hexdigest()
            for name in names - {"[Content_Types].xml"}:
                if not overrides.get(name, defaults.get(name.rsplit(".", 1)[-1])):
                    raise OfficeSafetyError(f"Missing content type for {name}")
            if set(overrides) - names:
                raise OfficeSafetyError("Content type references a missing part")
            for source, entries in rels.items():
                for target, _, mode in entries.values():
                    if mode == "Internal" and _target(source, target) not in names:
                        raise OfficeSafetyError(f"Missing relationship target from {source or 'package'}")
            for name, ids in refs.items():
                if ids - rels.get(name, {}).keys():
                    raise OfficeSafetyError(f"Unresolved relationship reference in {name}")
            roots = [(_target("", target), kind) for target, kind, mode in rels.get("", {}).values()
                     if mode == "Internal" and kind.endswith("/officeDocument")]
            if len(roots) != 1:
                raise OfficeSafetyError("Expected one Office document root")
            root = roots[0][0]
            kind = {"xl/workbook.xml": "xlsx", "ppt/presentation.xml": "pptx"}.get(root)
            if not kind or (workbook_edit and kind != "xlsx"):
                raise OfficeSafetyError("Unsupported Office document type")
            expected = ("spreadsheetml.sheet.main+xml" if kind == "xlsx" else "presentationml.presentation.main+xml")
            if not overrides.get(root, "").endswith(expected):
                raise OfficeSafetyError("Unsupported or mismatched document content type")
    except OfficeSafetyError:
        raise
    except (OSError, ValueError, RuntimeError, zipfile.BadZipFile, etree.XMLSyntaxError) as exc:
        raise OfficeSafetyError(f"Office package inspection failed: {exc}") from exc
    return PackageReport(kind, parts, external, raw_parts)


def _xml(value):
    if value is None:
        return None
    elem = value.to_tree() if hasattr(value, "to_tree") else value
    return (elem.tag, tuple(sorted((k, v) for k, v in elem.attrib.items() if k != "localSheetId")),
            elem.text or "", tuple(_xml(child) for child in elem))


def _value(value, *, wire=False):
    from openpyxl.cell.rich_text import CellRichText

    if isinstance(value, CellRichText):
        return _xml(value)
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return datetime.combine(value, time()).isoformat()
    if isinstance(value, float) and wire:
        return float(format(value, ".16g"))
    return None if value == "" else value


def _snapshot(wb, *, wire=False) -> dict:
    from copy import copy

    sheets = {}
    for ws in wb.worksheets:
        for dimension in ws.column_dimensions.values():
            dimension.reindex()  # the writer materializes min/max on serialization
        cells = {}
        for cell in ws._cells.values():
            if (_value(cell.value) is None and not cell.has_style and not getattr(cell, "hyperlink", None)
                    and not getattr(cell, "comment", None)):
                continue
            link = getattr(cell, "hyperlink", None)
            comment = getattr(cell, "comment", None)
            cells[cell.coordinate] = (
                _value(cell.value, wire=wire), tuple(_xml(copy(getattr(cell, attr))) for attr in
                                         ("font", "fill", "border", "alignment", "protection")),
                cell.number_format,
                tuple(getattr(link, k) for k in ("target", "location", "display", "tooltip")) if link else None,
                (comment.text, comment.author) if comment else None,
                cell.data_type if _value(cell.value) is not None else "empty",
            )
        metadata = (
            ws.sheet_state, _xml(ws.protection), str(ws.merged_cells), _xml(ws.data_validations),
            tuple((k, _xml(v)) for k, v in sorted(ws.defined_names.items())),
            tuple((str(k), tuple(_xml(r) for r in v)) for k, v in ws.conditional_formatting._cf_rules.items()),
            tuple((k, _xml(v)) for k, v in sorted(ws.row_dimensions.items())),
            tuple((k, _xml(v)) for k, v in sorted(ws.column_dimensions.items())),
            _xml(ws.sheet_properties), _xml(ws.sheet_format), _xml(ws.page_setup), _xml(ws.page_margins),
            str(ws.print_area), ws.print_title_rows, ws.print_title_cols, ws.freeze_panes,
        )
        sheets[ws.title] = {"cells": cells, "metadata": metadata}
    return {"sheets": sheets, "names": tuple((k, _xml(v)) for k, v in sorted(wb.defined_names.items())),
            "epoch": wb.epoch, "order": wb.sheetnames}


def load_workbook(path: Path | str):
    """Preflight a supported small workbook and retain its original preservation baseline."""
    import openpyxl
    from cpa import bigxlsx

    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        raise OfficeSafetyError("Only .xlsx template editing is supported")
    original_hash = _file_hash(path)
    report = inspect_package(path, workbook_edit=True)
    if bigxlsx.is_large(path):
        raise OfficeSafetyError("Large workbook: use the streaming workflow; whole-workbook editing is blocked")
    wb = openpyxl.load_workbook(path, rich_text=True, keep_links=True)
    wb._cpa_original = _snapshot(wb)
    wb._cpa_package = report
    with zipfile.ZipFile(path) as archive:
        wb._cpa_annotations = {name: archive.read(name) for name in report.parts if _annotation(name)}
    wb._cpa_source = (path, original_hash)
    try:
        _check_source(wb)
    except BaseException:
        wb.close()
        raise
    return wb


def _preservation(original, current, changed_cells, changed_sheets, added_sheets):
    if original["names"] != current["names"] or original["epoch"] != current["epoch"]:
        raise OfficeSafetyError("Workbook names/calendar preservation failed")
    before, after = original["sheets"], current["sheets"]
    if set(after) - set(before) - set(added_sheets) or set(before) - set(after):
        raise OfficeSafetyError("Sheet presence preservation failed")
    if [n for n in current["order"] if n in before] != original["order"]:
        # Replacing an explicitly owned sheet may move it; other sheets retain their relative order.
        if [n for n in current["order"] if n in before and n not in changed_sheets] != [
                n for n in original["order"] if n not in changed_sheets]:
            raise OfficeSafetyError("Sheet order preservation failed")
    for name, old in before.items():
        if name in changed_sheets:
            continue
        new = after[name]
        if old["metadata"] != new["metadata"]:
            raise OfficeSafetyError(f"Sheet metadata preservation failed: {name}")
        allowed = changed_cells.get(name, set())
        for coord in old["cells"].keys() | new["cells"].keys():
            a, b = old["cells"].get(coord), new["cells"].get(coord)
            if coord not in allowed and a != b:
                raise OfficeSafetyError(f"Cell preservation failed: {name}!{coord}")
            if coord in allowed and a and b and a[1][-1] != b[1][-1]:
                raise OfficeSafetyError(f"Cell protection preservation failed: {name}!{coord}")


def check_workbook_preservation(original: Path, candidate: Path, *, changed_cells=None,
                                changed_sheets=(), added_sheets=()) -> None:
    """Check a completed workflow against its original template, not just the last writer."""
    before = load_workbook(original)
    try:
        after = load_workbook(candidate)
        try:
            _preservation(before._cpa_original, after._cpa_original, changed_cells or {},
                          set(changed_sheets), set(added_sheets))
            _package_preservation(before._cpa_package, after._cpa_package)
        finally:
            after.close()
    finally:
        before.close()


def _annotation(name):
    return name.startswith("xl/comments/") or (name.startswith("xl/drawings/commentsDrawing") and name.endswith(".vml"))


def _comments(snapshot):
    return {(name, coord): value[4] for name, sheet in snapshot["sheets"].items()
            for coord, value in sheet["cells"].items() if value[4] is not None}


def _restore_annotations(path, parts):
    """Retain unchanged comment XML/VML byte-for-byte, including formatting and popup geometry."""
    import shutil
    from cpa import fsutil

    def write(tmp):
        with zipfile.ZipFile(path) as source, zipfile.ZipFile(tmp, "w") as target:
            if parts.keys() - set(source.namelist()):
                raise OfficeSafetyError("Annotation part mapping changed; preservation is not supported")
            for info in source.infolist():
                if info.filename in parts:
                    target.writestr(info, parts[info.filename])
                else:
                    with source.open(info) as src, target.open(info, "w") as dst:  # portability: ok ZIP binary members
                        shutil.copyfileobj(src, dst)
    fsutil.atomic_write(path, write)


def _package_preservation(source, report):
    for name, digest in source.parts.items():
        if _annotation(name) or name.startswith(("xl/charts/", "xl/drawings/", "xl/media/", "xl/theme/", "xl/tables/")):
            if report.parts.get(name) != digest:
                raise OfficeSafetyError(f"Package feature preservation failed: {name}")


def _difference(a, b, location="workbook"):
    if a == b:
        return ""
    if isinstance(a, dict) and isinstance(b, dict):
        for key in a.keys() | b.keys():
            if key not in a or key not in b:
                return f"{location}.{key}: presence changed"
            found = _difference(a[key], b[key], f"{location}.{key}")
            if found:
                return found
    if isinstance(a, (list, tuple)) and isinstance(b, (list, tuple)) and len(a) == len(b):
        for index, (first, second) in enumerate(zip(a, b)):
            found = _difference(first, second, f"{location}[{index}]")
            if found:
                return found
    return f"{location}: differs"


def save_workbook(wb, path: Path | str, *, changed_cells: dict[str, set[str]] | None = None,
                  changed_sheets=(), added_sheets=(), retries: int = 1, wait_s: float = 0.5) -> Path:
    """Save once, validate the candidate, then atomically replace; reject out-of-scope template edits.

    Whole-sheet ownership is explicit. Streaming/new workbooks get package checks; small normal
    workbooks also get serialized cell/metadata comparison. This never recalculates formulas.
    """
    import openpyxl
    from cpa import bigxlsx, fsutil

    _check_source(wb)
    expected = None if wb.write_only else _snapshot(wb, wire=True)
    original = getattr(wb, "_cpa_original", None)
    if original is not None:
        _preservation(original, _snapshot(wb), changed_cells or {}, set(changed_sheets), set(added_sheets))

    annotations = getattr(wb, "_cpa_annotations", {})
    if annotations and _comments(original) != _comments(expected):
        raise OfficeSafetyError("Editing existing annotations is unsupported; preservation would be ambiguous")

    def write(tmp: Path) -> None:
        wb.save(tmp)
        if annotations:
            _restore_annotations(tmp, annotations)
        report = inspect_package(tmp, workbook_edit=True)
        source = getattr(wb, "_cpa_package", None)
        if source:
            _package_preservation(source, report)
        if expected is not None:
            if bigxlsx.is_large(tmp):
                raise OfficeSafetyError("Candidate requires streaming; cannot prove whole-workbook preservation")
            result = openpyxl.load_workbook(tmp, rich_text=True, keep_links=True)
            try:
                difference = _difference(expected, _snapshot(result, wire=True))
                if difference:
                    raise OfficeSafetyError(f"Serialized workbook preservation failed: {difference}")
                if original is not None:
                    _preservation(original, _snapshot(result), changed_cells or {}, set(changed_sheets), set(added_sheets))
            finally:
                result.close()
        _check_source(wb)
    return fsutil.atomic_write(path, write, retries=retries, wait_s=wait_s)


def _presentation_state(prs):
    from copy import deepcopy

    return {"size": (prs.slide_width, prs.slide_height), "slides": [
        (str(slide.part.partname), deepcopy(slide._element),
         deepcopy(slide.notes_slide._element) if slide.has_notes_slide else None) for slide in prs.slides]}


def _presentation_signature(state, changed_shapes, allow_notes):
    from copy import deepcopy

    records = []
    for n, (name, root, notes) in enumerate(state["slides"], 1):
        root = deepcopy(root)
        allowed = changed_shapes.get(n, set())
        trees = root.xpath(".//p:spTree")
        if trees:
            for child in list(trees[0]):
                ids = child.xpath(".//p:cNvPr")
                if ids and int(ids[0].get("id")) in allowed:
                    trees[0].remove(child)
        records.append((name, _xml(root), None if allow_notes else _xml(notes)))
    return state["size"], records


def load_presentation(path: Path | str):
    """Inspect an existing deck and retain the original slide/part preservation baseline."""
    from pptx import Presentation

    path = Path(path)
    original_hash = _file_hash(path)
    report = inspect_package(path)
    if path.suffix.lower() != ".pptx" or report.kind != "pptx":
        raise OfficeSafetyError("Only .pptx editing is supported")
    prs = Presentation(path)
    prs._cpa_package = report
    prs._cpa_original = _presentation_state(prs)
    prs._cpa_source = (path, original_hash)
    _check_source(prs)
    return prs


def _relationship_name(name):
    parent, leaf = posixpath.split(name)
    return posixpath.join(parent, "_rels", leaf + ".rels")


def save_presentation(prs, path: Path | str, *, changed_shapes=None, changed_parts=(), allow_notes=False) -> Path:
    """Check scoped template preservation and serialized parts before replacing a deck.

    A mapped shape is explicitly owned by the caller; mixed-run replacement has its own narrower
    contract. Unknown original parts must survive, even when the object model did not read them.
    """
    from cpa import fsutil

    _check_source(prs)
    changed_shapes = changed_shapes or {}
    original = getattr(prs, "_cpa_original", None)
    if original and _presentation_signature(original, changed_shapes, allow_notes) != _presentation_signature(
            _presentation_state(prs), changed_shapes, allow_notes):
        raise OfficeSafetyError("Slide template preservation failed outside the mapped shapes")
    expected = {}
    package = prs.part.package
    for part in package.iter_parts():
        name = str(part.partname).lstrip("/")
        expected[name] = hashlib.sha256(part.blob).hexdigest()
        if len(part.rels):
            expected[_relationship_name(name)] = hashlib.sha256(part.rels.xml).hexdigest()
    expected["_rels/.rels"] = hashlib.sha256(package._rels.xml).hexdigest()

    def write(tmp: Path) -> None:
        prs.save(tmp)
        report = inspect_package(tmp)
        if report.kind != "pptx":
            raise OfficeSafetyError("Expected a presentation package")
        for name, digest in expected.items():
            if report.raw_parts.get(name) != digest:
                raise OfficeSafetyError(f"Serialized presentation preservation failed: {name}")
        source = getattr(prs, "_cpa_package", None)
        if source:
            owned = set(changed_parts) | {"[Content_Types].xml", "_rels/.rels", "docProps/core.xml",
                                         "ppt/presentation.xml", "ppt/_rels/presentation.xml.rels"}
            for n, slide in enumerate(prs.slides, 1):
                name = str(slide.part.partname).lstrip("/")
                if changed_shapes.get(n):
                    owned.add(name)
                if allow_notes:
                    owned.add(_relationship_name(name))
            for name, digest in source.parts.items():
                if name in owned or (allow_notes and name.startswith("ppt/notesSlides/")):
                    continue
                if report.parts.get(name) != digest:
                    raise OfficeSafetyError(f"Presentation part preservation failed: {name}")
        _check_source(prs)
    return fsutil.atomic_write(path, write)


def _file_hash(path: Path) -> str:
    from cpa import fsutil

    with fsutil.guard_long_path(path), path.open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


def _check_source(document):
    source = getattr(document, "_cpa_source", None)
    if source is not None and _file_hash(source[0]) != source[1]:
        raise OfficeSafetyError("Source changed after preflight; reload the approved file before editing")


def publish_artifacts(files: dict[Path, Path]) -> None:
    """Check every staged candidate before delivery; roll back ordinary copy failures.

    Multiple replacements are not a filesystem transaction. Concurrent readers may see the
    transition; a rollback failure is explicit and blocks approval. Sources stay in staging.
    """
    import shutil
    import tempfile
    from cpa import fsutil

    if not files:
        return
    pairs = [(Path(source), Path(target)) for source, target in files.items()]
    if len({target for _, target in pairs}) != len(pairs) or {s for s, _ in pairs} & {t for _, t in pairs}:
        raise OfficeSafetyError("Publication paths overlap")
    hashes = {}
    for source, _ in pairs:
        before = _file_hash(source)
        if source.suffix.lower() in (".xlsx", ".pptx"):
            report = inspect_package(source)
            if report.kind != source.suffix.lower()[1:]:
                raise OfficeSafetyError("Publication file type mismatch")
        if _file_hash(source) != before:
            raise OfficeSafetyError("Staged candidate changed during validation")
        hashes[source] = before

    def copy_checked(source, target, digest):
        def write(tmp):
            shutil.copyfile(source, tmp)
            if _file_hash(tmp) != digest:
                raise OfficeSafetyError("Candidate changed after validation; publication blocked")
        fsutil.atomic_write(target, write)

    folder = Path(tempfile.mkdtemp(prefix="cpa_publication_", dir=pairs[0][0].parent))
    keep_backups = False
    try:
        backups = {}
        for index, (_, target) in enumerate(pairs):
            target.parent.mkdir(parents=True, exist_ok=True)
            if target.exists():
                backup = Path(folder) / fsutil.safe_filename(f"backup_{index}.bin")
                digest = _file_hash(target)
                copy_checked(target, backup, digest)
                backups[target] = (backup, digest)
        delivered = []
        try:
            for source, target in pairs:
                copy_checked(source, target, hashes[source])
                delivered.append(target)
        except BaseException as exc:
            failures = []
            for target in reversed(delivered):
                try:
                    if target in backups:
                        backup, digest = backups[target]
                        copy_checked(backup, target, digest)
                    else:
                        target.unlink(missing_ok=True)
                except (OSError, OfficeSafetyError):
                    failures.append(target.name)
            if failures:
                keep_backups = True
                raise OfficeSafetyError("Incomplete publication rollback; DO NOT SEND: " + ", ".join(failures)
                                        + f". Recovery copies retained at {folder}") from exc
            raise
    finally:
        if not keep_backups:
            shutil.rmtree(folder, ignore_errors=True)


def deliver_artifacts(files: dict[Path, Path]) -> None:
    """Publish checked artifacts plus relocated sidecars, retaining valid staging manifests."""
    import tempfile
    from cpa import fsutil, manifest

    if not files:
        return
    expanded = dict(files)
    relocated = {manifest.to_rel(source): manifest.to_rel(target) for source, target in files.items()}
    with tempfile.TemporaryDirectory(prefix="cpa_delivery_", dir=next(iter(files)).parent) as directory:
        for index, (source, target) in enumerate(files.items()):
            if not manifest.sidecar(source).is_file():
                continue
            data = manifest.read(source)
            if data.get("sha256") != _file_hash(source):
                raise OfficeSafetyError("Staged manifest hash is stale; delivery blocked")
            data["path"] = manifest.to_rel(target)
            if isinstance(data.get("verification"), dict):
                ref = data["verification"].get("file")
                if ref in relocated:
                    data["verification"]["file"] = relocated[ref]
            side = Path(directory) / fsutil.safe_filename(f"manifest_{index}.json")
            fsutil.atomic_write(side, lambda tmp, data=data: tmp.write_text(
                json.dumps(data, indent=2, ensure_ascii=False) + "\n", encoding="utf-8", newline="\n"))
            expanded[side] = manifest.sidecar(target)
        publish_artifacts(expanded)


def acceptance(*, financial: str = "NOT_CHECKED", geometry_issues: int | None = None,
               package: str = "NOT_CHECKED", preservation: str = "NOT_CHECKED") -> dict[str, str]:
    """Report independent evidence dimensions; never infer financial or visual acceptance."""
    return {"package": package, "preservation": preservation, "financial": financial,
            "geometry": "NOT_CHECKED" if geometry_issues is None else (
                "REVIEW_REQUIRED" if geometry_issues else "NO_FINDINGS_IN_CHECKED_SCOPE"),
            "visual": "NOT_REVIEWED"}
