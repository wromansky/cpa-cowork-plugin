"""Large workbook streamer (C7): bounded-memory row iteration for .xlsx, .xlsm and .xlsb.

The single chokepoint for every workbook read at or above 15 MB (hard rule 11: "Workbooks above 15 MB stream
through cpa.bigxlsx. Never load whole."). `is_large(path)` is the one predicate callers branch on and means
15 MB *uncompressed* (on-disk size first, then the zip central directory; nothing is decompressed).
`iter_rows(path, sheet)` yields plain lists, never a DataFrame (D10: this module is pandas-free).

Build-list items: C7 large workbook streamer (build-list:287), guide step 7 / 13.7 (20 MB generated-file test),
MAX_PATH boundary (AGENTS.md, Windows-first), the U05 half of the pandas boundary (D10).
Hard rules enforced: 11 (stream above 15 MB, never load whole). Also D10 (openpyxl read_only, never load_workbook
without read_only, no pandas), D13 (temp file in the same directory + os.replace, bounded retry on
PermissionError, pathlib only), D03 (import and register() are cheap: no file, workspace or assumption access).

Design (U05 plan, decisions D1-D9):
- D1 the read path is openpyxl `read_only=True, data_only=True` (cached values only), not hand-rolled lxml
  iterparse; openpyxl degrades to xml.etree without lxml (speed, not correctness), so lxml is never required.
- D2 `is_large` answers the uncompressed question in two stages (see its docstring).
- D3 `.xlsb` is read through pyxlsb and converted by `to_xlsx` (pyxlsb -> openpyxl write_only, both streaming);
  no LibreOffice in the production path. Both .xlsb legs are proven only against a hand-encoded BIFF12 test
  workbook (the record subset the pinned pyxlsb parses); UNPROVEN against an Excel-written .xlsb.
- D4 dependencies are checked lazily on first use (`MissingDependency`), never at import.
- D6 interior gap rows are yielded as all-None lists; rows are padded, never truncated; the worksheet's
  <dimension> tag is distrusted (`reset_dimensions()` is called).
- D7 `write_rows` streams a flat sheet through a write_only workbook and saves via a temp file + os.replace.
- D8 `to_xlsx` converts every sheet, in workbook order, titles kept.
- D9 the CLI catches BigXlsxError and OSError only: one stderr line, exit 1, never a traceback. A file whose
  suffix is right but whose bytes are not a workbook (e.g. an SAP "Excel" export that is really HTML or text,
  renamed .xlsx) is wrapped into NotAWorkbook at open, so it too is one line, not a zipfile traceback.
- C13 boundary: `head`, `convert` and to_xlsx refuse any path with an Epic or quarantine component
  (EPIC_PATH_MARKERS, EpicPathRefused) before opening it, so no Epic row reaches the terminal or a copy except
  through `python -m cpa sources epic validate` (grain guard). iter_rows itself stays generic: the Epic adapter
  must run the grain guard on the header before it yields a row.

CLI: `python -m cpa bigxlsx sheets|head|count|convert` (on her machine `.venv\\Scripts\\python.exe -m cpa ...`).
`head` prints TSV: a None cell is an empty field; any other value is str(value) with each tab, CR and LF replaced
by one space, so one sheet row is one output line (no locale formatting).
"""
from __future__ import annotations

import argparse
import contextlib
import errno
import os
import sys
import time
import zipfile
from collections.abc import Iterable, Iterator, Sequence
from pathlib import Path
from types import ModuleType
from typing import Any

__all__ = [
    "STREAM_THRESHOLD_BYTES",
    "MAX_PATH",
    "SUPPORTED_SUFFIXES",
    "UNCOMPRESSED_PREFIXES",
    "UNCOMPRESSED_ENDINGS",
    "BigXlsxError",
    "UnsupportedFormat",
    "MissingDependency",
    "SheetNotFound",
    "LongPathError",
    "NotAWorkbook",
    "EpicPathRefused",
    "EPIC_PATH_MARKERS",
    "is_large",
    "sheet_names",
    "iter_rows",
    "row_count",
    "to_xlsx",
    "write_rows",
    "register",
]

STREAM_THRESHOLD_BYTES: int = 15 * 1024 * 1024  # hard rule 11; uncompressed sheet+SST bytes (guide:305, build-list:287)
"""Hard rule 11 threshold in bytes (15 MB). Read as a module global at call time so tests can monkeypatch it."""

MAX_PATH: int = 260  # Win32 default cap; only used to recognise a failure
"""Win32 default path-length cap; used only to recognise a long-path failure, never to refuse a path up front."""

SUPPORTED_SUFFIXES: tuple[str, ...] = (".xlsx", ".xlsm", ".xlsb")
"""Workbook suffixes this module reads, compared case-insensitively (SAP exports are often named EXPORT.XLSX)."""

UNCOMPRESSED_PREFIXES: tuple[str, ...] = ("xl/worksheets/", "xl/sharedStrings.")
"""Zip member name prefixes counted by is_large stage 2 (a member must also end with UNCOMPRESSED_ENDINGS)."""

UNCOMPRESSED_ENDINGS: tuple[str, ...] = (".xml", ".bin")
"""Zip member name endings counted by is_large stage 2 (.bin covers .xlsb sheet and string parts)."""

EPIC_PATH_MARKERS: tuple[str, ...] = ("epic", "quarantine")
"""Case-folded substrings that mark a path as Epic territory (C13): any path component containing one is refused by
the row-printing and row-copying entry points (`head`, `convert`, to_xlsx). Epic rows are read only through
`python -m cpa sources epic validate`, which runs the grain guard."""

_DIST = "cpa-automation"  # pyproject [project] name: MissingDependency reads the D09 pin from its metadata
_READY = object()  # priming sentinel: iter_rows advances its generator past open + sheet resolution


# --------------------------------------------------------------------------------------------------- exceptions


class BigXlsxError(RuntimeError):
    """Base for every cpa.bigxlsx failure."""


class UnsupportedFormat(BigXlsxError):
    """Suffix not in SUPPORTED_SUFFIXES; message names it and the supported set."""


class MissingDependency(BigXlsxError):
    """A pinned reader is not importable; names the package, its pin and scripts\\bootstrap.ps1.

    Raised on first use, never at import (D4)."""


class SheetNotFound(BigXlsxError):
    """Named sheet absent; message lists every sheet name present.

    Attributes: `sheet` (what was asked for) and `present` (tuple of sheet names in workbook order)."""

    def __init__(self, message: str, *, sheet: Any = None, present: Sequence[str] = ()) -> None:
        super().__init__(message)
        self.sheet = sheet
        self.present = tuple(present)


class NotAWorkbook(BigXlsxError):
    """Right suffix, wrong bytes: the file is not a real .xlsx/.xlsm/.xlsb (not a zip, or a zip without a
    workbook). Message names the file; the original exception is chained as __cause__."""


class EpicPathRefused(BigXlsxError):
    """The path lies under an Epic or quarantine folder (EPIC_PATH_MARKERS); rows from it would bypass the C13
    grain guard, so `head`, `convert` and to_xlsx refuse it before opening anything."""


class LongPathError(BigXlsxError):
    """OSError on a path >= MAX_PATH chars; names MAX_PATH, the length and the LongPathsEnabled/GPO fix.

    Raised only on Windows, only for a path of MAX_PATH or more characters, and only for the errors a disabled
    long-path setting produces (winerror 3 or 206, or FileNotFoundError); anything else is re-raised unchanged."""


# ------------------------------------------------------------------------------------------------ private helpers


def _pin(name: str) -> str:
    """The requirement spec for `name` from the installed cpa-automation metadata (pyproject.toml's pin), or the
    bare name when the metadata is absent. Read at raise time, so the pin is written only in pyproject.toml."""
    import importlib.metadata
    import re

    try:
        requires = importlib.metadata.requires(_DIST) or []
    except importlib.metadata.PackageNotFoundError:
        return name
    for spec in requires:
        spec = spec.split(";", 1)[0].strip()
        if re.split(r"[\s<>=!~\[(]", spec, maxsplit=1)[0].casefold() == name.casefold():
            return spec
    return name


def _require(name: str) -> ModuleType:
    """Import a pinned reader on first use; MissingDependency naming the pin and the bootstrap script otherwise."""
    import importlib

    try:
        return importlib.import_module(name)
    except ImportError as exc:
        raise MissingDependency(
            f"{name} is not installed (pinned as {_pin(name)} in pyproject.toml). From the repo root run "
            f"scripts\\bootstrap.ps1 in PowerShell (Linux/mac: scripts/bootstrap.sh) to reinstall the pinned set."
        ) from exc


def _is_long_path_failure(path: Path, exc: OSError) -> bool:
    if os.name != "nt" or len(str(path)) < MAX_PATH:
        return False
    return getattr(exc, "winerror", None) in (3, 206) or isinstance(exc, FileNotFoundError)


@contextlib.contextmanager
def _guard_long_path(path: Path) -> Iterator[None]:
    """Re-raise a Windows long-path OSError as LongPathError; every other OSError passes through unchanged."""
    try:
        yield
    except BigXlsxError:
        raise
    except OSError as exc:
        if not _is_long_path_failure(path, exc):
            raise
        raise LongPathError(
            f"cannot open {path}: the path is {len(str(path))} characters, at or above the Windows MAX_PATH limit "
            f"of {MAX_PATH}. Move the file to a shorter path, or enable long paths (Group Policy: Computer "
            "Configuration > Administrative Templates > System > Filesystem > Enable Win32 long paths; registry "
            "value LongPathsEnabled=1), then open a new PowerShell window."
        ) from exc


def _refuse_epic(path: Path) -> None:
    """EpicPathRefused when any component of path (made absolute, not resolved) contains an EPIC_PATH_MARKERS
    substring, case-folded. Called before any file access by every entry point that prints or copies rows."""
    for part in Path(os.path.abspath(path)).parts:
        folded = part.casefold()
        hit = next((m for m in EPIC_PATH_MARKERS if m in folded), None)
        if hit is not None:
            raise EpicPathRefused(
                f"{path}: path component {part!r} marks an Epic or quarantine file. Epic files are read only "
                "through `python -m cpa sources epic validate` (grain guard)"
            )


def _check_suffix(path: Path) -> str:
    """Return path's lower-cased suffix, or raise UnsupportedFormat naming it and SUPPORTED_SUFFIXES."""
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise UnsupportedFormat(
            f"unsupported workbook format {path.suffix or '(no suffix)'!r} for {path.name}; "
            f"supported: {', '.join(SUPPORTED_SUFFIXES)}"
        )
    return suffix


def _resolve_sheet(names: Sequence[str], sheet: str | int | None, *, where: str = "") -> int:
    """Zero-based index of `sheet` in `names`; SheetNotFound listing every name otherwise.

    None -> 0 (the first sheet); int -> must satisfy 0 <= index < len(names) (a negative index never counts from
    the end); str -> exact title, then a unique case-insensitive match (Excel titles are case-insensitive)."""
    present = ", ".join(repr(n) for n in names) or "(none)"
    source = f" in {where}" if where else ""
    if isinstance(sheet, bool) or not isinstance(sheet, (str, int, type(None))):
        raise TypeError(f"sheet must be a title (str), a zero-based index (int) or None, not {sheet!r}")
    if sheet is None:
        if names:
            return 0
        raise SheetNotFound(f"no worksheet{source}; sheets present: {present}", sheet=sheet, present=names)
    if isinstance(sheet, int):
        if 0 <= sheet < len(names):
            return sheet
        raise SheetNotFound(
            f"sheet index {sheet} out of range{source} (valid 0-{len(names) - 1}); sheets present: {present}",
            sheet=sheet, present=names,
        )
    if sheet in names:
        return list(names).index(sheet)
    folded = [i for i, n in enumerate(names) if n.casefold() == sheet.casefold()]
    if len(folded) == 1:
        return folded[0]
    raise SheetNotFound(f"sheet {sheet!r} not found{source}; sheets present: {present}", sheet=sheet, present=names)


def _open_read_only(path: Path) -> Any:
    """openpyxl read_only + data_only workbook (cached values only; D10: never load_workbook without read_only).

    load_workbook is called by module attribute at call time (the test seam); read_only/data_only are keywords."""
    openpyxl = _require("openpyxl")
    from openpyxl.utils.exceptions import InvalidFileException

    try:
        return openpyxl.load_workbook(str(path), read_only=True, data_only=True, keep_links=False)
    except (zipfile.BadZipFile, InvalidFileException, KeyError) as exc:
        raise _not_a_workbook(path, exc) from exc


def _open_xlsb(path: Path) -> Any:
    """pyxlsb Workbook; a non-zip file becomes NotAWorkbook (pyxlsb reads the file as a zip)."""
    pyxlsb = _require("pyxlsb")
    try:
        return pyxlsb.open_workbook(str(path))
    except (zipfile.BadZipFile, KeyError) as exc:
        raise _not_a_workbook(path, exc) from exc


@contextlib.contextmanager
def _open_book(path: Path, suffix: str) -> Iterator[tuple[list[str], Any]]:
    """Open path with the reader for its (lower-cased) suffix and yield (sheet names in workbook order, workbook);
    the workbook is closed on exit. .xlsb -> pyxlsb; .xlsx/.xlsm -> openpyxl read_only + data_only."""
    if suffix == ".xlsb":
        wb = _open_xlsb(path)
        names = list(wb.sheets)
    else:
        wb = _open_read_only(path)
        names = list(wb.sheetnames)
    try:
        yield names, wb
    finally:
        wb.close()


def _not_a_workbook(path: Path, exc: BaseException) -> NotAWorkbook:
    kind = path.suffix.lower() or "workbook"
    return NotAWorkbook(
        f"{path} is not a real {kind} file ({type(exc).__name__}: {exc}); it may be an export saved as "
        "HTML or text and renamed. Open it in Excel and Save As .xlsx, then retry"
    )


def _normalise_row(values: Sequence[Any], width: int) -> list:
    """List of the row's values padded with None to `width`; a wider row is kept whole (never truncated)."""
    row = list(values)
    if len(row) < width:
        row.extend([None] * (width - len(row)))
    return row


def _iter_xlsx(path: Path, sheet: str | int | None) -> Iterator[Any]:
    """Primed generator over an .xlsx/.xlsm sheet: yields _READY after open + resolve, then one list per row."""
    with _open_book(path, ".xlsx") as (names, wb):
        ws = wb[names[_resolve_sheet(names, sheet, where=path.name)]]
        if not hasattr(ws, "reset_dimensions"):
            raise BigXlsxError(f"sheet {ws.title!r} in {path.name} is a chart sheet and holds no rows")
        ws.reset_dimensions()  # method call: the <dimension> tag is distrusted (a bad A1:A1 would yield one cell)
        yield _READY
        width = 0
        for values in ws.iter_rows(min_row=1, min_col=1, values_only=True):
            row = _normalise_row(values, width)
            width = max(width, len(row))
            yield row


def _xlsb_sheet_rows(wb: Any, index: int) -> Iterator[list]:
    """Rows of one sheet of an open pyxlsb Workbook; `index` is zero-based (pyxlsb get_sheet is 1-based)."""
    ws = wb.get_sheet(index + 1)
    try:
        width = 0
        for cells in ws.rows(sparse=False):
            row = _normalise_row([cell.v for cell in cells], width)
            width = max(width, len(row))
            yield row
    finally:
        ws.close()


def _iter_xlsb(path: Path, sheet: str | int | None) -> Iterator[Any]:
    """Primed generator over an .xlsb sheet through the pinned pyxlsb: yields _READY, then one list per row.

    Checked against the installed pinned pyxlsb source: Workbook.get_sheet(idx) is 1-based (a str is matched
    case-insensitively there, so this module resolves names itself and adds 1); dates come back as raw Excel
    serial numbers and are NOT converted (no epoch/1904 guess without her file). get_sheet copies the sheet part
    to an on-disk temporary file before rows are parsed. Proven against a hand-encoded BIFF12 test workbook only;
    UNPROVEN against an Excel-written .xlsb."""
    with _open_book(path, ".xlsb") as (names, wb):
        rows = _xlsb_sheet_rows(wb, _resolve_sheet(names, sheet, where=path.name))
        try:
            yield _READY
            yield from rows
        finally:
            rows.close()


def _uncompressed_bytes(path: Path) -> int | None:
    """Summed ZipInfo.file_size of counted members (central directory only); None when not a readable zip."""
    try:
        with zipfile.ZipFile(path) as zf:
            infos = zf.infolist()
    except (zipfile.BadZipFile, zipfile.LargeZipFile, ValueError, EOFError):
        return None
    return sum(
        i.file_size for i in infos
        if i.filename.startswith(UNCOMPRESSED_PREFIXES) and i.filename.endswith(UNCOMPRESSED_ENDINGS)
    )


def _discard(wb: Any) -> None:
    """Best-effort cleanup of an unsaved write_only workbook: close each sheet's writer and remove its temp file.

    Without it an abandoned workbook leaves a temp file until interpreter exit and an open generator at GC."""
    for ws in getattr(wb, "_sheets", ()):
        with contextlib.suppress(Exception):
            if not ws.closed:
                ws.close()
        with contextlib.suppress(Exception):
            ws._writer.cleanup()


def _atomic_save(wb: Any, path: Path, retries: int, wait_s: float) -> Path:
    """Save a write_only workbook to a same-directory temp file once, then os.replace it onto path.

    wb.save runs exactly once (a write_only workbook cannot be saved twice). Only os.replace is retried: on
    PermissionError wait `wait_s` and retry, up to `retries` times, then remove the temp file and raise
    PermissionError naming Excel and the OneDrive sync client. The temp file is removed on every failure path."""
    import tempfile

    fd, name = tempfile.mkstemp(dir=path.parent, prefix="cpa_", suffix=".xlsx")
    os.close(fd)  # closed before save: an open handle blocks os.replace on Windows
    tmp = Path(name)
    try:
        wb.save(str(tmp))
    except BaseException:
        _discard(wb)
        with contextlib.suppress(OSError):
            tmp.unlink()
        raise
    attempt = 0
    while True:
        try:
            os.replace(tmp, path)
            return path
        except PermissionError as exc:
            if attempt >= retries:
                with contextlib.suppress(OSError):
                    tmp.unlink()
                raise PermissionError(
                    exc.errno if exc.errno is not None else errno.EACCES,
                    f"could not replace {path} after {attempt + 1} attempt(s): the file is open in Excel or locked "
                    "by the OneDrive sync client. Close it in Excel (or wait for OneDrive to finish syncing) and "
                    "run again",
                    str(path),
                ) from exc
            attempt += 1
            time.sleep(wait_s)
        except BaseException:
            with contextlib.suppress(OSError):
                tmp.unlink()
            raise


# ------------------------------------------------------------------------------------------------ public API


def is_large(path: str | os.PathLike) -> bool:
    """True when the workbook is at or above STREAM_THRESHOLD_BYTES uncompressed (hard rule 11).

    Two stages (plan D2; R142 and R191 both say *uncompressed*, and sheet XML deflates 8-20x):
    1. on-disk size >= STREAM_THRESHOLD_BYTES -> True;
    2. otherwise the zip central directory is read (nothing decompressed, handle closed before return) and the
       result is True when the summed uncompressed size of members named xl/worksheets/*.xml|.bin and
       xl/sharedStrings.xml|.bin is >= STREAM_THRESHOLD_BYTES.
    A non-zip or unreadable-zip file gets stage 1's answer. Never raises for a readable file; a Windows long-path
    OSError becomes LongPathError. The suffix is not checked. The threshold is read at call time."""
    path = Path(path)
    with _guard_long_path(path):
        if path.stat().st_size >= STREAM_THRESHOLD_BYTES:
            return True
        counted = _uncompressed_bytes(path)
    return counted is not None and counted >= STREAM_THRESHOLD_BYTES


def sheet_names(path: str | os.PathLike) -> list[str]:
    """Sheet titles in workbook order, without reading any cell.

    .xlsx/.xlsm through openpyxl read_only (chart sheets included, as Excel lists them); .xlsb through pyxlsb.
    Raises UnsupportedFormat, MissingDependency, LongPathError, or the OSError of a missing/locked file."""
    path = Path(path)
    suffix = _check_suffix(path)
    with _guard_long_path(path), _open_book(path, suffix) as (names, _):
        return names


def iter_rows(path: str | os.PathLike, sheet: str | int | None) -> Iterator[list]:
    """Yield one list of cell values per worksheet row, streaming, at any file size.

    `sheet` is a title (str), a zero-based index (int, 0 <= index < sheet count) or None for the first sheet; it
    has no default (interfaces.json). Cached values only (openpyxl read_only + data_only; formulas are never read).
    Row order is kept and no row is dropped: an interior row absent from the sheet XML is yielded as a list of
    None. Each row is padded with None to the widest row seen so far and never truncated, so rows before the
    first non-empty row may be shorter (an empty list for a leading gap row). The <dimension> tag is ignored.
    A sheet with no <dimension> tag at all (openpyxl write_only output, some non-Excel writers) is pre-scanned
    once by openpyxl when the workbook opens (streamed, bounded memory), so its first row costs one extra pass;
    Excel-written files carry the tag and open at once.

    Validation is eager: the suffix, the dependency, the file and the sheet are checked when iter_rows is called,
    so UnsupportedFormat, MissingDependency, SheetNotFound, LongPathError and OSError surface at call time, not at
    the first next(). The returned object is a generator whose close() releases the workbook (on Windows an
    abandoned open handle blocks os.replace and OneDrive), so consume it fully or close it.

    .xlsb streams through pyxlsb: dates arrive as raw Excel serial numbers (not converted), every number arrives
    as a float, and each row is as wide as the sheet's dimension record. The leg is proven against a hand-encoded
    BIFF12 test workbook only, not an Excel-written .xlsb."""
    path = Path(path)
    suffix = _check_suffix(path)
    gen = _iter_xlsb(path, sheet) if suffix == ".xlsb" else _iter_xlsx(path, sheet)
    with _guard_long_path(path):
        next(gen)  # opens the workbook and resolves the sheet now; errors close it and propagate
    return gen


def row_count(path: str | os.PathLike, sheet: str | int | None) -> int:
    """Rows iter_rows would yield, counted by streaming; never trusts the dimension tag.

    Reads the whole sheet (bounded memory), so its cost is one full pass. Same sheet rules and errors as
    iter_rows."""
    gen = iter_rows(path, sheet)
    try:
        return sum(1 for _ in gen)
    finally:
        gen.close()


def to_xlsx(path: str | os.PathLike, out_dir: str | os.PathLike, *, overwrite: bool = False) -> Path:
    """C7's conversion step: stream every sheet of a .xlsb through pyxlsb into a new .xlsx and return its path.

    A .xlsx/.xlsm input is returned unchanged (no copy). For a .xlsb, every sheet is converted in workbook order
    into one openpyxl write_only workbook, one output worksheet per source sheet with its title kept (D8: no sheet
    is silently lost). Reader and writer both stream; load_workbook is never called. The output is
    `<out_dir>/<stem>.xlsx` (short, for MAX_PATH), out_dir is created if absent, and the file is written through
    a same-directory temp file + os.replace with one retry on PermissionError. An existing target is refused with
    BigXlsxError unless overwrite=True. Dates stay raw Excel serial numbers; formulas become their cached values.
    An Epic or quarantine path (EPIC_PATH_MARKERS) of any suffix is refused first with EpicPathRefused. Proven against a hand-encoded BIFF12 test workbook only."""
    path = Path(path)
    _refuse_epic(path)
    suffix = _check_suffix(path)
    if suffix != ".xlsb":
        return path
    out_dir = Path(out_dir)
    target = out_dir / (path.stem + ".xlsx")
    _require("pyxlsb")  # fail fast (MissingDependency) before any directory is created
    openpyxl = _require("openpyxl")
    with _guard_long_path(target):
        out_dir.mkdir(parents=True, exist_ok=True)
        if target.exists() and not overwrite:
            raise BigXlsxError(f"{target} already exists; pass overwrite=True (CLI: --overwrite) to replace it")
    out = openpyxl.Workbook(write_only=True)
    try:
        with contextlib.ExitStack() as stack:
            with _guard_long_path(path):  # only the open is a long-path candidate, not the conversion loop
                names, src = stack.enter_context(_open_book(path, suffix))
            for index, title in enumerate(names):
                ws = out.create_sheet(title=title)
                rows = _xlsb_sheet_rows(src, index)
                try:
                    for row in rows:
                        ws.append(row)
                finally:
                    rows.close()
    except BaseException:
        _discard(out)
        raise
    with _guard_long_path(target):
        return _atomic_save(out, target, retries=1, wait_s=0.5)


def write_rows(path: str | os.PathLike, rows: Iterable[Sequence[Any]], sheet_title: str = "Sheet1", *,
               retries: int = 1, wait_s: float = 0.5) -> Path:
    """Stream an iterable of row sequences into a new .xlsx and return the final path.

    Flat data only: an openpyxl write_only workbook with one sheet, rows appended as given (no styles; a styled
    tab beside a large sheet belongs in its own module). The workbook is saved once to a temp file in path's
    directory (tempfile.mkstemp, handle closed before save) and moved with os.replace; on PermissionError the
    replace is retried `retries` times after `wait_s` seconds, then PermissionError is raised naming Excel and
    the OneDrive sync client. The temp file never survives a failure. path's suffix must be .xlsx and its parent
    directory must exist."""
    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        raise UnsupportedFormat(f"write_rows writes .xlsx only, not {path.suffix or '(no suffix)'!r} ({path.name})")
    openpyxl = _require("openpyxl")
    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet(title=sheet_title)
    try:
        for row in rows:
            ws.append(list(row))
    except BaseException:
        _discard(wb)
        raise
    with _guard_long_path(path):
        return _atomic_save(wb, path, retries=retries, wait_s=wait_s)


# ------------------------------------------------------------------------------------------------ CLI

_HELP = __doc__.splitlines()[0] if __doc__ else "Large workbook streamer (C7)."  # -OO strips docstrings
_SHEET_HELP = (
    "sheet title, or a zero-based sheet index when the value is all digits (a sheet whose title is digits "
    "cannot be addressed by title; use `bigxlsx sheets` to see the order); default: the first sheet"
)


def _sheet_arg(value: str) -> str | int:
    return int(value) if value.isascii() and value.isdigit() else value


def _tsv_cell(value: Any) -> str:
    """TSV field: None -> ''; else str(value) with each tab, CR and LF replaced by one space."""
    if value is None:
        return ""
    return str(value).replace("\t", " ").replace("\r", " ").replace("\n", " ")


def _fail(message: str) -> int:
    print(f"bigxlsx: error: {message}", file=sys.stderr)
    return 1


def _run(args: argparse.Namespace, action) -> int:
    """Run a CLI action, turning BigXlsxError and OSError into one stderr line and exit 1 (D9)."""
    try:
        return action()
    except BigXlsxError as exc:
        return _fail(str(exc))
    except OSError as exc:
        return _fail(f"{args.path}: {exc.strerror or exc}")


def _cli_sheets(args: argparse.Namespace) -> int:
    """`bigxlsx sheets PATH`: print every sheet title, one per line, without reading cells."""
    def action() -> int:
        for name in sheet_names(args.path):
            print(name)
        return 0
    return _run(args, action)


def _cli_head(args: argparse.Namespace) -> int:
    """`bigxlsx head PATH [--sheet S] [--rows N]`: print the first N rows as TSV.

    TSV rendering: a None cell prints as an empty field; any other value prints as str(value) with each tab, CR
    and LF replaced by one space, so one sheet row is one output line. A path under an Epic or quarantine folder
    (EPIC_PATH_MARKERS) is refused with exit 1 before the file is opened."""
    def action() -> int:
        _refuse_epic(args.path)  # C13: row values never reach the terminal from an Epic path
        gen = iter_rows(args.path, args.sheet)
        try:
            for _, row in zip(range(max(args.rows, 0)), gen):
                sys.stdout.write("\t".join(_tsv_cell(v) for v in row) + "\n")
        finally:
            gen.close()
        return 0
    return _run(args, action)


def _cli_count(args: argparse.Namespace) -> int:
    """`bigxlsx count PATH [--sheet S]`: print row_count as one integer line (a full streaming pass)."""
    def action() -> int:
        print(row_count(args.path, args.sheet))
        return 0
    return _run(args, action)


def _cli_convert(args: argparse.Namespace) -> int:
    """`bigxlsx convert PATH --out DIR [--overwrite]`: to_xlsx(path, out); print the resulting path.

    to_xlsx refuses an Epic or quarantine path (EpicPathRefused -> exit 1)."""
    def action() -> int:
        print(to_xlsx(args.path, args.out, overwrite=args.overwrite))
        return 0
    return _run(args, action)


def register(subparsers: argparse._SubParsersAction) -> None:
    """Register the `bigxlsx` subcommand tree with the cpa CLI registry.

    Import-cheap (cpa/cli.py rule 3): no file, workspace or assumption access; openpyxl, pyxlsb and tempfile are
    imported only inside the functions that use them (stdlib zipfile is imported at module level)."""
    top = subparsers.add_parser(
        "bigxlsx", help=_HELP,
        description=_HELP + " `head` prints TSV: an empty cell is an empty field; tabs and line breaks inside a "
        "cell print as spaces, so one sheet row is one line.",
    )
    sub = top.add_subparsers(dest="command", required=True, title="commands")

    p = sub.add_parser("sheets", help="List sheet names in workbook order without reading cells.")
    p.add_argument("path", type=Path, help="workbook (.xlsx, .xlsm or .xlsb)")
    p.set_defaults(func=_cli_sheets)

    p = sub.add_parser("head", help="Print the first rows of a sheet as TSV (streaming).")
    p.add_argument("path", type=Path, help="workbook (.xlsx, .xlsm or .xlsb)")
    p.add_argument("--sheet", type=_sheet_arg, default=None, help=_SHEET_HELP)
    p.add_argument("--rows", type=int, default=20, help="number of rows to print (default 20)")
    p.set_defaults(func=_cli_head)

    p = sub.add_parser("count", help="Print the sheet's row count; streams the whole sheet (full-pass check).")
    p.add_argument("path", type=Path, help="workbook (.xlsx, .xlsm or .xlsb)")
    p.add_argument("--sheet", type=_sheet_arg, default=None, help=_SHEET_HELP)
    p.set_defaults(func=_cli_count)

    p = sub.add_parser("convert", help="Convert every sheet of a .xlsb into a new .xlsx (streaming).")
    p.add_argument("path", type=Path, help="workbook to convert (.xlsb; .xlsx/.xlsm are returned unchanged)")
    p.add_argument("--out", type=Path, required=True,
                   help="directory to write the converted .xlsx into (created if absent)")
    p.add_argument("--overwrite", action="store_true", help="replace an existing <stem>.xlsx in --out")
    p.set_defaults(func=_cli_convert)
