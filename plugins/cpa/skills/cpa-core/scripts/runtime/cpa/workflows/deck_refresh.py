"""Committee deck refresh (B1 BOG, B2 FC, B14 Clinical Governance): one engine, one template map per deck.

Build-list items: B1, B2, B14 (guide 5 cpa-bog-refresh / cpa-fc-refresh / cpa-governance-refresh, guide 9
deck_refresh.py, guide 13 item 20). Steps (B1): normalize the exports through C4 (crosswalk) and C5 (paste-error
check), refresh every data tab of last cycle's workbook, run C1 (verify) against last cycle's workbook, update every
mapped text, chart and table figure on the fixed slide set, regenerate the speaker notes with C10 flagging what
changed, and write a slide-by-slide change log with C11. Hard rules enforced: never add, remove or reorder slides
(R067; checked before writing and again after); every figure carries period, status, source system and as-of on
the Deck Figures tab and in the manifest, and traces to the export cell or range it came from (R109, rule 7); a
figure the exports cannot supply is a yellow flag box on its slide and a yellow MISSING cell, never estimated and
never only in the notes (R068); unmatched department labels fail loudly; a workbook over 15 MB is refused (rule
11); nothing is sent anywhere; last cycle's files are never written.

Figures are read from the exports, not from template formula cells: a `lookup` figure is the one export row whose
department (after the crosswalk) equals the map key, a `sum` figure is a whole export column. Both are written as
constants to the workbook's `Deck Figures` tab, so verify re-reads the export cell or range and ties out exactly.
Maps live in reference/template_maps/{bog,fc,gov}.yaml (FIXTURE until she confirms them against her files).
Exit codes: 0 clean, 1 produced with issues (flags, unexplained variances, lint, verification), 2 stopped.
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
import zipfile
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import Any

__all__ = [
    "DECKS", "DECK_FIGURES_TAB", "FLAG_SHAPE_PREFIX", "CHANGE_LOG_NAME", "EXIT_OK", "EXIT_ISSUES", "EXIT_STOPPED",
    "DeckRefreshError", "TemplateMapError", "SlideSetChanged", "RefreshResult",
    "load_map", "check_map", "run", "register",
]

DECKS: dict[str, str] = {"BOG": "bog", "FC": "fc", "GOV": "gov"}  # deck -> template map name
DECK_FIGURES_TAB = "Deck Figures"
DECK_FIGURES_HEADER = ("Figure", "Label", "Value", "Period", "Status", "Source system", "As of", "Source file",
                       "Source ref", "Slides")
FLAG_SHAPE_PREFIX = "CPA flag"
CHANGE_LOG_NAME = "change_log.md"
MISSING = "MISSING"
BUILTIN_PLACEHOLDERS = ("month_name", "period", "fymm", "key", "as_of")
SLIDE_KINDS = ("title", "content", "takeaway")
LOST_PART_PREFIXES = ("xl/charts/", "xl/drawings/", "xl/pivotCache/", "xl/pivotTables/")
HIGH_BAR_AUDIENCES = ("dean", "board")
REPORT = "CPA deck refresh"
EXIT_OK, EXIT_ISSUES, EXIT_STOPPED = 0, 1, 2

_PLACEHOLDER_RE = re.compile(r"\{([A-Za-z_][A-Za-z0-9_]*)\}")
_RANGE_RE = re.compile(r"(?:'((?:[^']|'')+)'|([A-Za-z0-9_.]+))!\$?[A-Za-z]{1,3}\$?\d+:\$?[A-Za-z]{1,3}\$?(\d+)")
_DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")


class DeckRefreshError(RuntimeError):
    """The refresh stopped before writing anything (bad input, base file or configuration); the message says why."""


class TemplateMapError(DeckRefreshError):
    """The template map is inconsistent with itself or with last cycle's deck or workbook."""


class SlideSetChanged(DeckRefreshError):
    """Last cycle's deck does not hold the map's slides in the map's order (R067: never add or reorder slides)."""


# ---------------------------------------------------------------- data classes


@dataclass
class _Input:
    tab: str
    path: Path
    sheet: str
    header: list[str]
    rows: list[tuple[int, list]]  # (export row number, values), non-blank rows only
    source: str
    as_of: str
    status: str
    ref_prefix: str  # "Sheet!" for a workbook, "" for CSV
    dept_column: str | None = None


@dataclass
class _Figure:
    fid: str
    label: str
    fmt: str
    source_hint: str
    value: float | None = None
    text: str = MISSING
    status: str = ""
    source: str = ""
    as_of: str = ""
    source_file: Path | None = None
    ref: str = ""
    reason: str = ""
    slides: list[int] = field(default_factory=list)
    prior: Any = None

    @property
    def missing(self) -> bool:
        return self.value is None

    def to_json(self) -> dict:
        return {"figure": self.fid, "label": self.label, "value": self.value, "text": self.text,
                "status": self.status, "source": self.source, "as_of": self.as_of,
                "source_file": str(self.source_file) if self.source_file else "", "ref": self.ref,
                "missing": self.missing, "reason": self.reason, "slides": self.slides, "prior": self.prior}


@dataclass
class RefreshResult:
    """What one refresh produced (guide 9: workbook, deck, change_log, verify) and every finding."""

    deck_name: str
    key: str
    fymm: str
    workbook: Path
    deck: Path
    change_log: Path
    verify: dict | None = None
    issues: list[dict] = field(default_factory=list)
    flags: list[dict] = field(default_factory=list)
    unexplained: list[str] = field(default_factory=list)
    figures: list[dict] = field(default_factory=list)
    lint: list[dict] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    def exit_code(self) -> int:
        """0 clean; 1 produced with issues, flags, unexplained variances or a verification that is not CLEAN."""
        dirty = self.issues or self.flags or self.unexplained or (self.verify is not None and not self.verify["clean"])
        return EXIT_ISSUES if dirty else EXIT_OK

    def to_json(self) -> dict:
        return {"deck_name": self.deck_name, "key": self.key, "fymm": self.fymm, "workbook": str(self.workbook),
                "deck": str(self.deck), "change_log": str(self.change_log), "verify": self.verify,
                "issues": self.issues, "flags": self.flags, "unexplained": self.unexplained,
                "figures": self.figures, "lint": self.lint, "warnings": self.warnings,
                "exit_code": self.exit_code()}


# ---------------------------------------------------------------- small helpers


def _issue(code: str, tab: str, cell: str, expected: str, found: str, figure: str = "") -> dict:
    return {"code": code, "tab": tab, "cell": cell, "expected": expected, "found": found, "figure": figure}


def _letters(index: int) -> str:
    out = ""
    while index:
        index, rem = divmod(index - 1, 26)
        out = chr(65 + rem) + out
    return out


def _quote(sheet: str) -> str:
    return sheet if re.fullmatch(r"[A-Za-z0-9_]+", sheet) else "'" + sheet.replace("'", "''") + "'"


def _number(value: Any) -> float | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        text = value.strip().replace(",", "").replace("$", "")
        if text.startswith("(") and text.endswith(")"):
            text = "-" + text[1:-1]
        try:
            return float(text)
        except ValueError:
            return None
    return None


def _blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _period_label(fymm: str) -> str:
    from cpa import periods

    fy, fm = periods.parse_fymm(fymm)
    return f"FY{fy % 100:02d} M{fm:02d}"


def _month_name(fymm: str) -> str:
    import calendar

    from cpa import periods

    year, month = periods.calendar_month(fymm)
    return f"{calendar.month_name[month]} {year}"


def _prev_fymm(fymm: str) -> str:
    from cpa import periods

    fy, fm = periods.parse_fymm(fymm)
    fy, fm = (fy, fm - 1) if fm > 1 else (fy - 1, 12)
    return f"{fy % 100:02d}{fm:02d}"


def _render(fmt: str, value: float) -> str:
    return fmt.format(value)


# ---------------------------------------------------------------- template map


def load_map(deck: str) -> dict:
    """The template map for BOG, FC or GOV (workspace copy first, cpa.config.template_map)."""
    from cpa import config

    name = DECKS.get(str(deck).upper())
    if name is None:
        raise DeckRefreshError(f"deck must be one of {', '.join(DECKS)}, got {deck!r}")
    return config.template_map(name)


def _slide_items(slide: dict) -> tuple[list, list, list]:
    return list(slide.get("texts") or []), list(slide.get("charts") or []), list(slide.get("tables") or [])


def check_map(tmap: dict, deck: str) -> None:
    """Raise TemplateMapError naming the first inconsistency; the map is never repaired or defaulted."""
    from cpa.pptx import notes

    where = f"template map {DECKS.get(str(deck).upper(), deck)}.yaml"
    for key in ("audience", "outbox", "outputs", "base", "inputs", "figures", "slides", "period_key"):
        if tmap.get(key) in (None, "", [], {}):
            raise TemplateMapError(f"{where}: `{key}` is missing or empty")
    if tmap["audience"] not in notes.AUDIENCES:
        raise TemplateMapError(f"{where}: audience must be one of {notes.AUDIENCES}, got {tmap['audience']!r}")
    if tmap["period_key"] not in ("fymm", "date"):
        raise TemplateMapError(f"{where}: period_key must be fymm or date, got {tmap['period_key']!r}")
    for part in ("outputs", "base"):
        for kind in ("deck", "workbook"):
            if not isinstance((tmap[part] or {}).get(kind), str):
                raise TemplateMapError(f"{where}: `{part}.{kind}` must name a file")
    inputs = tmap["inputs"]
    for tab, spec in inputs.items():
        for key in ("path", "columns"):
            if not spec.get(key):
                raise TemplateMapError(f"{where}: input {tab!r} lacks `{key}`")
        dept = spec.get("department_column")
        if dept and dept not in spec["columns"]:
            raise TemplateMapError(f"{where}: input {tab!r} department_column {dept!r} is not in its columns")
    for fid, spec in tmap["figures"].items():
        kinds = [k for k in ("lookup", "sum") if spec.get(k)]
        if len(kinds) != 1:
            raise TemplateMapError(f"{where}: figure {fid!r} needs exactly one of lookup or sum")
        if not spec.get("label") or not spec.get("format"):
            raise TemplateMapError(f"{where}: figure {fid!r} needs a label and a format")
        how = spec[kinds[0]]
        tab = how.get("tab")
        if tab not in inputs:
            raise TemplateMapError(f"{where}: figure {fid!r} reads tab {tab!r}, which is not an input")
        if how.get("column") not in inputs[tab]["columns"]:
            raise TemplateMapError(f"{where}: figure {fid!r} column {how.get('column')!r} is not in {tab!r} columns")
        if kinds[0] == "lookup" and (not how.get("key") or not inputs[tab].get("department_column")):
            raise TemplateMapError(f"{where}: lookup figure {fid!r} needs a key and a department_column on {tab!r}")
    slides = tmap["slides"]
    for n, slide in enumerate(slides, start=1):
        if not isinstance(slide, dict) or not slide.get("title"):
            raise TemplateMapError(f"{where}: slide {n} needs a title")
        if slide.get("kind", "content") not in SLIDE_KINDS:
            raise TemplateMapError(f"{where}: slide {n} kind must be one of {SLIDE_KINDS}")
        texts, charts, tables = _slide_items(slide)
        for item in texts:
            for name in _PLACEHOLDER_RE.findall(str(item.get("text", ""))):
                if name not in tmap["figures"] and name not in BUILTIN_PLACEHOLDERS:
                    raise TemplateMapError(f"{where}: slide {n} text names unknown figure {{{name}}}")
        refs = [f for c in charts for s in (c.get("series") or []) for f in (s.get("figures") or [])]
        refs += [cell.get("figure") for t in tables for cell in (t.get("cells") or [])]
        for fid in refs:
            if fid not in tmap["figures"]:
                raise TemplateMapError(f"{where}: slide {n} names unknown figure {fid!r}")
        for c in charts:
            for s in c.get("series") or []:
                if len(s.get("figures") or []) != len(c.get("categories") or []):
                    raise TemplateMapError(f"{where}: slide {n} chart {c.get('shape')!r} series {s.get('name')!r} "
                                           "needs one figure per category")
    if tmap["audience"] in HIGH_BAR_AUDIENCES and slides[-1].get("kind") != "takeaway":
        raise TemplateMapError(f"{where}: a {tmap['audience']} deck must close on a `kind: takeaway` slide "
                               "(R145/R185); the refresh never adds one")


# ---------------------------------------------------------------- inputs


def _resolve_input(pattern: str, fymm: str, ws: Path, inbox: Path) -> Path:
    from cpa import manifest

    parts = pattern.format(fymm=fymm).split("/")
    base, rest = (inbox, parts[1:]) if parts[0] == "inbox" else (ws, parts)
    folder = base.joinpath(*rest[:-1])
    hits = sorted(p for p in folder.glob(rest[-1]) if p.is_file() and not manifest.is_sidecar(p)) \
        if folder.is_dir() else []
    if not hits:
        raise DeckRefreshError(f"input not found: {folder / rest[-1]} (map pattern {pattern!r})")
    if len(hits) > 1:
        raise DeckRefreshError(f"more than one file matches {folder / rest[-1]}: {', '.join(p.name for p in hits)}")
    return hits[0]


def _read_rows(path: Path, sheet: str | None) -> tuple[str, list[list]]:
    from cpa import bigxlsx

    suffix = path.suffix.lower()
    if suffix in (".csv", ".tsv", ".txt"):
        with path.open(encoding="utf-8-sig", newline="") as fh:
            return "", [list(r) for r in csv.reader(fh, delimiter="\t" if suffix == ".tsv" else ",")]
    names = bigxlsx.sheet_names(path)
    title = sheet if sheet else names[0]
    gen = bigxlsx.iter_rows(path, title)
    try:
        return title, [list(r) for r in gen]
    finally:
        gen.close()


def _load_input(tab: str, spec: dict, path: Path, cw, issues: list[dict]) -> _Input:
    import pandas as pd

    from cpa import crosswalk, manifest, reconcile, verify

    try:
        meta = manifest.read(path)
    except manifest.MissingManifest as exc:
        raise DeckRefreshError(f"{path.name} has no manifest; every deck input must carry one (source system and "
                               "as-of, rule 7). Land it through its pull skill or `python -m cpa manifest write`") \
            from exc
    if not meta.get("source") or not meta.get("as_of"):
        raise DeckRefreshError(f"{path.name}'s manifest lacks source or as_of")
    status = str(meta.get("status") or spec.get("status") or "")
    if status not in verify.FIGURE_STATUSES:
        raise DeckRefreshError(f"{path.name}: status {status!r} is not one of {verify.FIGURE_STATUSES} (set it in "
                               "the export manifest or the map input's `status`)")
    sheet, grid = _read_rows(path, spec.get("sheet"))
    if not grid:
        raise DeckRefreshError(f"{path.name} is empty")
    header = [str(h).strip() if h is not None else "" for h in grid[0]]
    lacking = [c for c in spec["columns"] if c not in header]
    if lacking:
        raise DeckRefreshError(f"{path.name} header lacks {lacking} (found {header}); the export layout changed")
    rows = [(n, list(r) + [None] * (len(header) - len(r))) for n, r in enumerate(grid[1:], start=2)
            if not all(_blank(v) for v in r)]
    frame = pd.DataFrame([r[:len(header)] for _, r in rows], columns=header)
    for finding in reconcile.detect_paste_errors(frame[[c for c in header if c]], spec["columns"]):
        issues.append(_issue("PASTE_ERROR", tab, "", f"a clean paste of {path.name}", finding))
    dept = spec.get("department_column")
    if dept:
        col = header.index(dept)
        offenders = [str(r[col]) if r[col] is not None else "" for _, r in rows if cw.lookup(r[col]) is None]
        if offenders:
            raise crosswalk.UnmatchedDepartment(offenders, cw.path)
        for _, r in rows:
            r[col] = cw.lookup(r[col])
    return _Input(tab=tab, path=path, sheet=sheet, header=header, rows=rows, source=str(meta["source"]),
                  as_of=str(meta["as_of"]), status=status, ref_prefix=f"{_quote(sheet)}!" if sheet else "",
                  dept_column=dept or None)


def _resolve_figure(fid: str, spec: dict, inputs: dict[str, _Input]) -> _Figure:
    from cpa import crosswalk

    fig = _Figure(fid=fid, label=str(spec["label"]), fmt=str(spec["format"]), source_hint=str(spec.get("source") or ""))
    kind = "lookup" if spec.get("lookup") else "sum"
    how = spec[kind]
    src = inputs[how["tab"]]
    col = src.header.index(how["column"])
    letter = _letters(col + 1)
    fig.status, fig.source, fig.as_of, fig.source_file = src.status, src.source, src.as_of, src.path
    if kind == "lookup":
        dcol = src.header.index(str(src.dept_column))
        want = crosswalk.fold(how["key"])
        hits = [(n, r) for n, r in src.rows if crosswalk.fold(r[dcol]) == want]
        if len(hits) > 1:
            raise DeckRefreshError(f"figure {fid!r}: {len(hits)} rows of {src.path.name} are {how['key']!r} "
                                   f"(rows {', '.join(str(n) for n, _ in hits)}); one row per department expected")
        if not hits:
            fig.reason = f"no {how['key']} row in {src.path.name}"
            return fig
        n, row = hits[0]
        fig.ref = f"{src.ref_prefix}{letter}{n}"
        value = _number(row[col])
        if value is None:
            fig.reason = f"{src.path.name} {fig.ref} holds {row[col]!r}, not a number"
            return fig
        fig.value = value
    else:
        if not src.rows:
            fig.reason = f"{src.path.name} has no data rows"
            return fig
        first, last = src.rows[0][0], src.rows[-1][0]
        fig.ref = f"{src.ref_prefix}{letter}{first}:{letter}{last}"
        bad = [n for n, r in src.rows if not _blank(r[col]) and _number(r[col]) is None]
        if bad:
            fig.reason = f"{src.path.name} column {how['column']} has non-numbers on rows {bad[:5]}"
            return fig
        fig.value = sum(_number(r[col]) or 0.0 for _, r in src.rows)
    fig.text = _render(fig.fmt, fig.value)
    return fig


# ---------------------------------------------------------------- deck helpers


def _slide_title(slide) -> str:
    if slide.shapes.title is not None:
        return slide.shapes.title.text_frame.text.strip()
    for shape in slide.shapes:
        if shape.name == "Title" and shape.has_text_frame:
            return shape.text_frame.text.strip()
    return ""


def _shape(slide, name: str):
    return next((s for s in slide.shapes if s.name == name), None)


def _check_slides(slides, tmap: dict, where: str) -> None:
    want = [str(s["title"]).strip() for s in tmap["slides"]]
    have = [_slide_title(s) for s in slides]
    if len(have) != len(want):
        raise SlideSetChanged(f"{where} has {len(have)} slide(s); the map lists {len(want)}. Slides are never added, "
                              "removed or reordered (R067): fix the map or the deck, then run again")
    for n, (h, w) in enumerate(zip(have, want), start=1):
        if h.casefold() != w.casefold():
            raise SlideSetChanged(f"{where} slide {n} is titled {h!r}; the map expects {w!r} there. Slides are never "
                                  "reordered (R067): fix the map or the deck, then run again")


def _check_shapes(slides, tmap: dict) -> None:
    for n, (slide, spec) in enumerate(zip(slides, tmap["slides"]), start=1):
        texts, charts, tables = _slide_items(spec)
        for item, attr in [(t, "has_text_frame") for t in texts] + [(c, "has_chart") for c in charts] + \
                [(t, "has_table") for t in tables]:
            shape = _shape(slide, str(item.get("shape")))
            if shape is None or not getattr(shape, attr, False):
                raise TemplateMapError(f"slide {n} ({spec['title']}): no {attr[4:].replace('_', ' ')} shape named "
                                       f"{item.get('shape')!r} (Selection Pane names)")
        for c in charts:
            chart = _shape(slide, c["shape"]).chart
            cats = [str(x) for x in chart.plots[0].categories]
            if cats != [str(x) for x in c["categories"]]:
                raise TemplateMapError(f"slide {n} chart {c['shape']!r} categories are {cats}; the map lists "
                                       f"{c['categories']}")
            names = [s.name for s in chart.series]
            if names != [s["name"] for s in c["series"]]:
                raise TemplateMapError(f"slide {n} chart {c['shape']!r} series are {names}; the map lists "
                                       f"{[s['name'] for s in c['series']]}")
        for t in tables:
            table = _shape(slide, t["shape"]).table
            for cell in t.get("cells") or []:
                if not (0 <= int(cell["row"]) < len(table.rows) and 0 <= int(cell["col"]) < len(table.columns)):
                    raise TemplateMapError(f"slide {n} table {t['shape']!r} has no cell row {cell['row']} col "
                                           f"{cell['col']}")


def _set_text(text_frame, text: str) -> None:
    """Replace the text, keeping each paragraph's first-run formatting (so the refresh never drifts the format)."""
    from copy import deepcopy

    lines = text.split("\n")
    paras = list(text_frame.paragraphs)
    template_run = next((r for p in paras for r in p.runs), None)
    for i, line in enumerate(lines):
        if i < len(paras):
            para = paras[i]
        else:
            para = text_frame.add_paragraph()
        runs = list(para.runs)
        if runs:
            run = runs[0]
            for extra in runs[1:]:
                extra._r.getparent().remove(extra._r)
        else:
            run = para.add_run()
            if template_run is not None and template_run._r.find(
                    "{http://schemas.openxmlformats.org/drawingml/2006/main}rPr") is not None:
                run._r.insert(0, deepcopy(template_run._r.find(
                    "{http://schemas.openxmlformats.org/drawingml/2006/main}rPr")))
        run.text = line
    for para in paras[len(lines):]:
        para._p.getparent().remove(para._p)


def _add_flag_box(slide, prs, lines: list[str], hex_color: str, n: int) -> None:
    from pptx.dml.color import RGBColor
    from pptx.util import Emu, Inches, Pt

    width = Emu(prs.slide_width - Inches(1))
    box = slide.shapes.add_textbox(Inches(0.5), Emu(prs.slide_height - Inches(1.6)), width, Inches(1.1))
    box.name = f"{FLAG_SHAPE_PREFIX} {n}"
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor.from_string(hex_color.lstrip("#"))
    box.text_frame.word_wrap = True
    for i, line in enumerate(lines):
        para = box.text_frame.paragraphs[0] if i == 0 else box.text_frame.add_paragraph()
        run = para.add_run()
        run.text = line
        run.font.size = Pt(18)


# ---------------------------------------------------------------- workbook helpers


def _short_ranges(wb, tab: str, old_last: int, new_last: int) -> list[dict]:
    """Formula ranges on `tab` that covered the old data but now end above its last row."""
    from cpa import verify

    found = []
    if new_last <= old_last:
        return found
    for sheet in wb.worksheets:
        if sheet.title == verify.VERIFICATION_SHEET:
            continue
        for row in sheet.iter_rows():
            for cell in row:
                value = cell.value
                if not (isinstance(value, str) and value.startswith("=")):
                    continue
                for m in _RANGE_RE.finditer(value):
                    name = m.group(1).replace("''", "'") if m.group(1) is not None else m.group(2)
                    end = int(m.group(3))
                    if name == tab and old_last <= end < new_last:
                        found.append(_issue("FORMULA_RANGE_SHORT", sheet.title, cell.coordinate,
                                            f"a range reaching row {new_last} (the refreshed {tab} data)",
                                            f"{value} stops at row {end}"))
    return found


def _zip_parts(path: Path) -> set[str]:
    with zipfile.ZipFile(path) as z:
        return {n for n in z.namelist() if n.startswith(LOST_PART_PREFIXES)}


def _prior_figures(wb) -> dict[str, Any]:
    if DECK_FIGURES_TAB not in wb.sheetnames:
        return {}
    return {str(r[0]): r[2] for r in wb[DECK_FIGURES_TAB].iter_rows(min_row=2, values_only=True) if r and r[0]}


# ---------------------------------------------------------------- run


def run(deck: str, fymm: str, templates_dir: Path | str | None = None, inbox: Path | str | None = None,
        outdir: Path | str | None = None, *, key: str | None = None, base_deck: Path | str | None = None,
        base_workbook: Path | str | None = None, explanations: dict[str, str] | None = None,
        verify_output: bool = True) -> RefreshResult:
    """Refresh last cycle's deck and workbook for `deck` (BOG, FC, GOV) from `fymm`'s exports.

    Raises DeckRefreshError (or a cpa.config / crosswalk error) before anything is written when it cannot run."""
    import openpyxl
    from openpyxl.styles import PatternFill
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData

    from cpa import bigxlsx, config, crosswalk, fsutil, manifest, periods
    from cpa.pptx import brand, lint, notes

    deck = str(deck).upper()
    tmap = load_map(deck)
    check_map(tmap, deck)
    periods.parse_fymm(fymm)
    if tmap["period_key"] == "date":
        key = key or date.today().isoformat()
        if not _DATE_RE.match(key):
            raise DeckRefreshError(f"the {deck} output key is a date YYYY-MM-DD, got {key!r}")
        date.fromisoformat(key)
    else:
        key = fymm
    names = {"key": key, "fymm": fymm, "prev_fymm": _prev_fymm(fymm)}
    ws = config.workspace()
    templates = Path(templates_dir) if templates_dir is not None else ws / "templates" / tmap["outbox"]
    inbox_dir = Path(inbox) if inbox is not None else ws / "inbox"
    out_dir = Path(outdir) if outdir is not None else ws / "outbox" / tmap["outbox"] / key
    b_deck = Path(base_deck) if base_deck is not None else templates / tmap["base"]["deck"].format(**names)
    b_wb = Path(base_workbook) if base_workbook is not None else templates / tmap["base"]["workbook"].format(**names)
    for path, what in ((b_deck, "last cycle's deck"), (b_wb, "last cycle's workbook")):
        if not path.is_file():
            raise DeckRefreshError(f"{what} not found: {path} (rollover copies last month's outputs into "
                                   f"templates/{tmap['outbox']}/; or pass it explicitly)")
    if b_deck.suffix.lower() != ".pptx" or b_wb.suffix.lower() != ".xlsx":
        raise DeckRefreshError(f"last cycle's files must be .pptx and .xlsx, got {b_deck.name} and {b_wb.name}")
    if bigxlsx.is_large(b_wb):
        raise DeckRefreshError(f"{b_wb.name} is over 15 MB; it is never loaded whole (hard rule 11)")

    # 2-3: inputs and figures (nothing written yet)
    issues: list[dict] = []
    warnings: list[str] = []
    paths = {tab: _resolve_input(spec["path"], fymm, ws, inbox_dir) for tab, spec in tmap["inputs"].items()}
    cw = crosswalk.load()
    inputs: dict[str, _Input] = {}
    for tab, spec in tmap["inputs"].items():
        inputs[tab] = _load_input(tab, spec, paths[tab], cw, issues)
    figures = {fid: _resolve_figure(fid, spec, inputs) for fid, spec in tmap["figures"].items()}
    period = _period_label(fymm)
    for n, spec in enumerate(tmap["slides"], start=1):
        texts, charts, tables = _slide_items(spec)
        used = [f for t in texts for f in _PLACEHOLDER_RE.findall(str(t.get("text", "")))]
        used += [f for c in charts for s in c["series"] for f in s["figures"]]
        used += [cell["figure"] for t in tables for cell in t.get("cells") or []]
        for fid in dict.fromkeys(used):
            if fid in figures:
                figures[fid].slides.append(n)
    missing = [f for f in figures.values() if f.missing]
    as_of = max(i.as_of for i in inputs.values())
    titles = [str(spec["title"]) for spec in tmap["slides"]]
    flag_lines = {f.fid: brand.flag_text(f.label, (" and ".join(f"the {titles[n - 1]}" for n in f.slides) + " slide")
                                         if f.slides else "", f.source_hint or f.source) for f in missing}

    # 4: base files
    prs = Presentation(str(b_deck))
    if tmap.get("fixture") and not str(prs.core_properties.subject or "").startswith("FIXTURE"):
        raise DeckRefreshError(f"the {DECKS[deck]}.yaml map is still the FIXTURE map, and {b_deck.name} is not a "
                               "FIXTURE deck (document subject). Confirm the map against her deck and set "
                               "`fixture: false` before refreshing a real deck")
    slides = list(prs.slides)
    _check_slides(slides, tmap, b_deck.name)
    _check_shapes(slides, tmap)
    wb = openpyxl.load_workbook(str(b_wb))
    try:
        absent = [tab for tab in tmap["inputs"] if tab not in wb.sheetnames]
        if absent:
            raise TemplateMapError(f"{b_wb.name} has no data tab {absent}; tabs are {wb.sheetnames}")
        prior = _prior_figures(wb)
        yellow = brand.color("flag_yellow") if missing else ""  # MissingAssumption before any write (D24)
        for fig in figures.values():
            fig.prior = prior.get(fig.fid)
        explanations = dict(explanations or {})

        # 5: workbook
        for tab, src in inputs.items():
            sh = wb[tab]
            old_last = sh.max_row
            new_last = 1 + len(src.rows)
            issues.extend(_short_ranges(wb, tab, old_last, new_last))
            sh.delete_rows(1, sh.max_row)
            sh.append(src.header)
            for _, row in src.rows:
                sh.append(row[:len(src.header)])
        if DECK_FIGURES_TAB in wb.sheetnames:
            fsh = wb[DECK_FIGURES_TAB]
            fsh.delete_rows(1, fsh.max_row)
        else:
            fsh = wb.create_sheet(DECK_FIGURES_TAB)
        fsh.append(list(DECK_FIGURES_HEADER))
        cells: dict[str, str] = {}
        for r, fig in enumerate(figures.values(), start=2):
            fsh.append([fig.fid, fig.label, flag_lines[fig.fid] if fig.missing else fig.value, period, fig.status,
                        fig.source, fig.as_of, fig.source_file.name if fig.source_file else "", fig.ref,
                        ", ".join(map(str, fig.slides))])
            if fig.missing:
                fsh.cell(row=r, column=3).fill = PatternFill("solid", fgColor=yellow.lstrip("#"))
            cells[fig.fid] = f"{_quote(DECK_FIGURES_TAB)}!C{r}"
        wb.calculation.fullCalcOnLoad = True
        out_dir.mkdir(parents=True, exist_ok=True)
        out_wb = out_dir / fsutil.safe_filename(tmap["outputs"]["workbook"].format(**names))
        out_deck = out_dir / fsutil.safe_filename(tmap["outputs"]["deck"].format(**names))
        fsutil.atomic_write(out_wb, lambda tmp: wb.save(str(tmp)))
    finally:
        wb.close()
    for part in sorted(_zip_parts(b_wb) - _zip_parts(out_wb)):
        issues.append(_issue("PART_LOST", "", "", f"{part} kept from {b_wb.name}",
                             "dropped by the openpyxl round trip; re-add the chart or object by hand"))
    manifest.write(out_wb, "derived", f"{REPORT} {deck} workbook", f"{deck} {key}", as_of,
                   inputs=[*paths.values(), b_wb], period=period, fymm=fymm, deck=deck)
    for fig in figures.values():
        if not fig.missing:
            manifest.add_figure(out_wb, fig.fid, fig.value, fig.source_file, fig.ref, cell=cells[fig.fid],
                                status=fig.status, period=period)

    # 6: deck
    values = {"month_name": _month_name(fymm), "period": period, "fymm": fymm, "key": key, "as_of": as_of}
    flags: list[dict] = []
    for n, (slide, spec) in enumerate(zip(slides, tmap["slides"]), start=1):
        for shape in [s for s in slide.shapes if s.name.startswith(FLAG_SHAPE_PREFIX)]:
            shape._element.getparent().remove(shape._element)
        texts, charts, tables = _slide_items(spec)
        for item in texts:
            text = _PLACEHOLDER_RE.sub(lambda m: figures[m.group(1)].text if m.group(1) in figures
                                       else str(values[m.group(1)]), str(item["text"]))
            _set_text(_shape(slide, item["shape"]).text_frame, text)
        for c in charts:
            data = CategoryChartData()
            data.categories = list(c["categories"])
            for s in c["series"]:
                data.add_series(s["name"], [figures[f].value for f in s["figures"]])
            _shape(slide, c["shape"]).chart.replace_data(data)
        for t in tables:
            table = _shape(slide, t["shape"]).table
            for cell in t.get("cells") or []:
                _set_text(table.cell(int(cell["row"]), int(cell["col"])).text_frame, figures[cell["figure"]].text)
        lost = [figures[f] for f in dict.fromkeys(
            [f for f in figures if n in figures[f].slides]) if figures[f].missing]
        if lost:
            lines = [flag_lines[f.fid] for f in lost]
            _add_flag_box(slide, prs, lines, yellow, n)
            flags += [{"slide": n, "figure": f.fid, "label": f.label, "text": line, "reason": f.reason}
                      for f, line in zip(lost, lines)]
    for fig in missing:
        if not fig.slides:
            flags.append({"slide": None, "figure": fig.fid, "label": fig.label, "reason": fig.reason,
                          "text": flag_lines[fig.fid]})
    fsutil.atomic_write(out_deck, lambda tmp: prs.save(str(tmp)))
    _check_slides(list(Presentation(str(out_deck)).slides), tmap, out_deck.name)  # R067, after the write

    # 7: verify against last cycle's workbook
    verification = None
    unexplained: list[str] = []
    if verify_output:
        from cpa import verify

        try:
            vr = verify.build_verification_tab(out_wb, prior=b_wb)
        except verify.VerifyError as exc:
            issues.append(_issue("VERIFY_FAILED", "", "", "a Verification tab", str(exc)))
        else:
            verification = {"summary": vr.summary, "clean": vr.clean, "issues": [i.line() for i in vr.issues],
                            "figures": [r.to_json() for r in vr.figure_rows],
                            "file": manifest.to_rel(vr.verification_file)}
            unexplained = [i.figure_id for i in vr.issues
                           if i.code == "VARIANCE_ABOVE_THRESHOLD" and i.figure_id not in explanations]

    # 8: notes, lint, change log
    note_slides = []
    for n, spec in enumerate(tmap["slides"], start=1):
        on = [f for f in figures.values() if n in f.slides]
        figs = []
        for f in on:
            if f.missing:
                continue
            before = _number(f.prior)
            changed = before is not None and before != f.value
            figs.append(notes.Figure(label=f.label, value=f.text, period=period, status=f.status, changed=changed,
                                     prior_value=_render(f.fmt, before) if changed else None))
        points = [f"{f.label}: {explanations[f.fid]}" for f in on if f.fid in explanations]
        note_slides.append(notes.Slide(title=str(spec["title"]), points=points, figures=figs,
                                       flagged_metrics=[f.label for f in on if f.missing]))
    notes.write_notes(out_deck, note_slides, tmap["audience"])
    lint_issues = [i.to_json() for i in lint.lint_deck(out_deck, audience=tmap["audience"])]
    for li in lint_issues:
        issues.append(_issue(f"LINT_{li['code']}", li["location"], "", "format lint clean", li["detail"]))

    result = RefreshResult(deck_name=deck, key=key, fymm=fymm, workbook=out_wb, deck=out_deck,
                           change_log=out_dir / CHANGE_LOG_NAME, verify=verification, issues=issues, flags=flags,
                           unexplained=unexplained, figures=[f.to_json() for f in figures.values()],
                           lint=lint_issues, warnings=warnings)
    _write_change_log(result, tmap, figures, b_deck, b_wb, explanations)
    return result


# ---------------------------------------------------------------- change log


def _write_change_log(result: RefreshResult, tmap: dict, figures: dict[str, _Figure], b_deck: Path, b_wb: Path,
                      explanations: dict[str, str]) -> None:
    from cpa import bigxlsx, fsutil, manifest, verify
    from cpa.pptx import diff

    deck_changes = diff.diff_decks(b_deck, result.deck)
    lines = [f"# {result.deck_name} change log, {_period_label(result.fymm)} ({result.key})", "",
             f"- Last cycle's deck: {manifest.to_rel(b_deck)}",
             f"- Last cycle's workbook: {manifest.to_rel(b_wb)}",
             f"- Refreshed deck: {result.deck.name}; workbook: {result.workbook.name}",
             f"- Slides: {len(tmap['slides'])}, same titles in the same order (R067)",
             f"- Figures: {len(figures)}; changed since last cycle: "
             f"{sum(1 for f in figures.values() if _changed(f))}; flagged MISSING: {len(result.flags)}",
             f"- Verification: {result.verify['summary'] if result.verify else 'not run'}",
             "- Unexplained variances above threshold: " + (", ".join(result.unexplained) or "none"), ""]
    for n, spec in enumerate(tmap["slides"], start=1):
        lines.append(f"## Slide {n}: {spec['title']}")
        body = []
        for f in figures.values():
            if n not in f.slides:
                continue
            if f.missing:
                body.append(f"- FLAGGED {f.label}: {MISSING} ({f.reason})")
            elif _changed(f):
                note = f"; why: {explanations[f.fid]}" if f.fid in explanations else ""
                body.append(f"- numeric {f.label}: {_render(f.fmt, _number(f.prior))} -> {f.text}{note}")
            elif f.prior is None:
                body.append(f"- {f.label}: {f.text} (no prior value)")
        for c in deck_changes:
            if c.location == f"slide {n}" or c.location.startswith(f"slide {n} "):
                body.append(f"- {c.kind} {c.location} {c.field}: {c.before!r} -> {c.after!r}")
        lines += (body or ["- no change"]) + [""]
    lines.append("## Workbook")
    if bigxlsx.is_large(b_wb) or bigxlsx.is_large(result.workbook):
        lines.append("- over 15 MB: cell-by-cell diff skipped")
    else:
        per_tab: dict[str, int] = {}
        for c in diff.diff_workbooks(b_wb, result.workbook):
            tab = c.location.split("!")[0]
            if tab != verify.VERIFICATION_SHEET:
                per_tab[tab] = per_tab.get(tab, 0) + 1
        lines += [f"- {tab}: {count} cell change(s)" for tab, count in per_tab.items()] or ["- no change"]
    lines.append("")
    for title, rows in (("Flags", [f"- slide {f['slide']}: {f['text']}" for f in result.flags]),
                        ("Lint", [f"- {i['location']}: {i['code']} {i['detail']}" for i in result.lint]),
                        ("Issues", [f"- {i['tab'] or 'workbook'}{'!' + i['cell'] if i['cell'] else ''}: "
                                    f"{i['code']} expected {i['expected']}, found {i['found']}"
                                    for i in result.issues]),
                        ("Verification issues", [f"- {x}" for x in (result.verify or {}).get("issues", [])])):
        lines += [f"## {title}", *(rows or ["- none"]), ""]
    text = "\n".join(lines)
    fsutil.atomic_write(result.change_log, lambda tmp: tmp.write_text(text, encoding="utf-8", newline="\n"))


def _changed(f: _Figure) -> bool:
    before = _number(f.prior)
    return not f.missing and before is not None and before != f.value


# ---------------------------------------------------------------- CLI


def _cmd_run(args: argparse.Namespace) -> int:
    from cpa import config, crosswalk, fsutil, manifest

    explanations = None
    try:
        if args.explain:
            import yaml

            explanations = yaml.safe_load(Path(args.explain).read_text(encoding="utf-8-sig")) or {}
            if not isinstance(explanations, dict):
                raise DeckRefreshError(f"{args.explain} must map figure id to explanation text")
            explanations = {str(k): str(v) for k, v in explanations.items()}
        result = run(args.deck, args.month, args.templates, args.inbox, args.out, key=args.date,
                     base_deck=args.base_deck, base_workbook=args.base_workbook, explanations=explanations,
                     verify_output=not args.no_verify)
    except (DeckRefreshError, config.ConfigError, crosswalk.CrosswalkError, manifest.ManifestError,
            fsutil.UnsafeFilename, FileNotFoundError, ValueError) as exc:
        print(f"stopped: {exc}", file=sys.stderr)
        return EXIT_STOPPED
    if args.json:
        print(json.dumps(result.to_json(), indent=2, ensure_ascii=False, default=str))
        return result.exit_code()
    print(f"deck: {result.deck}")
    print(f"workbook: {result.workbook}")
    print(f"change log: {result.change_log}")
    if result.verify is not None:
        print(f"verification: {result.verify['summary']}")
    for f in result.flags:
        print(f"flag: slide {f['slide']}: {f['text']}")
    for fid in result.unexplained:
        print(f"unexplained variance above threshold: {fid} (explain it with --explain)")
    for i in result.issues:
        where = f"{i['tab']}!{i['cell']}" if i["cell"] else (i["tab"] or "workbook")
        print(f"issue: {where}: {i['code']} expected {i['expected']}, found {i['found']}")
    return result.exit_code()


def register(subparsers: argparse._SubParsersAction) -> None:
    """Register `deck_refresh run`. Import-cheap: no file access here (D03)."""
    top = subparsers.add_parser("deck_refresh", help="Refresh the BOG, FC or governance deck and workbook.")
    sub = top.add_subparsers(dest="command", required=True, title="commands")
    p = sub.add_parser(
        "run", help="Refresh last cycle's deck and workbook from this month's exports.",
        description="Refill every data tab, write the Deck Figures tab, update every mapped figure on the fixed "
                    "slide set (never adds, removes or reorders slides), verify against last cycle's workbook, "
                    "regenerate speaker notes and write change_log.md. Exit 0 clean, 1 produced with issues, "
                    "2 stopped.")
    p.add_argument("--deck", type=str.upper, choices=tuple(DECKS), required=True, help="BOG, FC or GOV.")
    p.add_argument("--month", required=True, help="Fiscal month of the exports, FYMM (e.g. 2703).")
    p.add_argument("--date", default=None, help="GOV only: output folder date YYYY-MM-DD (default today).")
    p.add_argument("--templates", type=Path, default=None,
                   help="Folder with last cycle's deck and workbook (default templates/<bog|fc|governance>/).")
    p.add_argument("--inbox", type=Path, default=None, help="Inbox folder (default the workspace inbox).")
    p.add_argument("--out", type=Path, default=None, help="Output folder (default outbox/<deck>/<key>/).")
    p.add_argument("--base-deck", dest="base_deck", type=Path, default=None, help="Last cycle's deck, explicitly.")
    p.add_argument("--base-workbook", dest="base_workbook", type=Path, default=None,
                   help="Last cycle's workbook, explicitly.")
    p.add_argument("--explain", type=Path, default=None,
                   help="YAML file mapping figure id to why it moved; explained variances go in the notes.")
    p.add_argument("--no-verify", action="store_true", help="Skip the Verification tab (run cpa verify yourself).")
    p.add_argument("--json", action="store_true", help="Print the full result as JSON.")
    p.set_defaults(func=_cmd_run)
