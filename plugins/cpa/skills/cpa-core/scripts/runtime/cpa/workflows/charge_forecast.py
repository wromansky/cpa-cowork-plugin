"""Charge forecast (B13): paste the Tableau Charges by Day export into the forecast workbook and read base, P10 and P90.

Build-list items: B13 charge forecast (guide 5.4 cpa-charge-forecast, guide 13.14), its one-time items (wire the
Prior-Year cell to a SUMIFS against Historical Data; check the Workday Calendar against the holidays file) and the
MAPE tracking table. Hard rules enforced: 7 (every figure carries period, status, source system and as-of through
the manifests and the Verification tab), 9 (a partial export is labelled, never silently used), 11 (the export is
streamed; a template over 15 MB is refused), 15 (the five input cells come from assumptions.yaml
charge_forecast.input_cells and the run stops while any is null; nothing is guessed), and "never hard-code the
prior-year value".

`run`: copy the template in memory, paste the export into Paste MTD, set the fiscal month and total workdays
(`cpa.periods.workdays`, from reference/holidays_jhm.csv) in the mapped input cells, wire the Prior-Year cell, append
the month's daily rows to Historical Data once the month has closed, save the output copy (her template is never
written), recalculate (`cpa.recalc`), read base and P10-P90, record figures and lineage (`cpa.manifest`), run
`cpa.verify`, and write forecast_summary.json. `track`: fold one summary into <workspace>/logs/forecast_mape.csv.
Cell addresses live in reference/template_maps/charge_forecast.yaml (FIXTURE until she confirms them).
Exit codes: 0 produced and clean (a partial export alone is still 0, labelled), 1 produced with issues, 2 stopped.
"""
from __future__ import annotations

import argparse
import csv
import io
import json
import math
import re
import shutil
import sys
import tempfile
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

MAP_NAME = "charge_forecast"
INPUT_KEY = ("charge_forecast", "input_cells")
REQUIRED_INPUTS = ("fiscal_month", "total_workdays")
EXPECTED_INPUTS = 5  # inventory section J: "five yellow input cells"
FISCAL_MONTH_FORMATS = ("month_start_date", "fiscal_month_int", "calendar_month_int")
OUTPUT_STEM = "CPA_Charge_Forecast_"
SUMMARY_NAME = "forecast_summary.json"
FIGURES_NAME = "forecast_figures.csv"
MAPE_PARTS = ("logs", "forecast_mape.csv")
MAPE_COLUMNS = ("fymm", "as_of", "run_at", "workday_fraction", "partial", "mtd_charges", "base_forecast",
                "p10_forecast", "p90_forecast", "actual", "ape_pct", "within_p10_p90")
REQUIRED_OUTPUTS = ("mtd_charges", "workday_fraction", "base_forecast", "p10_forecast", "p90_forecast")
FORECAST_FIGURES = ("workday_fraction", "base_forecast", "p10_forecast", "p90_forecast")
PARTIAL_LABEL = "POSSIBLE PARTIAL EXPORT"
REPORT = "CPA Charge Forecast"
EXIT_OK, EXIT_ISSUES, EXIT_STOPPED = 0, 1, 2

_REF_RE = re.compile(r"^(?:'((?:[^']|'')+)'|([^!]+))!\$?([A-Za-z]{1,3})\$?(\d+)$")
_COLS_RE = re.compile(r"^([A-Za-z]{1,3}):([A-Za-z]{1,3})$")

__all__ = [
    "ForecastError", "InputCellMapMissing", "input_cells", "fiscal_month_value", "read_export", "run", "exit_code",
    "track", "register", "OUTPUT_STEM", "SUMMARY_NAME", "FIGURES_NAME", "MAPE_PARTS", "MAPE_COLUMNS",
]


class ForecastError(RuntimeError):
    """The run stopped before producing anything (bad input, template or configuration); the message says why."""


class InputCellMapMissing(ForecastError):
    """charge_forecast.input_cells is null, holds a null, lacks a required name, or holds a bad cell reference."""


# ---------------------------------------------------------------- small helpers


def _split_ref(text: Any) -> tuple[str, str] | None:
    """'Forecast!C4' or "'Source & Notes'!A1" -> (sheet, 'C4'); anything else None."""
    if not isinstance(text, str):
        return None
    m = _REF_RE.match(text.strip())
    if not m:
        return None
    sheet = m.group(1).replace("''", "'") if m.group(1) is not None else m.group(2)
    return sheet, f"{m.group(3).upper()}{m.group(4)}"


def _quote(sheet: str) -> str:
    return sheet if re.fullmatch(r"[A-Za-z0-9_]+", sheet) else "'" + sheet.replace("'", "''") + "'"


def _col(letters: str) -> int:
    n = 0
    for ch in letters.upper():
        n = n * 26 + ord(ch) - 64
    return n


def _letters(index: int) -> str:
    out = ""
    while index:
        index, rem = divmod(index - 1, 26)
        out = chr(65 + rem) + out
    return out


def _as_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        text = value.strip()
        for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y"):
            try:
                return datetime.strptime(text, fmt).date()
            except ValueError:
                continue
    return None


def _as_number(value: Any) -> float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        text = value.strip().replace(",", "").replace("$", "")
        if text.startswith("(") and text.endswith(")"):
            text = "-" + text[1:-1]
        try:
            return float(text)
        except ValueError:
            return None
    return None


def _money(value: float) -> str:
    return f"${value:,.0f}"


def _pct(fraction: float) -> int:
    return int(math.floor(fraction * 100 + 0.5))


def _issue(code: str, tab: str, cell: str, expected: str, found: str) -> dict:
    return {"code": code, "tab": tab, "cell": cell, "expected": expected, "found": found}


def _period_label(fymm: str) -> str:
    from cpa import periods

    fy, fm = periods.parse_fymm(fymm)
    return f"FY{fy % 100:02d} M{fm:02d}"


def _month_bounds(fymm: str) -> tuple[date, date]:
    import calendar

    from cpa import periods

    year, month = periods.calendar_month(fymm)
    return date(year, month, 1), date(year, month, calendar.monthrange(year, month)[1])


# ---------------------------------------------------------------- configuration


def _ask(detail: str) -> str:
    return (f"{detail}. Stop: ask the analyst to open the forecast workbook and name the five yellow input cells on "
            "the Forecast tab (fiscal_month and total_workdays are required; use Sheet!A1 references), then propose "
            "them with `python -m cpa config propose charge_forecast.input_cells \"{fiscal_month: 'Forecast!C4', "
            "total_workdays: 'Forecast!C5', ...}\"` and have her paste the proposal from reference/proposed into "
            "assumptions.yaml (automation never edits that file)")


def input_cells() -> dict[str, tuple[str, str]]:
    """The named input cells from assumptions charge_forecast.input_cells (R042: any null stops the run)."""
    from cpa import config

    try:
        cells = config.assumption(*INPUT_KEY)
    except config.MissingAssumption as exc:
        raise InputCellMapMissing(_ask(str(exc))) from exc
    if not isinstance(cells, dict):
        raise InputCellMapMissing(_ask("charge_forecast.input_cells must be a mapping of name to cell"))
    nulls = [f"charge_forecast.input_cells.{k}" for k, v in cells.items() if v is None or str(v).strip() == ""]
    if nulls:
        raise InputCellMapMissing(_ask(f"{', '.join(nulls)} is null (unconfirmed)"))
    missing = [k for k in REQUIRED_INPUTS if k not in cells]
    if missing:
        raise InputCellMapMissing(_ask(f"charge_forecast.input_cells has no {', '.join(missing)}"))
    parsed: dict[str, tuple[str, str]] = {}
    for name, ref in cells.items():
        loc = _split_ref(ref)
        if loc is None:
            raise InputCellMapMissing(_ask(f"charge_forecast.input_cells.{name} is {ref!r}, not a Sheet!A1 reference"))
        parsed[str(name)] = loc
    return parsed


def _load_map() -> dict:
    from cpa import config

    tmap = config.template_map(MAP_NAME)
    need = ("template", "tabs", "inputs", "export", "paste", "outputs", "prior_year", "workday_calendar",
            "historical", "notes", "curve")
    missing = [k for k in need if k not in tmap]
    if missing:
        raise ForecastError(f"template map {MAP_NAME}.yaml lacks {', '.join(missing)}")
    for name in REQUIRED_OUTPUTS:
        if _split_ref(tmap["outputs"].get(name)) is None:
            raise ForecastError(f"template map {MAP_NAME}.yaml outputs.{name} is not a Sheet!A1 reference")
    if _split_ref(tmap["prior_year"].get("cell")) is None:
        raise ForecastError(f"template map {MAP_NAME}.yaml prior_year.cell is not a Sheet!A1 reference")
    if not _COLS_RE.match(str(tmap["paste"].get("clear_columns") or "")):
        raise ForecastError(f"template map {MAP_NAME}.yaml paste.clear_columns must look like A:D")
    return tmap


def fiscal_month_value(fymm: str, fmt: str) -> Any:
    """What the fiscal-month input cell receives, per the map's inputs.fiscal_month_format."""
    from cpa import periods

    fy, fm = periods.parse_fymm(fymm)
    year, month = periods.calendar_month(fymm)
    if fmt == "month_start_date":
        return datetime(year, month, 1)
    if fmt == "fiscal_month_int":
        return fm
    if fmt == "calendar_month_int":
        return month
    raise ForecastError(f"template map inputs.fiscal_month_format is {fmt!r}; use one of "
                        f"{', '.join(FISCAL_MONTH_FORMATS)}")


# ---------------------------------------------------------------- the export


@dataclass
class ExportFacts:
    """What the Charges by Day export holds for the month (all dates inside the fiscal month)."""

    path: Path
    sheet: str
    header: list
    header_row: int
    date_col: int  # 1-based
    charges_col: int  # 1-based
    rows: list[tuple[int, date, float, list]]  # (export row number, date, charges, the full row)
    as_of: date
    as_of_source: str
    source: str
    filters: str
    days_elapsed: int
    month_days: int
    partial: bool
    closed: bool
    exceeds: bool
    warnings: list[str] = field(default_factory=list)

    @property
    def total(self) -> float:
        return sum(r[2] for r in self.rows)

    @property
    def charges_ref(self) -> str:
        first, last = self.rows[0][0], self.rows[-1][0]
        col = _letters(self.charges_col)
        return f"{_quote(self.sheet)}!{col}{first}:{col}{last}"


def read_export(path: Path, fymm: str, tmap: dict, as_of: date | None = None) -> ExportFacts:
    """Stream the export (hard rule 11) and work out as-of, days elapsed, partial and closed."""
    from cpa import bigxlsx, manifest

    spec = tmap["export"]
    header_row = int(spec.get("header_row") or 1)
    names = bigxlsx.sheet_names(path)
    sheet = spec.get("sheet") or names[0]
    if sheet not in names:
        raise ForecastError(f"export {path.name} has no sheet {sheet!r} (sheets: {', '.join(names)})")
    wanted = {role: str(text).strip().casefold() for role, text in (spec.get("columns") or {}).items()}
    for role in ("date", "charges"):
        if role not in wanted:
            raise ForecastError(f"template map export.columns lacks {role}")
    start, end = _month_bounds(fymm)
    label = _period_label(fymm)
    header: list = []
    rows: list[tuple[int, date, float, list]] = []
    warnings: list[str] = []
    seen: dict[date, int] = {}
    date_col = charges_col = 0
    gen = bigxlsx.iter_rows(path, sheet)
    try:
        for r, row in enumerate(gen, start=1):
            if r < header_row:
                continue
            if r == header_row:
                header = list(row)
                found = {str(v).strip().casefold(): i + 1 for i, v in enumerate(header) if v is not None}
                absent = [spec["columns"][role] for role in ("date", "charges") if wanted[role] not in found]
                if absent:
                    raise ForecastError(f"export {path.name} row {header_row} has no column {', '.join(absent)} "
                                        f"(found: {', '.join(str(v) for v in header if v is not None)})")
                date_col, charges_col = found[wanted["date"]], found[wanted["charges"]]
                continue
            if all(v is None or str(v).strip() == "" for v in row):
                continue
            d = _as_date(row[date_col - 1] if date_col - 1 < len(row) else None)
            if d is None:
                warnings.append(f"export row {r} skipped (no date): {row[date_col - 1]!r}")
                continue
            if not start <= d <= end:
                raise ForecastError(f"export {path.name} row {r} is dated {d.isoformat()}, outside {label} "
                                    f"({start.isoformat()} to {end.isoformat()}); check the file's period")
            if d in seen:
                raise ForecastError(f"export {path.name} has {d.isoformat()} twice (rows {seen[d]} and {r})")
            seen[d] = r
            raw = row[charges_col - 1] if charges_col - 1 < len(row) else None
            value = _as_number(raw)
            if value is None:
                if raw not in (None, ""):
                    raise ForecastError(f"export {path.name} row {r} charges {raw!r} is not a number")
                warnings.append(f"export row {r} ({d.isoformat()}) has no charges; counted as 0")
                value = 0.0
            rows.append((r, d, value, list(row)))
    finally:
        gen.close()
    if not header:
        raise ForecastError(f"export {path.name} has no header row {header_row}")
    if not rows:
        raise ForecastError(f"export {path.name} has no dated rows")
    rows.sort(key=lambda t: t[0])
    meta = manifest.read(path) if manifest.exists(path) else {}
    if as_of is not None:
        as_of_source = "--as-of"
    elif _as_date(meta.get("as_of")) is not None:
        as_of, as_of_source = _as_date(meta.get("as_of")), "manifest"
    else:
        as_of, as_of_source = max(seen), "last export date"
        warnings.append("export has no manifest as_of; as-of taken from its last date, so a truncated export "
                        "cannot be detected")
    if as_of < start:
        raise ForecastError(f"as-of {as_of.isoformat()} is before {label} starts ({start.isoformat()})")
    days_elapsed = (min(as_of, end) - start).days + 1
    month_days = (end - start).days + 1
    n = len(rows)
    closed = as_of >= end and n == month_days
    if as_of > end and not closed:
        warnings.append(f"as-of is past month end but {month_days - n} day(s) are absent; history not appended")
    return ExportFacts(
        path=path, sheet=sheet, header=header, header_row=header_row, date_col=date_col, charges_col=charges_col,
        rows=rows, as_of=as_of, as_of_source=as_of_source, source=str(meta.get("source") or ""),
        filters=str(meta.get("filters") or ""), days_elapsed=days_elapsed, month_days=month_days,
        partial=n < days_elapsed, closed=closed, exceeds=n > days_elapsed or max(seen) > as_of, warnings=warnings,
    )


# ---------------------------------------------------------------- workbook steps


def _check_template(wb, tmap: dict, cells: dict[str, tuple[str, str]], path: Path) -> None:
    tabs = [str(t) for t in tmap["tabs"]]
    absent = [t for t in tabs if t not in wb.sheetnames]
    if absent:
        raise ForecastError(f"template {path.name} has no tab {', '.join(absent)} (map {MAP_NAME}.yaml tabs; found "
                            f"{', '.join(wb.sheetnames)})")
    if tmap.get("fixture", True):
        marker = tmap.get("fixture_marker") or {}
        loc = _split_ref(marker.get("cell"))
        value = wb[loc[0]][loc[1]].value if loc and loc[0] in wb.sheetnames else None
        if not str(value or "").startswith(str(marker.get("startswith") or "FIXTURE")):
            raise ForecastError(
                f"template map {MAP_NAME}.yaml is still the FIXTURE layout and {path.name} is not the fixture "
                "workbook. Confirm every address in the map against her workbook, then set fixture: false")
    for name, (sheet, _) in cells.items():
        if sheet not in wb.sheetnames:
            raise InputCellMapMissing(_ask(f"charge_forecast.input_cells.{name} names tab {sheet!r}, which "
                                           f"{path.name} does not have"))


def _paste(wb, tmap: dict, ex: ExportFacts) -> int:
    """Clear the paste columns, paste the export header and its dated rows, and return the dated rows read back."""
    spec = tmap["paste"]
    sh = wb[spec["sheet"]]
    header_row = int(spec.get("header_row") or 1)
    first_col, last_col = (_col(x) for x in _COLS_RE.match(str(spec["clear_columns"])).groups())
    width = max(i + 1 for r in [ex.header] + [row for *_, row in ex.rows] for i, v in enumerate(r) if v is not None)
    if width > last_col - first_col + 1:
        raise ForecastError(f"the export is {width} columns wide but paste.clear_columns "
                            f"{spec['clear_columns']} holds {last_col - first_col + 1}")
    for r in range(header_row, max(sh.max_row, header_row) + 1):
        for c in range(first_col, last_col + 1):
            sh.cell(row=r, column=c).value = None
    for i, v in enumerate(ex.header[:width]):
        sh.cell(row=header_row, column=first_col + i).value = v
    for n, (_, d, charges, row) in enumerate(ex.rows, start=1):
        for i in range(width):
            v = row[i] if i < len(row) else None
            if i + 1 == ex.date_col:
                v = datetime(d.year, d.month, d.day)
            elif i + 1 == ex.charges_col:
                v = charges
            cell = sh.cell(row=header_row + n, column=first_col + i)
            cell.value = v
            if isinstance(v, datetime) and cell.number_format == "General":
                cell.number_format = "yyyy-mm-dd"
    date_c = first_col + ex.date_col - 1
    return sum(1 for r in range(header_row + 1, sh.max_row + 1)
               if isinstance(sh.cell(row=r, column=date_c).value, (datetime, date)))


def _check_calendar(wb, tmap: dict, fymm: str, workdays: int, issues: list, warnings: list) -> None:
    spec = tmap["workday_calendar"]
    sh = wb[spec["sheet"]]
    start, _ = _month_bounds(fymm)
    mc, wc = _col(spec["month_column"]), _col(spec["workdays_column"])
    for r in range(int(spec.get("first_row") or 2), sh.max_row + 1):
        if _as_date(sh.cell(row=r, column=mc).value) == start:
            found = sh.cell(row=r, column=wc).value
            if _as_number(found) != workdays:
                issues.append(_issue("WORKDAY_CALENDAR_MISMATCH", sh.title, f"{spec['workdays_column']}{r}",
                                     f"{workdays} workdays (reference/holidays_jhm.csv)",
                                     f"{found} in the Workday Calendar tab (not edited; update her workbook)"))
            return
    warnings.append(f"Workday Calendar has no row for {start.isoformat()}; not checked")


def _wire_prior_year(wb, tmap: dict, issues: list, must_act: list, template: Path) -> str:
    spec = tmap["prior_year"]
    sheet, coord = _split_ref(spec["cell"])
    cell = wb[sheet][coord]
    current = cell.value
    hist = str(tmap["historical"]["sheet"])
    if isinstance(current, str) and current.startswith("=") and "SUMIFS" in current.upper() and hist in current:
        return "already wired"
    formula = spec.get("formula")
    if not formula:
        issues.append(_issue("PRIOR_YEAR_HARDCODED", sheet, coord,
                             f"a SUMIFS against {hist} (template map prior_year.formula)",
                             f"{current!r}; prior_year.formula is null in the map"))
        return "hard-coded"
    cell.value = str(formula)
    must_act.append(f"The template's Prior-Year cell ({spec['cell']}) holds {current!r}; this output wires it to a "
                    f"SUMIFS against {hist}. To make that permanent, put the same formula in {template.name}.")
    return "wired"


def _append_history(wb, tmap: dict, ex: ExportFacts, label: str, template: Path, issues: list, warnings: list,
                    must_act: list) -> int:
    spec = tmap["historical"]
    sh = wb[spec["sheet"]]
    dc, cc = _col(spec["date_column"]), _col(spec["charges_column"])
    first = int(spec.get("first_row") or 2)
    present: set[date] = set()
    last_row, last_date = first - 1, None
    for r in range(first, sh.max_row + 1):
        d = _as_date(sh.cell(row=r, column=dc).value)
        if d is not None:
            present.add(d)
            last_row, last_date = r, d if last_date is None or d > last_date else last_date
    month = {d for _, d, _, _ in ex.rows}
    overlap = month & present
    if overlap == month:
        warnings.append(f"{label} is already in Historical Data; nothing appended")
        return 0
    if overlap:
        issues.append(_issue("HISTORY_PARTIAL_OVERLAP", sh.title, "",
                             f"{label} either absent from {sh.title} or complete",
                             f"{len(overlap)} of {len(month)} days already there; nothing appended"))
        return 0
    start = ex.rows[0][1].replace(day=1)
    if last_date is not None and last_date < start - timedelta(days=1):
        warnings.append(f"{sh.title} ends {last_date.isoformat()}; appending {label} leaves a gap before it")
    fmt = sh.cell(row=last_row, column=dc).number_format if last_row >= first else "yyyy-mm-dd"
    for i, (_, d, charges, _) in enumerate(ex.rows, start=1):
        cell = sh.cell(row=last_row + i, column=dc, value=datetime(d.year, d.month, d.day))
        cell.number_format = fmt if fmt != "General" else "yyyy-mm-dd"
        sh.cell(row=last_row + i, column=cc, value=charges)
    must_act.append(f"Historical Data in this workbook now includes {label}. To carry it into next month's run, "
                    f"replace templates/{template.name} with this workbook (keep a copy of the old template).")
    return len(ex.rows)


def _append_note(wb, tmap: dict, text: str) -> None:
    spec = tmap["notes"]
    sh = wb[spec["sheet"]]
    c = _col(spec["column"])
    last = max([r for r in range(1, sh.max_row + 1) if sh.cell(row=r, column=c).value not in (None, "")] or [0])
    sh.cell(row=last + 1, column=c).value = text


def _check_curve(wb, tmap: dict, issues: list) -> None:
    spec = tmap["curve"]
    sh = wb[spec["sheet"]]
    first = int(spec.get("first_row") or 2)
    n = sum(1 for r in range(first, sh.max_row + 1) if _as_number(sh.cell(row=r, column=1).value) is not None)
    if spec.get("rows") and n != int(spec["rows"]):
        issues.append(_issue("CURVE_ROWS", sh.title, f"A{first}", f"{spec['rows']} curve rows", f"{n} rows"))


def _recalc_outputs(path: Path, tmap: dict, issues: list) -> tuple[dict[str, Any], str]:
    """Recalculate a temp copy and read the mapped outputs; ({}, reason) when that is not possible."""
    import openpyxl

    from cpa import recalc

    work = Path(tempfile.mkdtemp(prefix="cpa_cf_"))
    try:
        rc = recalc.recalc(path, work)
        if not rc.recalculated or rc.output is None:
            code = "NOT_RECALCULATED" if rc.status == recalc.NOT_RECALCULATED else "RECALC_FAILED"
            issues.append(_issue(code, "", "", "verified recalculation", rc.reason or rc.status))
            return {}, rc.reason or rc.status
        for e in rc.errors:
            issues.append(_issue("FORMULA_ERROR", e.sheet, e.cell, "a value (zero formula errors)", e.value))
        wb = openpyxl.load_workbook(str(rc.output), data_only=True)
        try:
            values: dict[str, Any] = {}
            for name, ref in (tmap["outputs"] or {}).items():
                loc = _split_ref(ref)
                if loc is None or loc[0] not in wb.sheetnames:
                    continue
                raw = wb[loc[0]][loc[1]].value
                values[name] = raw
                if name in REQUIRED_OUTPUTS:
                    num = _as_number(raw)
                    if num is None:
                        issues.append(_issue("OUTPUT_NOT_NUMERIC", loc[0], loc[1], f"a number for {name}", repr(raw)))
                    values[name] = num
            return values, ""
        finally:
            wb.close()
    finally:
        shutil.rmtree(work, ignore_errors=True)


def _write_figures_csv(path: Path, values: dict[str, Any]) -> dict[str, str]:
    """The numbers run read and reports, as a CSV grid verify can tie the workbook to; figure -> ref."""
    from cpa import fsutil

    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(["figure", "value"])
    refs = {}
    for i, name in enumerate(FORECAST_FIGURES, start=2):
        v = values.get(name)
        w.writerow([name, "" if v is None else repr(float(v))])
        refs[name] = f"B{i}"
    text = buf.getvalue()

    def write(tmp: Path) -> None:
        with tmp.open("w", encoding="utf-8", newline="") as fh:
            fh.write(text)

    fsutil.atomic_write(path, write)
    return refs


def _line(label: str, values: dict[str, Any], ex: ExportFacts, reason: str, out_name: str) -> str:
    base, p10, p90 = (values.get(k) for k in ("base_forecast", "p10_forecast", "p90_forecast"))
    frac = values.get("workday_fraction")
    if None in (base, p10, p90, frac):
        line = (f"{label} forecast: not calculated ({reason or 'outputs unreadable'}); open {out_name} in Excel "
                "for analyst review. Automatic recalculation is unsupported; do not treat cached outputs as verified.")
    else:
        line = (f"{label} forecast: base {_money(base)}, range {_money(p10)}–{_money(p90)}, "
                f"at {_pct(frac)}% of workdays.")
    if ex.partial:
        line += f" {PARTIAL_LABEL}: {len(ex.rows)} of {ex.days_elapsed} days in the export."
    return line


# ---------------------------------------------------------------- run


def run(export: Path | str, out_dir: Path | str, *, template: Path | str | None = None,
        as_of: date | None = None, verify_output: bool = True) -> dict:
    """Build the month's forecast workbook and summary; return the summary (exit_code() gives the CLI code).

    Raises ForecastError (or a cpa.config ConfigError) before writing anything when the run cannot proceed."""
    import openpyxl

    from cpa import bigxlsx, config, fsutil, manifest, periods

    export, out_dir = Path(export), Path(out_dir)
    tmap = _load_map()
    cells = input_cells()
    if not export.is_file():
        raise ForecastError(f"export not found: {export}")
    try:
        fymm = periods.fymm_from_filename(export.name)
    except periods.AmbiguousFilename as exc:
        raise ForecastError(f"{exc}; name the export charges_by_day_<FYMM>.xlsx") from exc
    label = _period_label(fymm)
    fm_value = fiscal_month_value(fymm, str((tmap.get("inputs") or {}).get("fiscal_month_format")))
    tpath = Path(template) if template is not None else config.workspace() / "templates" / str(tmap["template"])
    if not tpath.is_file():
        raise ForecastError(f"forecast template not found: {tpath} (guide 5.4: templates/{tmap['template']})")
    if tpath.suffix.lower() != ".xlsx":
        raise ForecastError(f"{tpath.name}: the forecast template must be .xlsx")
    if bigxlsx.is_large(tpath):
        raise ForecastError(f"{tpath.name} is over 15 MB; it is never loaded whole (hard rule 11)")
    ex = read_export(export, fymm, tmap, as_of)
    workdays = periods.workdays(fymm)  # MissingReference when the holidays file has no rows for the year
    run_at = manifest.utc_now_iso()
    issues: list[dict] = []
    warnings: list[str] = list(ex.warnings)
    must_act: list[str] = []
    source = ex.source or "SOURCE UNKNOWN (the export has no manifest)"
    if not ex.source:
        warnings.append(f"{export.name} has no manifest; its figures will show SOURCE UNKNOWN in verification")

    wb = openpyxl.load_workbook(str(tpath))
    try:
        _check_template(wb, tmap, cells, tpath)
        if ex.exceeds:
            issues.append(_issue("EXPORT_ROWS_EXCEED_DAYS", str(tmap["paste"]["sheet"]), "",
                                 f"at most {ex.days_elapsed} rows ({_month_bounds(fymm)[0].isoformat()} through "
                                 f"as-of {ex.as_of.isoformat()})", f"{len(ex.rows)} rows"))
        paste_rows = _paste(wb, tmap, ex)
        if paste_rows != len(ex.rows):
            issues.append(_issue("PASTE_ROWS_MISMATCH", str(tmap["paste"]["sheet"]), "",
                                 f"{len(ex.rows)} dated rows (the export)", f"{paste_rows} rows"))
        fm_sheet, fm_cell = cells["fiscal_month"]
        wb[fm_sheet][fm_cell].value = fm_value
        wd_sheet, wd_cell = cells["total_workdays"]
        wb[wd_sheet][wd_cell].value = workdays
        _check_calendar(wb, tmap, fymm, workdays, issues, warnings)
        prior_state = _wire_prior_year(wb, tmap, issues, must_act, tpath)
        _check_curve(wb, tmap, issues)
        appended = _append_history(wb, tmap, ex, label, tpath, issues, warnings, must_act) if ex.closed else 0
        if len(cells) < EXPECTED_INPUTS:
            warnings.append(f"charge_forecast.input_cells names {len(cells)} cells; the workbook has "
                            f"{EXPECTED_INPUTS} yellow input cells")
        inputs_found = {name: _jsonable(wb[s][c].value) for name, (s, c) in cells.items()}
        note = (f"Run {run_at}: {label} ({periods.sap_label(fymm)}) status forecast; MTD charges are actual as of "
                f"{ex.as_of.isoformat()} ({ex.as_of_source}); source {source}, {export.name}; "
                f"{len(ex.rows)} of {ex.days_elapsed} days.")
        if ex.partial:
            note += f" {PARTIAL_LABEL}: {len(ex.rows)} of {ex.days_elapsed} days in the export."
        _append_note(wb, tmap, note)
        wb.calculation.fullCalcOnLoad = True
        out_dir.mkdir(parents=True, exist_ok=True)
        out_path = out_dir / fsutil.safe_filename(f"{OUTPUT_STEM}{fymm}.xlsx")
        fsutil.atomic_write(out_path, lambda tmp: wb.save(str(tmp)))
    finally:
        wb.close()

    values, reason = _recalc_outputs(out_path, tmap, issues)
    figures_path = out_dir / FIGURES_NAME
    refs = _write_figures_csv(figures_path, values)
    filters = ex.filters or f"{label}; Charges by Day"
    manifest.write(figures_path, "derived", f"{REPORT} figures", filters, ex.as_of.isoformat(), inputs=[out_path])
    for name in FORECAST_FIGURES:
        manifest.add_figure(figures_path, name, values.get(name), export, ex.charges_ref, status="forecast",
                            period=label)
    manifest.write(out_path, "derived", REPORT, filters, ex.as_of.isoformat(), inputs=[export, tpath],
                   status="forecast", period=label, fymm=fymm, partial_export=ex.partial)
    manifest.add_figure(out_path, "mtd_charges", values.get("mtd_charges"), export, ex.charges_ref,
                        cell=tmap["outputs"]["mtd_charges"], status="actual", period=label)
    for name in FORECAST_FIGURES:
        manifest.add_figure(out_path, name, values.get(name), figures_path, refs[name],
                            cell=tmap["outputs"][name], status="forecast", period=label)

    verification = None
    if verify_output:
        from cpa import verify

        result = verify.build_verification_tab(out_path)
        verification = {"summary": result.summary, "clean": result.clean,
                        "issues": [i.line() for i in result.issues],
                        "file": manifest.to_rel(result.verification_file)}

    summary = {
        "fymm": fymm, "period": label, "sap_label": periods.sap_label(fymm), "status": "forecast",
        "as_of": ex.as_of.isoformat(), "as_of_source": ex.as_of_source, "run_at": run_at, "source": source,
        "export": manifest.to_rel(export), "template": manifest.to_rel(tpath), "workbook": manifest.to_rel(out_path),
        "figures_csv": manifest.to_rel(figures_path),
        "days_elapsed": ex.days_elapsed, "month_days": ex.month_days, "export_rows": len(ex.rows),
        "paste_rows": paste_rows, "partial": ex.partial, "closed": ex.closed,
        "charges_total": round(ex.total, 2), "total_workdays": workdays,
        "workdays_elapsed": _jsonable(values.get("workdays_elapsed")),
        "workday_fraction": values.get("workday_fraction"), "mtd_charges": values.get("mtd_charges"),
        "base_forecast": values.get("base_forecast"), "p10_forecast": values.get("p10_forecast"),
        "p90_forecast": values.get("p90_forecast"),
        "prior_year": prior_state, "history_appended": appended, "inputs_found": inputs_found,
        "line": _line(label, values, ex, reason, out_path.name),
        "issues": issues, "warnings": warnings, "analyst_must_act": must_act, "verification": verification,
    }
    summary_path = out_dir / SUMMARY_NAME
    text = json.dumps(summary, indent=2, ensure_ascii=False) + "\n"
    fsutil.atomic_write(summary_path, lambda tmp: tmp.write_text(text, encoding="utf-8", newline="\n"))
    manifest.write(summary_path, "derived", f"{REPORT} summary", filters, ex.as_of.isoformat(), row_count=1,
                   inputs=[out_path])
    return summary


def _jsonable(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


def exit_code(summary: dict) -> int:
    """0 clean, 1 produced with issues (including a verification that is not CLEAN)."""
    ver = summary.get("verification")
    return EXIT_ISSUES if summary.get("issues") or (ver is not None and not ver.get("clean")) else EXIT_OK


# ---------------------------------------------------------------- MAPE tracking


def _num_text(value: Any) -> str:
    if value is None or value == "":
        return ""
    text = f"{float(value):.6f}".rstrip("0").rstrip(".")
    return "0" if text in ("-0", "") else text


def _read_log(path: Path) -> list[dict]:
    if not path.is_file():
        return []
    with path.open(encoding="utf-8-sig", newline="") as fh:
        reader = csv.DictReader(fh)
        if tuple(reader.fieldnames or ()) != MAPE_COLUMNS:
            raise ForecastError(f"{path} has columns {reader.fieldnames}; expected {', '.join(MAPE_COLUMNS)}. "
                                "It was not rewritten; fix or move it and run again")
        return [dict(r) for r in reader]


def track(summary_path: Path | str) -> dict:
    """R122: fold one forecast_summary.json into <workspace>/logs/forecast_mape.csv (idempotent)."""
    from cpa import config, fsutil

    path = Path(summary_path)
    if not path.is_file():
        raise ForecastError(f"summary not found: {path} (run writes outbox/forecast/<FYMM>/{SUMMARY_NAME})")
    s = json.loads(path.read_text(encoding="utf-8-sig"))
    fymm = str(s.get("fymm") or "")
    if not re.fullmatch(r"\d{4}", fymm):
        raise ForecastError(f"{path.name} has no fymm")
    log = config.workspace().joinpath(*MAPE_PARTS)
    rows = _read_log(log)
    before = [dict(r) for r in rows]
    result = {"log": str(log), "fymm": fymm, "tracked": True, "note": ""}
    if s.get("closed"):
        actual = s.get("charges_total")
        if actual is None:
            raise ForecastError(f"{path.name} is a closed month without charges_total")
        result["note"] = f"actual {_money(float(actual))} recorded for {fymm}"
    else:
        actual = next((float(r["actual"]) for r in rows if r["fymm"] == fymm and r["actual"] != ""), None)
        if s.get("base_forecast") is None:
            result.update(tracked=False, note=f"{path.name} has no base forecast (not recalculated); nothing tracked")
        else:
            new = {
                "fymm": fymm, "as_of": str(s.get("as_of") or ""), "run_at": str(s.get("run_at") or ""),
                "workday_fraction": _num_text(s.get("workday_fraction")),
                "partial": "yes" if s.get("partial") else "no", "mtd_charges": _num_text(s.get("mtd_charges")),
                "base_forecast": _num_text(s.get("base_forecast")), "p10_forecast": _num_text(s.get("p10_forecast")),
                "p90_forecast": _num_text(s.get("p90_forecast")), "actual": "", "ape_pct": "", "within_p10_p90": "",
            }
            for i, r in enumerate(rows):
                if r["fymm"] == fymm and r["as_of"] == new["as_of"]:
                    rows[i] = new
                    break
            else:
                rows.append(new)
    if actual is not None:
        for r in rows:
            if r["fymm"] != fymm:
                continue
            base = float(r["base_forecast"])
            r["actual"] = _num_text(actual)
            r["ape_pct"] = _num_text(abs(base - float(actual)) / abs(float(actual)) * 100) if float(actual) else ""
            lo, hi = r["p10_forecast"], r["p90_forecast"]
            r["within_p10_p90"] = ("yes" if float(lo) <= float(actual) <= float(hi) else "no") if lo and hi else ""
    if rows != before:
        buf = io.StringIO()
        w = csv.DictWriter(buf, fieldnames=list(MAPE_COLUMNS), lineterminator="\n")
        w.writeheader()
        w.writerows(rows)
        text = buf.getvalue()

        def write(tmp: Path) -> None:
            with tmp.open("w", encoding="utf-8", newline="") as fh:
                fh.write(text)

        log.parent.mkdir(parents=True, exist_ok=True)
        fsutil.atomic_write(log, write)
    apes = [float(r["ape_pct"]) for r in rows if r["ape_pct"] != ""]
    result["rows"] = len(rows)
    result["mape_pct"] = sum(apes) / len(apes) if apes else None
    result["scored"] = len(apes)
    return result


# ---------------------------------------------------------------- CLI


def _cmd_run(args: argparse.Namespace) -> int:
    from cpa import config, fsutil

    try:
        as_of = date.fromisoformat(args.as_of) if args.as_of else None
    except ValueError:
        print(f"stopped: --as-of must be YYYY-MM-DD, got {args.as_of!r}", file=sys.stderr)
        return EXIT_STOPPED
    try:
        summary = run(args.export, args.out, template=args.template, as_of=as_of, verify_output=not args.no_verify)
    except (ForecastError, config.ConfigError, fsutil.UnsafeFilename, FileNotFoundError) as exc:
        print(f"stopped: {exc}", file=sys.stderr)
        return EXIT_STOPPED
    print(summary["line"])
    for i in summary["issues"]:
        where = f"{i['tab']}!{i['cell']}" if i["cell"] else (i["tab"] or "workbook")
        print(f"issue: {where}: {i['code']} expected {i['expected']}, found {i['found']}")
    if summary["verification"] is not None:
        print(f"verification: {summary['verification']['summary']}")
    for w in summary["warnings"]:
        print(f"warning: {w}")
    for a in summary["analyst_must_act"]:
        print(f"for the analyst: {a}")
    print(f"summary: {Path(args.out) / SUMMARY_NAME}")
    return exit_code(summary)


def _cmd_track(args: argparse.Namespace) -> int:
    from cpa import config

    try:
        result = track(args.summary)
    except (ForecastError, config.ConfigError, json.JSONDecodeError) as exc:
        print(f"stopped: {exc}", file=sys.stderr)
        return EXIT_STOPPED
    if result["note"]:
        print(result["note"])
    mape = result["mape_pct"]
    print(f"{result['log']}: {result['rows']} forecast row(s); MAPE "
          + ("n/a (no month closed yet)" if mape is None else f"{mape:.2f}% over {result['scored']} scored row(s)"))
    return EXIT_OK if result["tracked"] else EXIT_ISSUES


def register(subparsers: argparse._SubParsersAction) -> None:
    """Register `charge_forecast run|track`. Import-cheap: no file access here (D03)."""
    top = subparsers.add_parser("charge_forecast", help=__doc__.splitlines()[0])
    sub = top.add_subparsers(dest="command", required=True, title="commands")
    p = sub.add_parser(
        "run", help="Build the month's forecast workbook from a Charges by Day export and report base and P10-P90.",
        description="Paste the export into Paste MTD, set fiscal month and total workdays, wire the Prior-Year "
                    "SUMIFS, append history once the month closes, recalculate, verify, and write "
                    "forecast_summary.json. Exit 0 clean, 1 produced with issues, 2 stopped.")
    p.add_argument("--export", type=Path, required=True, help="inbox/tableau/charges_by_day_<FYMM>.xlsx")
    p.add_argument("--out", type=Path, required=True, help="Output folder, e.g. outbox/forecast/<FYMM>/")
    p.add_argument("--template", type=Path, default=None,
                   help="Forecast workbook (default: templates/CPA_Charge_Forecast_Rebuild.xlsx in the workspace).")
    p.add_argument("--as-of", dest="as_of", default=None,
                   help="Last day the export covers, YYYY-MM-DD (default: the export manifest's as_of).")
    p.add_argument("--no-verify", action="store_true", help="Skip the Verification tab (run cpa verify yourself).")
    p.set_defaults(func=_cmd_run)
    t = sub.add_parser("track", help="Add one forecast summary to logs/forecast_mape.csv and score closed months.",
                       description="Upsert the summary's forecast row; a closed month's summary fills the actual, "
                                   "APE and P10-P90 hit for every forecast of that month. Exit 0 tracked, 1 nothing "
                                   "to track, 2 stopped.")
    t.add_argument("--summary", type=Path, required=True,
                   help="outbox/forecast/<FYMM>/forecast_summary.json written by run.")
    t.set_defaults(func=_cmd_track)
