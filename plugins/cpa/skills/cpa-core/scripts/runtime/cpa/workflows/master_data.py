"""Departmental master data refresh (B20, guide 655 cpa-master-data): refresh the Power BI Load tab,
reconcile it against the Trend Department tab (C5, `cpa.reconcile`), and normalize department labels
through the crosswalk (C4, `cpa.crosswalk`) so an abbreviation mismatch fails loudly instead of comparing
silently.

Build-list item B20 (CPA_Cowork_Build_List.md:241, guide:655): Inputs the A4 Power BI export and the
existing master data workbook; refreshes the workbook's `Load` sheet from the export; optionally recalcs
the workbook (`cpa.recalc`, skipped with a note rather than failing when LibreOffice is absent -- recalc
is a convenience here, not this module's contract); runs C5 `reconcile.compare` between `Load` and
`Trend Department`; runs C4 `crosswalk.normalize` on the department column of both. Acceptance (R076,
R131): Load and Trend reconcile within tolerance, or every difference is listed -- `compare()` already
never drops a row, so this module only wires the two tabs together and adds the department-label gate.

R223: paste errors, out-of-period dates and department abbreviation inconsistencies are the named
failure modes for this reconciliation. `reconcile.detect_paste_errors` and `reconcile.check_dates`
findings are collected but never fatal by themselves (they are reported alongside the reconciliation,
same as `deck_refresh`'s issue lists); an unmatched department label IS fatal (`crosswalk.UnmatchedDepartment`
propagates, hard rule "unmatched department labels fail loudly").

Rule 11 (large workbooks stream through cpa.bigxlsx, never load whole): both the export and the master
data workbook are checked with `bigxlsx.is_large` before any `openpyxl`/`pandas` read; a large file is
refused rather than streamed here -- a departmental master list is not expected to approach 15 MB, and
this module stays pandas-sized like `cpa.reconcile` and `cpa.crosswalk` do.
"""
from __future__ import annotations

import argparse
import re
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from cpa import config

__all__ = [
    "LOAD_TAB", "TREND_TAB", "MasterDataError", "MasterDataResult",
    "refresh_load_tab", "read_tab", "run", "register",
]

LOAD_TAB = "Load"
TREND_TAB = "Trend Department"
_DATE_COL_RE = re.compile(r"date", re.I)


class MasterDataError(RuntimeError):
    """Base for every master data refresh failure."""


@dataclass
class MasterDataResult:
    """What one refresh produced: the reconciliation between Load and Trend Department plus findings."""

    workbook: Path
    rows_loaded: int
    reconcile: "object"  # cpa.reconcile.ReconResult
    paste_findings: list[str] = field(default_factory=list)
    date_findings: list[str] = field(default_factory=list)
    recalc_note: "str | None" = None

    @property
    def tied(self) -> bool:
        return self.reconcile.tied

    def summary(self) -> str:
        head = self.reconcile.summary()
        extra = []
        if self.paste_findings:
            extra.append(f"{len(self.paste_findings)} paste-error finding(s)")
        if self.date_findings:
            extra.append(f"{len(self.date_findings)} date finding(s)")
        if self.recalc_note:
            extra.append(self.recalc_note)
        return head if not extra else f"{head} ({'; '.join(extra)})"

    def to_json(self) -> dict:
        return {
            "workbook": str(self.workbook),
            "rows_loaded": self.rows_loaded,
            "tied": self.tied,
            "crosswalk_version": self.reconcile.crosswalk_version,
            "notes": list(self.reconcile.notes),
            "differences": self.reconcile.differences.to_dict("records"),
            "only_in_a": self.reconcile.only_in_a.to_dict("records"),
            "only_in_b": self.reconcile.only_in_b.to_dict("records"),
            "paste_findings": self.paste_findings,
            "date_findings": self.date_findings,
            "recalc_note": self.recalc_note,
        }


def _refuse_if_large(path: Path) -> None:
    from cpa import bigxlsx

    if path.suffix.lower() in bigxlsx.SUPPORTED_SUFFIXES and bigxlsx.is_large(path):
        raise MasterDataError(
            f"{path} is a large workbook (>= {bigxlsx.STREAM_THRESHOLD_BYTES} bytes uncompressed); "
            "cpa.bigxlsx streams files this size, not this module (hard rule 11). Refused."
        )


def _read_export(path: Path, sheet):
    import pandas as pd

    _refuse_if_large(path)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        return pd.read_excel(
            path, sheet_name=sheet if sheet is not None else 0, engine="openpyxl", dtype_backend="numpy_nullable"
        )
    return pd.read_csv(path, encoding="utf-8-sig", dtype_backend="numpy_nullable")


def refresh_load_tab(workbook_path: Path, export_path: Path, *, export_sheet: "str | None" = None) -> int:
    """Replace the workbook's `Load` sheet with the export's header+rows; every other sheet is untouched.

    Creates the workbook (with only a `Load` sheet) if it does not exist yet; creates the `Load` sheet if
    the workbook exists but lacks one. Returns the number of data rows written. Both files are refused if
    they are a large workbook (rule 11)."""
    import openpyxl

    from cpa import fsutil

    workbook_path = Path(workbook_path)
    export_path = Path(export_path)
    if workbook_path.is_file():
        _refuse_if_large(workbook_path)
    df = _read_export(export_path, export_sheet)
    header = list(df.columns)
    rows = df.to_numpy().tolist()

    if workbook_path.is_file():
        wb = openpyxl.load_workbook(workbook_path)
    else:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
    if LOAD_TAB in wb.sheetnames:
        del wb[LOAD_TAB]
    ws = wb.create_sheet(LOAD_TAB)
    ws.append(header)
    for row in rows:
        ws.append(row)

    def _write(tmp: Path) -> None:
        wb.save(tmp)

    fsutil.atomic_write(workbook_path, _write)
    return len(rows)


def _load_figure_records(export_path: Path, export_sheet: str | None, load_df) -> list[dict]:
    """Describe numeric Load cells and their matching incoming-export cell references."""
    from numbers import Real

    from openpyxl.utils import get_column_letter

    from cpa import manifest

    try:
        source_manifest = manifest.read(export_path)
        date.fromisoformat(str(source_manifest["as_of"]))
        source_available = bool(source_manifest.get("source"))
    except (manifest.MissingManifest, KeyError, TypeError, ValueError):
        source_manifest = {}
        source_available = False

    source_sheet = None
    if source_available and export_path.suffix.lower() in (".xlsx", ".xlsm"):
        import openpyxl

        source_wb = openpyxl.load_workbook(export_path, read_only=True, data_only=True)
        try:
            source_sheet = export_sheet or source_wb.sheetnames[0]
        finally:
            source_wb.close()

    figures = []
    for row_index, row in enumerate(load_df.itertuples(index=False, name=None), start=2):
        for column_index, value in enumerate(row, start=1):
            if isinstance(value, Real) and not isinstance(value, bool):
                source_ref = None
                if source_available:
                    source_ref = f"{get_column_letter(column_index)}{row_index}"
                    if source_sheet is not None:
                        source_ref = f"{source_sheet}!{source_ref}"
                figures.append({
                    "figure_id": f"master_data.load.r{row_index}.c{column_index}",
                    "value": value.item() if hasattr(value, "item") else value,
                    "source_ref": source_ref,
                    "cell": f"{LOAD_TAB}!{get_column_letter(column_index)}{row_index}",
                    "status": source_manifest.get("status"),
                    "period": source_manifest.get("period"),
                })
    return figures


def _record_load_lineage(workbook_path: Path, export_path: Path, figures: list[dict]) -> None:
    """Replace this refresh's Load lineage from genuine source metadata and refresh the file hash."""
    from cpa import manifest

    if manifest.exists(workbook_path):
        data = manifest.read(workbook_path)
        old_figures = data.get("figures") or {}
        retained = {}
        for fid, figure_record in old_figures.items():
            if (fid.startswith("master_data.load.") and isinstance(figure_record, dict)
                    and (not figure_record.get("cell") or str(figure_record["cell"]).startswith(f"{LOAD_TAB}!"))):
                continue
            if isinstance(figure_record, dict):
                figure_record = dict(figure_record)
                for key in ("status", "period"):
                    if figure_record.get(key) is None and data.get(key) is not None:
                        figure_record[key] = data[key]
            retained[fid] = figure_record
        inputs = list(data.get("inputs") or [])
        source_path = manifest.to_rel(export_path)
        if source_path not in inputs:
            inputs.append(source_path)
        for figure in figures:
            record = {"value": figure["value"], "cell": figure["cell"]}
            if figure.get("source_ref") is not None:
                record.update(source_file=source_path, ref=figure["source_ref"])
            if figure.get("status") is not None:
                record["status"] = figure["status"]
            if figure.get("period") is not None:
                record["period"] = figure["period"]
            retained[figure["figure_id"]] = record

        # D21 applies workbook-level status/period to every figure. Materialize the previous values
        # onto preserved figures, then omit those globals so stale Trend metadata cannot label new Load.
        standard = {"source", "report", "filters", "as_of", "exported_at", "row_count", "sha256",
                    "path", "inputs", "figures"}
        extra = {key: value for key, value in data.items()
                 if key not in standard and key not in {"status", "period", "verification"}}
        manifest.write(workbook_path, data["source"], data["report"], data["filters"], data["as_of"],
                       inputs=inputs, figures=retained, **extra)
        return

    if not figures or not any(figure.get("source_ref") is not None for figure in figures):
        return

    source_manifest = manifest.read(export_path)
    source_path = manifest.to_rel(export_path)
    manifest.write(workbook_path, "derived", "master data refresh", "", source_manifest["as_of"],
                   inputs=[export_path], **{key: source_manifest[key] for key in ("status", "period")
                                           if source_manifest.get(key) is not None})
    for figure in figures:
        manifest.add_figure(workbook_path, figure["figure_id"], figure["value"], export_path,
                            figure["source_ref"], cell=figure["cell"], status=figure.get("status"),
                            period=figure.get("period"))
    manifest.update(workbook_path)


def read_tab(workbook_path: Path, tab_name: str):
    """Read one sheet of the master data workbook as a DataFrame (refuses a large workbook, rule 11)."""
    import pandas as pd

    workbook_path = Path(workbook_path)
    _refuse_if_large(workbook_path)
    return pd.read_excel(
        workbook_path, sheet_name=tab_name, engine="openpyxl", dtype_backend="numpy_nullable"
    )


def _find_date_column(columns) -> "str | None":
    for col in columns:
        if _DATE_COL_RE.search(str(col)):
            return col
    return None


def run(
    workbook_path: Path,
    export_path: Path,
    *,
    keys,
    measures,
    dept_key: str,
    tolerance=None,
    crosswalk=None,
    period_start: "date | None" = None,
    period_end: "date | None" = None,
    export_sheet: "str | None" = None,
    recalc: bool = True,
) -> MasterDataResult:
    """B20: refresh Load, then reconcile Load vs Trend Department with the department crosswalk applied.

    Raises `crosswalk.UnmatchedDepartment` (never caught here -- fails loudly, hard rule) when either
    tab has a department label the crosswalk has no row for. Paste-error and date findings never stop
    the run; they are carried on the result alongside the reconciliation (R223)."""
    from cpa import crosswalk as crosswalk_mod
    from cpa import reconcile as reconcile_mod

    keys = list(keys)
    measures = list(measures)
    workbook_path = Path(workbook_path)

    rows_loaded = refresh_load_tab(workbook_path, export_path, export_sheet=export_sheet)
    load_df = read_tab(workbook_path, LOAD_TAB)
    export_df = _read_export(Path(export_path), export_sheet)
    load_figures = _load_figure_records(Path(export_path), export_sheet, export_df)
    trend_df = read_tab(workbook_path, TREND_TAB)

    paste_findings = reconcile_mod.detect_paste_errors(
        trend_df, expected_columns=list(load_df.columns), expected_rows=len(load_df)
    )

    date_findings: list[str] = []
    if period_start is not None and period_end is not None:
        date_col = _find_date_column(measures) or _find_date_column(keys) or _find_date_column(trend_df.columns)
        if date_col is not None and date_col in trend_df.columns:
            date_findings = reconcile_mod.check_dates(trend_df[date_col], period_start, period_end)

    cw = crosswalk_mod.load(crosswalk)
    offenders = sorted(
        set(crosswalk_mod.unmatched(load_df[dept_key], crosswalk=cw))
        | set(crosswalk_mod.unmatched(trend_df[dept_key], crosswalk=cw))
    )
    if offenders:
        raise crosswalk_mod.UnmatchedDepartment(offenders, cw.path)
    load_df = load_df.copy()
    trend_df = trend_df.copy()
    load_df[dept_key] = crosswalk_mod.normalize(load_df[dept_key], crosswalk=cw)
    trend_df[dept_key] = crosswalk_mod.normalize(trend_df[dept_key], crosswalk=cw)

    recalc_note = None
    if recalc:
        from cpa import recalc as recalc_mod

        soffice = recalc_mod.find_soffice()
        if soffice is None:
            recalc_note = "soffice not found; Trend Department recalc skipped"
        else:
            recalc_mod.recalc_in_place(workbook_path, soffice=soffice)

    tol = tolerance or reconcile_mod.tolerance_from_assumptions()
    result = reconcile_mod.compare(load_df, trend_df, keys, measures, tol, crosswalk_version=cw.version)
    _record_load_lineage(workbook_path, Path(export_path), load_figures)

    return MasterDataResult(
        workbook=workbook_path, rows_loaded=rows_loaded, reconcile=result,
        paste_findings=paste_findings, date_findings=date_findings, recalc_note=recalc_note,
    )


# ---------------------------------------------------------------- CLI


def _cmd_refresh(a: argparse.Namespace) -> int:
    from cpa import crosswalk as crosswalk_mod
    from cpa import reconcile as reconcile_mod

    keys = [k.strip() for k in a.keys.split(",") if k.strip()]
    measures = [m.strip() for m in a.measures.split(",") if m.strip()]

    if (a.tolerance_abs is None) != (a.tolerance_pct is None):
        print("--tolerance-abs and --tolerance-pct must be given together")
        return 2
    tolerance = None
    if a.tolerance_abs is not None:
        tolerance = reconcile_mod.Tolerance(abs=a.tolerance_abs, pct=a.tolerance_pct)

    period_start = date.fromisoformat(a.period_start) if a.period_start else None
    period_end = date.fromisoformat(a.period_end) if a.period_end else None

    try:
        result = run(
            Path(a.workbook), Path(a.export), keys=keys, measures=measures, dept_key=a.dept_key,
            tolerance=tolerance, crosswalk=a.crosswalk, period_start=period_start, period_end=period_end,
            export_sheet=a.export_sheet, recalc=not a.no_recalc,
        )
    except (MasterDataError, crosswalk_mod.CrosswalkError, config.MissingReference,
            config.MissingAssumption, reconcile_mod.ReconcileError) as exc:
        print(str(exc))
        return 1

    if a.out:
        from cpa import fsutil

        rows = result.reconcile.to_rows()
        import pandas as pd

        def _write(tmp: Path) -> None:
            pd.DataFrame(rows).to_csv(tmp, index=False, encoding="utf-8", lineterminator="\n")

        fsutil.atomic_write(Path(a.out), _write)

    if a.json:
        import json as _json

        print(_json.dumps(result.to_json(), default=str))
    else:
        print(f"{result.rows_loaded} row(s) loaded into {LOAD_TAB}")
        print(result.summary())
        for row in result.reconcile.to_rows():
            print(row)
        for finding in result.paste_findings:
            print(f"paste-error: {finding}")
        for finding in result.date_findings:
            print(f"date: {finding}")
        for note in result.reconcile.notes:
            print(f"note: {note}")

    return 0 if (result.tied and not result.paste_findings and not result.date_findings) else 1


def register(subparsers: argparse._SubParsersAction) -> None:
    """Register `master_data refresh`. Import-cheap: no file access here (D03)."""
    top = subparsers.add_parser("master_data", help=__doc__.splitlines()[0])
    sub = top.add_subparsers(dest="command", required=True, title="commands")

    p = sub.add_parser(
        "refresh", help="Refresh the Load tab from a Power BI export and reconcile it against Trend Department.",
        description="Exit 0 tied with no paste-error/date findings, 1 differences or findings listed, "
                    "2 on argument errors. An unmatched department label stops the run before anything "
                    "is compared (nothing written past the Load tab refresh).",
    )
    p.add_argument("--workbook", required=True, type=Path, help="Master data workbook (Load + Trend Department tabs).")
    p.add_argument("--export", required=True, type=Path, help="Power BI departmental export (A4), csv or xlsx.")
    p.add_argument("--export-sheet", default=None, help="Sheet name for an xlsx export (default: first sheet).")
    p.add_argument("--keys", required=True, help="Comma-separated shared key columns (usually the department column).")
    p.add_argument("--measures", required=True, help="Comma-separated measure columns to compare.")
    p.add_argument("--dept-key", required=True, help="A --keys column to normalize through the crosswalk first.")
    p.add_argument("--crosswalk", default=None, type=Path, help="Crosswalk CSV (default: reference/dept_crosswalk.csv).")
    p.add_argument("--tolerance-abs", type=float, default=None)
    p.add_argument("--tolerance-pct", type=float, default=None)
    p.add_argument("--period-start", default=None, help="YYYY-MM-DD; with --period-end, checks a date column in Trend.")
    p.add_argument("--period-end", default=None, help="YYYY-MM-DD")
    p.add_argument("--no-recalc", action="store_true", help="Skip the LibreOffice recalc step.")
    p.add_argument("--out", default=None, type=Path, help="Write every listed difference/one-sided row here (CSV).")
    p.add_argument("--json", action="store_true", help="Print the full result as JSON.")
    p.set_defaults(func=_cmd_refresh)
