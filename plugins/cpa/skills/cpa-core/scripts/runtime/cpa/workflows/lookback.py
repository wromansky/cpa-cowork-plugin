"""Lookback (B10-B12): cohort assembly on the MedVitals record ID, variance workbook, provider support deck.

Build-list items: B10 lookback data assembly, B11 lookback variance workbook, B12 lookback provider slides and
deck (guide 9 `lookback.py`: assemble, workbook, deck; guide 7.2 cpa-lookback readiness). Hard rules enforced:
TCC = base + supplements, never fringe (cpa.benchmarks.tcc); gross vs net collections labelled (the plan and the
Epic column are both net collections); every figure carries period, status, source system and as-of (sidecar
manifests; an input without one stops the run, since an as-of is never invented); missing data is a yellow
formula-ready placeholder with a note, never blank, never estimated; period mismatches are prorated through
cpa.periods with its label (months when both sides are whole months, days otherwise, withheld with a stated
reason when a bound is absent or the logic looks wrong, never capped); unmatched department labels fail loudly
(cpa.crosswalk); discrepancies (offer letter, prior lookback, department) are listed OPEN and never resolved;
nothing is sent or written to a source system. Framing: retention support, never a performance review.

Layouts marked FIXTURE (cohort.csv columns, bp line labels, letters.csv shape, prior Plan tab, workbook tabs)
wait on her prior lookback template, offer letter and recommendation letter (docs/OPEN_ITEMS.md NFH-P5-01/02).
"""
from __future__ import annotations

import argparse
import csv
import io
import json
import sys
from dataclasses import asdict, dataclass, field
from datetime import date
from pathlib import Path
from typing import Any

from cpa import state

__all__ = [
    "WORKFLOW", "COHORT_COLUMNS", "LINES", "JOINED_COLUMNS", "RETENTION_FRAMING", "RETENTION_SHORT",
    "FRAMING_BANNED", "LookbackError", "FramingViolation", "DeckNumberMismatch", "LintFailed", "Finding",
    "Proration", "AssembleResult", "assemble", "assemble_result", "workbook", "deck", "READINESS_RULES",
    "register",
]

WORKFLOW = "cpa-lookback"
COHORT_FILE = "cohort.csv"
LETTERS_FILE = "letters.csv"
NARRATIVE_DIR = "narrative"
PRIOR_FILE = "prior_lookback.xlsx"
JOINED_FILE = "joined.csv"  # Build List says joined.parquet; pyarrow is not pinned (D09), so CSV (DECISIONS)
REPORT_FILE = "join_report.md"
FINDINGS_FILE = "findings.json"
PNL_PREFIX = "P&L "
MAX_RECORD_ID = 31 - len(PNL_PREFIX)  # Excel's 31-character sheet-name limit
INPUT_SUFFIXES = (".xlsx", ".csv")

# FIXTURE - confirm against her file: the cohort list the cpa-lookback skill (or she) writes per cohort.
COHORT_COLUMNS = ("record_id", "provider_label", "department", "sap_provider_id", "epic_provider_id",
                  "plan_start", "plan_end", "actual_start", "actual_end")
BOUND_COLUMNS = ("plan_start", "plan_end", "actual_start", "actual_end")
LETTER_COLUMNS = ("record_id", "field", "offer_value", "source_file", "page")

RETENTION_FRAMING = (
    "Purpose: retention support. This lookback sets each provider's first-year plan beside actuals so the "
    "practice can give operational support to reach incentive targets and break the hire, ramp-up, leave cycle. "
    "It is not an evaluation of the provider."
)
RETENTION_SHORT = "Retention support: operational help to reach incentive targets; not an evaluation of the provider."
# Phrasing that turns a lookback into a performance review (R120/R209); a hit stops the deck, never rewritten.
FRAMING_BANNED = ("performance review", "underperform", "under-perform", "poor performance", "low performer",
                  "deficien", "disciplin", "probation", "terminat", "corrective action", "failed to",
                  "failing to", "performance improvement plan")

FLOW, RATE, DERIVED, OFFER = "flow", "rate", "derived", "offer"
RATE_LABEL = "not prorated (annual rate on both sides)"
STATUS = {"plan": "projected", "actual": "actual", "offer": "budget"}


@dataclass(frozen=True)
class Line:
    """One lookback line. `bp_line` is the MedVitals business plan label (FIXTURE); `actual_column` the
    canonical adapter column; `counterpart` the actual line an offer term is compared with."""

    key: str
    label: str
    kind: str
    bp_line: str = ""
    actual_source: str = ""
    actual_column: str = ""
    counterpart: str = ""


LINES: tuple[Line, ...] = (
    Line("wrvu", "wRVUs", FLOW, "wRVU", "epic", "wrvu"),
    Line("pb_encounters", "PB encounters", FLOW, "PB Encounters", "epic", "pb_encounters"),
    Line("charges", "Charges", FLOW, "Charges", "epic", "charges"),
    Line("net_collections", "Collections (net)", FLOW, "Net Collections", "epic", "net_collections"),
    Line("base_salary", "Base salary (annual)", RATE, "Base Salary", "sap", "base_salary"),
    Line("supplements", "Supplements (annual)", RATE, "Supplements", "sap", "supplements"),
    Line("tcc", "TCC (base + supplements, fringe excluded)", DERIVED),
    Line("offer_base", "Base salary vs offer letter", OFFER, counterpart="base_salary"),
    Line("offer_supplements", "Supplements vs offer letter", OFFER, counterpart="supplements"),
)
LINE_BY_KEY = {ln.key: ln for ln in LINES}
OFFER_FIELDS = {ln.key: ln for ln in LINES if ln.kind == OFFER}
MONEY_LINES = {"charges", "net_collections", "base_salary", "supplements", "tcc", "offer_base", "offer_supplements"}
DECK_LINES = tuple(ln.key for ln in LINES if ln.kind in (FLOW, RATE))
SUMMARY_LINES = ("wrvu", "charges", "net_collections")
DECK_COLUMNS = {"D": "prorated_plan", "E": "actual", "F": "variance"}

JOINED_COLUMNS = (
    "record_id", "provider_label", "department", "line", "label", "kind",
    "plan_value", "plan_file", "plan_ref", "plan_period", "plan_missing", "plan_supplier",
    "actual_value", "actual_file", "actual_ref", "actual_period", "actual_as_of", "actual_missing",
    "actual_supplier", "proration_unit", "factor", "prorated_plan", "proration_label", "variance",
)


class LookbackError(RuntimeError):
    """A data stop: nothing downstream is written. Names the file or value to fix."""


class FramingViolation(LookbackError):
    """Slide text reads as a performance review (R120/R209). Names the phrase and where it came from."""


class DeckNumberMismatch(LookbackError):
    """A deck number would differ from its recalculated workbook cell (R121)."""


class LintFailed(LookbackError):
    """The generated deck failed cpa.pptx.lint (format lint is a hard deliverable)."""


@dataclass
class Finding:
    """One gap or discrepancy. Always OPEN: this module lists, it never resolves (R072/R118)."""

    kind: str
    record_id: str
    line: str
    detail: str
    left: str = ""
    right: str = ""
    status: str = "OPEN"

    def to_json(self) -> dict:
        return asdict(self)


@dataclass
class Proration:
    """How one provider's plan maps onto the actual period. `factor` None means withheld (label says why)."""

    unit: str
    plan_n: int | None
    actual_n: int | None
    factor: float | None
    label: str
    plan_period: str
    actual_period: str


@dataclass
class AssembleResult:
    cohort: str
    joined: Path
    report: Path
    findings_file: Path
    record_ids: tuple[str, ...]
    findings: list[Finding] = field(default_factory=list)

    def to_json(self) -> dict:
        return {"cohort": self.cohort, "joined": str(self.joined), "report": str(self.report),
                "findings_file": str(self.findings_file), "record_ids": list(self.record_ids),
                "findings": [f.to_json() for f in self.findings]}


# ---------------------------------------------------------------- small helpers


def _root(root: Path | str | None) -> Path:
    from cpa import config

    return Path(root) if root is not None else config.workspace()


def _check_cohort_name(cohort: str) -> str:
    from cpa import fsutil

    try:
        return fsutil.safe_filename(str(cohort))
    except fsutil.UnsafeFilename as exc:
        raise LookbackError(f"cohort name {cohort!r} cannot name a folder or file: {exc}") from exc


def _staging(root: Path, cohort: str) -> Path:
    return root / "staging" / "lookback" / cohort


def _outbox(root: Path, cohort: str) -> Path:
    return root / "outbox" / "lookback" / cohort


def _num(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        text = str(value).strip().replace(",", "")
        if not text or text.casefold() in ("nan", "<na>", "none"):
            return None
        return float(text)
    except (TypeError, ValueError):
        return None


def _num_cell(value: Any) -> float | None:
    try:
        import pandas as pd

        if value is pd.NA:
            return None
    except ImportError:  # pragma: no cover - pandas is pinned
        pass
    n = _num(value)
    return None if n is None or n != n else n


def _text(value: Any) -> str:
    return "" if value is None else ("" if isinstance(value, float) and value != value else str(value))


def _read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as fh:
        return [{(k or "").strip(): (v or "").strip() for k, v in row.items()} for row in csv.DictReader(fh)]


def _csv_text(header: tuple[str, ...], rows: list[dict[str, Any]]) -> str:
    buf = io.StringIO()
    writer = csv.writer(buf, lineterminator="\n")
    writer.writerow(header)
    for row in rows:
        writer.writerow([_text(row.get(c)) for c in header])
    return buf.getvalue()


def _write_text(path: Path, text: str) -> Path:
    from cpa import fsutil

    path.parent.mkdir(parents=True, exist_ok=True)
    fsutil.atomic_write(path, lambda tmp: tmp.write_text(text, encoding="utf-8", newline=""))
    return path


def _inputs(folder: Path, pattern: str) -> list[Path]:
    from cpa import manifest

    if not folder.is_dir():
        return []
    return sorted(p for p in folder.glob(pattern) if p.is_file() and not manifest.is_sidecar(p)
                  and p.suffix.lower() in INPUT_SUFFIXES)


def _as_of(path: Path) -> tuple[str, str]:
    """(source system, as_of) from the input's sidecar; stops when there is none (as-of is never invented)."""
    from cpa import manifest

    try:
        data = manifest.read(path)
    except manifest.MissingManifest as exc:
        raise LookbackError(f"{path.name} has no sidecar manifest, so its as-of date is unknown; write one with "
                            "`python -m cpa manifest write` (every lookback figure carries its as-of)") from exc
    source, as_of = str(data.get("source") or "").strip(), str(data.get("as_of") or "").strip()
    try:
        parsed = date.fromisoformat(as_of)
    except ValueError as exc:
        raise LookbackError(f"{path.name} manifest as-of date is missing or invalid; use an ISO date "
                            "YYYY-MM-DD from the source export") from exc
    if not source or parsed.isoformat() != as_of:
        missing = "source system" if not source else "as-of date"
        raise LookbackError(f"{path.name} manifest {missing} is missing or invalid; use source provenance "
                            "from the export")
    return source, as_of


def _period_text(start: date | None, end: date | None) -> str:
    from cpa import periods

    if start is None or end is None:
        return "not stated"
    return f"{start.isoformat()} to {end.isoformat()} (FYMM {periods.fymm(start)}-{periods.fymm(end)})"


def _fmt_amount(value: float, line: str, *, signed: bool = False) -> str:
    text = f"{abs(value):,.0f}"
    if line in MONEY_LINES:
        text = "$" + text
    sign = "-" if value < 0 and round(abs(value)) != 0 else ("+" if signed and value > 0 else "")
    return sign + text


# ---------------------------------------------------------------- B10: assemble


def _read_cohort(root: Path, cohort: str) -> list[dict[str, str]]:
    path = _staging(root, cohort) / COHORT_FILE
    if not path.is_file():
        raise LookbackError(f"{path} not found: cohort.csv lists the cohort's MedVitals record IDs "
                            f"({', '.join(COHORT_COLUMNS)})")
    rows = _read_csv(path)
    if not rows:
        raise LookbackError(f"{path} lists no providers")
    absent = [c for c in COHORT_COLUMNS if c not in rows[0]]
    if absent:
        raise LookbackError(f"{path} is missing column(s) {', '.join(absent)}")
    seen: dict[str, str] = {}
    for row in rows:
        rid = row["record_id"]
        if not rid:
            raise LookbackError(f"{path}: a row has no record_id (the MedVitals record ID is the join key)")
        if len(rid) > MAX_RECORD_ID or any(c in rid for c in "[]:*?/\\"):
            raise LookbackError(f"{path}: record_id {rid!r} cannot name a workbook tab (at most {MAX_RECORD_ID} "
                                "characters, none of [ ] : * ? / \\)")
        if rid.casefold() in seen:
            raise LookbackError(f"{path}: record_id {rid!r} appears twice (as {seen[rid.casefold()]!r})")
        seen[rid.casefold()] = rid
    return rows


def _proration(rid: str, row: dict[str, str], findings: list[Finding]) -> Proration:
    from cpa import periods

    parsed: dict[str, date] = {}
    bad: list[str] = []
    for col in BOUND_COLUMNS:
        text = row.get(col, "").strip()
        if not text:
            bad.append(f"{col} not stated")
            continue
        try:
            parsed[col] = date.fromisoformat(text)
        except ValueError:
            bad.append(f"{col} {text!r} is not an ISO date")
    plan_period = _period_text(parsed.get("plan_start"), parsed.get("plan_end"))
    actual_period = _period_text(parsed.get("actual_start"), parsed.get("actual_end"))

    def withheld(kind: str, reason: str, detail: str) -> Proration:
        findings.append(Finding(kind, rid, "*", detail))
        return Proration("withheld", None, None, None, f"withheld ({kind}: {reason})", plan_period, actual_period)

    if bad:
        reason = "; ".join(bad)
        return withheld("PERIOD_BOUNDS", reason, f"{reason}; proration and variance withheld (an explicit start "
                                                 "and end are required on both the plan and the actual side)")
    for side in ("plan", "actual"):
        start, end = parsed[f"{side}_start"], parsed[f"{side}_end"]
        if end < start:
            return withheld("LOGIC", f"{side} end before start",
                            f"{side} period end {end} is before its start {start}; proration and variance withheld")
    try:
        plan_n = periods.months_between(parsed["plan_start"], parsed["plan_end"])
        actual_n = periods.months_between(parsed["actual_start"], parsed["actual_end"])
        unit = "months"
    except periods.PartialPeriod:
        plan_n = (parsed["plan_end"] - parsed["plan_start"]).days + 1
        actual_n = (parsed["actual_end"] - parsed["actual_start"]).days + 1
        unit = "days"
    if actual_n > plan_n:
        return withheld("LOGIC", f"actual {actual_n} {unit} longer than plan {plan_n} {unit}",
                        f"actual period ({actual_n} {unit}) is longer than the plan period ({plan_n} {unit}); "
                        "variance withheld, never capped")
    fn = periods.prorate if unit == "months" else periods.prorate_days
    factor, label = fn(1.0, plan_n, actual_n)
    return Proration(unit, plan_n, actual_n, factor, label, plan_period, actual_period)


def _prorate_value(p: Proration, value: float) -> tuple[float, str]:
    from cpa import periods

    fn = periods.prorate if p.unit == "months" else periods.prorate_days
    return fn(value, p.plan_n, p.actual_n)


def _load(name: str, path: Path):
    from cpa import sources

    return sources.get(name).load(path)


def _match(frame, column: str, wanted: str) -> list[int]:
    if column not in frame.columns or not wanted:
        return []
    keys = [(_text(v) if v is not None else "").strip().casefold() for v in frame[column].tolist()]
    return [i for i, k in enumerate(keys) if k == wanted.strip().casefold()]


def _bp_values(root: Path, rows: list[dict[str, str]], findings: list[Finding]) -> dict[str, dict]:
    """record_id -> {'file': Path, 'lines': {bp_line folded: (value, ref)}} for present bp files."""
    from cpa.sources import medvitals

    inbox = root / "inbox"
    check = medvitals.check_cohort([r["record_id"] for r in rows], inbox=inbox)
    if check.duplicated:
        raise LookbackError("business plan files collide on a record ID (case-insensitive): "
                            + "; ".join(", ".join(v) for v in check.duplicated.values()))
    for rid in check.missing:
        findings.append(Finding("MISSING_INPUT", rid, "*", f"no MedVitals business plan file bp_{rid} in "
                                                          "inbox/medvitals; every plan line is a placeholder"))
    out: dict[str, dict] = {}
    for rid, name in check.present.items():
        path = inbox / "medvitals" / name
        _as_of(path)
        frame = _load("medvitals.bp", path)
        lines: dict[str, list[tuple[float | None, str]]] = {}
        for i, (label, value) in enumerate(zip(frame["line"].tolist(), frame["year_1"].tolist())):
            lines.setdefault(_text(label).strip().casefold(), []).append((_num_cell(value), f"B{i + 2}"))
        out[rid] = {"file": path, "lines": lines}
    return out


def _epic_frame(root: Path, cohort: str, findings: list[Finding]):
    files = _inputs(root / "inbox" / "epic", f"*_{cohort}.*")
    if len(files) > 1:
        raise LookbackError(f"{len(files)} Epic extracts for cohort {cohort}: {', '.join(p.name for p in files)}; "
                            "exactly one is expected (never summed, never picked). Archive the extra one")
    if not files:
        findings.append(Finding("MISSING_INPUT", "*", "*", f"no Epic aggregate extract inbox/epic/<report>_{cohort} "
                                                           "for this cohort; every Epic actual is a placeholder"))
        return None, None, ""
    path = files[0]
    _, as_of = _as_of(path)
    frame = _load("epic.aggregate", path)
    if "provider_id" not in frame.columns:
        raise LookbackError(f"{path.name} has no provider ID column; the lookback joins Epic rows only on the "
                            "epic_provider_id cohort.csv gives for each MedVitals record ID")
    return path, frame, as_of


def _salary_frame(root: Path, cohort: str, findings: list[Finding]):
    files = _inputs(root / "inbox" / "sap", f"salary_{cohort}.*")
    if len(files) > 1:
        raise LookbackError(f"{len(files)} salary extracts for cohort {cohort}: {', '.join(p.name for p in files)}; "
                            "exactly one is expected")
    if not files:
        findings.append(Finding("MISSING_INPUT", "*", "*", f"no SAP salary extract inbox/sap/salary_{cohort}; "
                                                           "every compensation actual is a placeholder"))
        return None, None, ""
    path = files[0]
    _, as_of = _as_of(path)
    return path, _load("sap.salary", path), as_of


def _letters(root: Path, cohort: str) -> tuple[Path | None, dict[str, list[dict[str, str]]]]:
    path = _staging(root, cohort) / LETTERS_FILE
    if not path.is_file():
        return None, {}
    _as_of(path)
    rows = _read_csv(path)
    absent = [c for c in LETTER_COLUMNS if rows and c not in rows[0]]
    if absent:
        raise LookbackError(f"{path} is missing column(s) {', '.join(absent)}")
    by_id: dict[str, list[dict[str, str]]] = {}
    for i, row in enumerate(rows):
        row["_ref"] = f"C{i + 2}"
        by_id.setdefault(row["record_id"].casefold(), []).append(row)
    return path, by_id


def _prior(path: Path | None) -> dict[tuple[str, str], float]:
    if path is None:
        return {}
    from cpa import bigxlsx

    try:
        rows = bigxlsx.iter_rows(path, "Plan")
    except bigxlsx.BigXlsxError as exc:
        raise LookbackError(f"prior lookback {path.name}: {exc}") from exc
    header = [str(v).strip() if v is not None else "" for v in next(rows, [])]
    need = ("record_id", "line", "plan_value")
    if any(n not in header for n in need):
        raise LookbackError(f"prior lookback {path.name}: its Plan tab needs columns {', '.join(need)} "
                            "(FIXTURE layout until her template is mapped)")
    idx = [header.index(n) for n in need]
    out: dict[tuple[str, str], float] = {}
    for row in rows:
        vals = [row[i] if i < len(row) else None for i in idx]
        value = _num(vals[2])
        if vals[0] and vals[1] and value is not None:
            out[(str(vals[0]).strip().casefold(), str(vals[1]).strip())] = value
    return out


def assemble(cohort: str, *, root: Path | str | None = None, prior: Path | str | None = None) -> Path:
    """B10: join the cohort's inputs on the MedVitals record ID; write joined.csv, findings.json and
    join_report.md (each with a manifest) under staging/lookback/<cohort>/; return joined.csv."""
    return assemble_result(cohort, root=root, prior=prior).joined


def assemble_result(cohort: str, *, root: Path | str | None = None,
                    prior: Path | str | None = None) -> AssembleResult:
    """B10 with the findings returned. See `assemble`."""
    import pandas as pd

    from cpa import benchmarks, crosswalk, manifest, reconcile

    cohort = _check_cohort_name(cohort)
    root = _root(root)
    staging = _staging(root, cohort)
    findings: list[Finding] = []
    cohort_path = staging / COHORT_FILE
    cohort_rows = _read_cohort(root, cohort)
    _as_of(cohort_path)

    cw_path = root / "reference" / crosswalk.CROSSWALK_NAME
    cw = crosswalk.load(cw_path if cw_path.is_file() else None)
    departments = crosswalk.normalize(pd.Series([r["department"] for r in cohort_rows]), crosswalk=cw).tolist()

    bp = _bp_values(root, cohort_rows, findings)
    epic_path, epic, epic_as_of = _epic_frame(root, cohort, findings)
    sap_path, sap, sap_as_of = _salary_frame(root, cohort, findings)
    letters_path, letters = _letters(root, cohort)
    prior_path = Path(prior) if prior is not None else (staging / PRIOR_FILE if (staging / PRIOR_FILE).is_file()
                                                         else None)
    if prior_path is not None:
        _as_of(prior_path)
    prior_plan = _prior(prior_path)
    tolerance = reconcile.tolerance_from_assumptions() if (letters_path or prior_plan) else None

    joined: list[dict[str, Any]] = []
    upstream: dict[str, dict[str, Any]] = {}  # joined-manifest figures: fid -> export file and ref
    epic_notes: list[str] = []
    provider_notes: list[dict[str, str]] = []
    for crow, dept in zip(cohort_rows, departments):
        rid, label = crow["record_id"], crow["provider_label"]
        p = _proration(rid, crow, findings)
        provider_notes.append({"record_id": rid, "provider": label, "department": dept, "plan": p.plan_period,
                               "actual": p.actual_period, "proration": p.label})
        values: dict[str, dict[str, Any]] = {}

        # plan side: the MedVitals business plan, keyed on the record ID
        info = bp.get(rid)
        for ln in LINES:
            if ln.kind not in (FLOW, RATE):
                continue
            v: dict[str, Any] = {}
            if info is None:
                v.update(plan_missing=f"{ln.label} plan (no business plan file bp_{rid})",
                         plan_supplier=f"MedVitals business plan file bp_{rid} (line {ln.bp_line!r})")
            else:
                hits = info["lines"].get(ln.bp_line.casefold(), [])
                if len(hits) == 1 and hits[0][0] is not None:
                    v.update(plan_value=hits[0][0], plan_file=info["file"], plan_ref=hits[0][1])
                else:
                    why = "not in the file" if not hits else ("blank" if len(hits) == 1 else f"{len(hits)} rows")
                    v.update(plan_missing=f"{ln.label} plan (line {ln.bp_line!r} {why} in bp_{rid})",
                             plan_supplier=f"MedVitals business plan file bp_{rid} (line {ln.bp_line!r})")
                    findings.append(Finding("MISSING_INPUT", rid, ln.key, f"business plan line {ln.bp_line!r} "
                                                                           f"{why} in {info['file'].name}"))
            values[ln.key] = v

        # actual side: Epic on epic_provider_id, SAP on sap_provider_id, both taken from this record's row
        epic_idx = _match(epic, "provider_id", crow["epic_provider_id"]) if epic is not None else []
        if epic is not None and not epic_idx:
            findings.append(Finding("MISSING_INPUT", rid, "*", f"no Epic rows for provider id "
                                                               f"{crow['epic_provider_id']!r} in {epic_path.name}"))
        if len(epic_idx) > 1:
            epic_notes.append(f"{label} ({rid}): {len(epic_idx)} Epic rows summed "
                              f"(rows {', '.join(str(i + 2) for i in epic_idx)})")
        if epic_idx and "department" in epic.columns:
            labels = sorted({_text(epic["department"].iloc[i]).strip() for i in epic_idx} - {""})
            if labels:
                canon = crosswalk.normalize(pd.Series(labels), crosswalk=cw).tolist()
                other = sorted({c for c in canon if c != dept})
                if other:
                    findings.append(Finding("DEPT_MISMATCH", rid, "*", f"Epic department {', '.join(other)} "
                                                                       f"differs from the cohort department {dept}",
                                            left=dept, right=", ".join(other)))
        sap_idx = _match(sap, "provider_id", crow["sap_provider_id"]) if sap is not None else []
        if sap is not None and len(sap_idx) != 1:
            findings.append(Finding("MISSING_INPUT" if not sap_idx else "AMBIGUOUS_INPUT", rid, "*",
                                    f"{len(sap_idx)} SAP salary rows for provider id {crow['sap_provider_id']!r} "
                                    f"in {sap_path.name}; compensation actuals are placeholders"))
        for ln in LINES:
            if ln.kind not in (FLOW, RATE):
                continue
            v = values[ln.key]
            if ln.actual_source == "epic":
                supplier = (f"Epic aggregate extract inbox/epic/<report>_{cohort} (provider id "
                            f"{crow['epic_provider_id']}, column {ln.actual_column})")
                frame, idx, path, as_of = epic, epic_idx, epic_path, epic_as_of
            else:
                supplier = (f"SAP salary extract inbox/sap/salary_{cohort} (provider id {crow['sap_provider_id']}, "
                            f"column {ln.actual_column})")
                frame, idx, path, as_of = sap, (sap_idx if len(sap_idx) == 1 else []), sap_path, sap_as_of
            cells = ([_num_cell(frame[ln.actual_column].iloc[i]) for i in idx]
                     if frame is not None and ln.actual_column in frame.columns else [])
            if idx and cells and all(c is not None for c in cells):
                v.update(actual_value=sum(cells), actual_file=path, actual_as_of=as_of,
                         actual_ref=f"{ln.actual_column} rows {','.join(str(i + 2) for i in idx)}")
            else:
                why = ("no rows for the provider" if not idx else
                       "column not in the extract" if not cells else f"blank in {cells.count(None)} of {len(cells)} rows")
                v.update(actual_missing=f"{ln.label} actual ({why})", actual_supplier=supplier)
                if idx and cells:
                    findings.append(Finding("MISSING_INPUT", rid, ln.key, f"{ln.actual_column} {why} in {path.name}"))

        # derived TCC = base + supplements (fringe never included)
        base, supp = values["base_salary"], values["supplements"]
        tcc: dict[str, Any] = {}
        if base.get("plan_value") is not None and supp.get("plan_value") is not None:
            tcc["plan_value"] = benchmarks.tcc(base["plan_value"], supp["plan_value"])
        else:
            tcc.update(plan_missing="TCC plan (base or supplements plan absent)", plan_supplier="the lines above")
        if base.get("actual_value") is not None and supp.get("actual_value") is not None:
            tcc["actual_value"] = benchmarks.tcc(base["actual_value"], supp["actual_value"])
            tcc["actual_as_of"] = base.get("actual_as_of", "")
        else:
            tcc.update(actual_missing="TCC actual (base or supplements actual absent)", actual_supplier="the lines above")
        values["tcc"] = tcc

        # offer letter terms against SAP (listed, never resolved)
        offered = letters.get(rid.casefold(), [])
        known = {}
        for lrow in offered:
            fld = lrow["field"]
            if fld not in OFFER_FIELDS:
                findings.append(Finding("OFFER_STRUCTURE", rid, "offer", f"offer letter term {fld} = "
                                        f"{lrow['offer_value']} has no SAP counterpart in the lookback",
                                        left=lrow["offer_value"]))
                continue
            known.setdefault(fld, []).append(lrow)
        for key, ln in OFFER_FIELDS.items():
            v = {}
            sap_value = values[ln.counterpart].get("actual_value")
            supplier = (f"offer letter PDF; the cpa-lookback skill writes staging/lookback/{cohort}/letters.csv "
                        f"(field {key})")
            rows_for = known.get(key, [])
            value = _num(rows_for[0]["offer_value"]) if len(rows_for) == 1 else None
            if letters_path is None or not offered:
                v.update(plan_missing=f"{ln.label}: offer letter terms not extracted", plan_supplier=supplier)
                if key == "offer_base":
                    findings.append(Finding("MISSING_INPUT", rid, "offer", "no offer letter terms for this record ID "
                                            f"({LETTERS_FILE} {'absent' if letters_path is None else 'has no rows'})"))
            elif not rows_for:
                v.update(plan_missing=f"{ln.label}: the offer letter states no {key} term", plan_supplier=supplier)
                findings.append(Finding("OFFER_STRUCTURE", rid, key, f"the offer letter states no {key} term; SAP "
                                        f"{ln.counterpart} is {'not available' if sap_value is None else format(sap_value, ',.0f')}",
                                        right="" if sap_value is None else str(sap_value)))
            elif value is None:
                v.update(plan_missing=f"{ln.label}: offer value unreadable or repeated", plan_supplier=supplier)
                findings.append(Finding("AMBIGUOUS_INPUT", rid, key, f"{len(rows_for)} {key} row(s) in {LETTERS_FILE} "
                                        "without one numeric offer_value"))
            else:
                v.update(plan_value=value, plan_file=letters_path, plan_ref=rows_for[0]["_ref"])
                if sap_value is not None and not tolerance.ties(value, sap_value):
                    findings.append(Finding("COMP_VS_OFFER", rid, key, f"offer letter {key} {value:,.0f} differs "
                                            f"from SAP {ln.counterpart} {sap_value:,.0f}", left=str(value),
                                            right=str(sap_value)))
            if sap_value is not None:
                v.update(actual_value=sap_value, actual_file=values[ln.counterpart]["actual_file"],
                         actual_ref=values[ln.counterpart]["actual_ref"],
                         actual_as_of=values[ln.counterpart].get("actual_as_of", ""))
            else:
                v.update(actual_missing=values[ln.counterpart].get("actual_missing", f"{ln.label} actual"),
                         actual_supplier=values[ln.counterpart].get("actual_supplier", ""))
            values[key] = v

        # prior lookback plan figures (the MedVitals value stays the plan; the prior is shown beside it)
        for (prid, pline), pvalue in prior_plan.items():
            if prid != rid.casefold() or pline not in values:
                continue
            current = values[pline].get("plan_value")
            if current is None or not tolerance.ties(current, pvalue):
                shown = "absent" if current is None else f"{current:,.0f}"
                findings.append(Finding("PLAN_VS_PRIOR", rid, pline, f"MedVitals plan {shown} differs from the prior "
                                        f"lookback {pvalue:,.0f} ({prior_path.name})",
                                        left="" if current is None else str(current), right=str(pvalue)))

        for ln in LINES:
            v = values[ln.key]
            plan_v, actual_v = v.get("plan_value"), v.get("actual_value")
            row: dict[str, Any] = {
                "record_id": rid, "provider_label": label, "department": dept, "line": ln.key, "label": ln.label,
                "kind": ln.kind, "plan_value": plan_v, "plan_file": manifest.to_rel(v["plan_file"]) if v.get("plan_file")
                else "", "plan_ref": v.get("plan_ref", ""), "plan_period": p.plan_period,
                "plan_missing": v.get("plan_missing", ""), "plan_supplier": v.get("plan_supplier", ""),
                "actual_value": actual_v, "actual_file": manifest.to_rel(v["actual_file"]) if v.get("actual_file")
                else "", "actual_ref": v.get("actual_ref", ""), "actual_period": p.actual_period,
                "actual_as_of": v.get("actual_as_of", ""), "actual_missing": v.get("actual_missing", ""),
                "actual_supplier": v.get("actual_supplier", ""),
            }
            if ln.kind == FLOW:
                row["proration_unit"] = p.unit
                row["factor"] = p.factor
                if p.factor is not None and plan_v is not None:
                    row["prorated_plan"], row["proration_label"] = _prorate_value(p, plan_v)
                else:
                    row["prorated_plan"], row["proration_label"] = None, p.label
            else:
                row.update(proration_unit="none", factor=None, proration_label=RATE_LABEL,
                           prorated_plan=plan_v)
            if row["prorated_plan"] is not None and actual_v is not None:
                row["variance"] = actual_v - row["prorated_plan"]
            joined.append(row)
            for side in ("plan", "actual"):
                if v.get(f"{side}_file") and v.get(f"{side}_value") is not None:
                    side_key = "offer" if (ln.kind == OFFER and side == "plan") else side
                    upstream[f"{rid}.{ln.key}.{side_key}"] = {
                        "value": v[f"{side}_value"], "source_file": manifest.to_rel(v[f"{side}_file"]),
                        "ref": v.get(f"{side}_ref", ""), "status": STATUS[side_key],
                        "period": p.plan_period if side == "plan" else p.actual_period}

    joined_path = _write_text(staging / JOINED_FILE, _csv_text(JOINED_COLUMNS, joined))
    inputs = [cohort_path]
    inputs += [x for x in (epic_path, sap_path, letters_path, prior_path) if x is not None]
    inputs += [info["file"] for info in bp.values()]
    input_as_of = []
    for path in inputs:
        _, dated = _as_of(path)
        if dated:
            input_as_of.append(dated)
    if not input_as_of:
        raise LookbackError("no dated source manifest is available; an as-of date cannot be invented")
    as_of = max(input_as_of)
    manifest.write(joined_path, source="lookback", report="B10 lookback assembly", filters=f"cohort={cohort}",
                   as_of=as_of, inputs=inputs, join_key="MedVitals record ID")
    manifest.update(joined_path, figures=upstream)

    findings_path = _write_text(staging / FINDINGS_FILE,
                                json.dumps([f.to_json() for f in findings], indent=2, ensure_ascii=False) + "\n")
    report = _write_text(staging / REPORT_FILE, _report(cohort, inputs, provider_notes, epic_notes, findings))
    for path, name in ((findings_path, "B10 lookback findings"), (report, "B10 lookback join report")):
        manifest.write(path, source="lookback", report=name, filters=f"cohort={cohort}", as_of=as_of, row_count=len(findings),
                       inputs=[joined_path])
    return AssembleResult(cohort, joined_path, report, findings_path, tuple(r["record_id"] for r in cohort_rows),
                          findings)


def _report(cohort: str, inputs: list[Path], providers: list[dict[str, str]], epic_notes: list[str],
            findings: list[Finding]) -> str:
    from cpa import manifest

    lines = [f"# Lookback join report: {cohort}", "",
             "Join key: the MedVitals record ID (cohort.csv record_id). SAP and Epic rows are matched only through the "
             "ids cohort.csv gives for each record ID, never by row order or name.",
             "Every gap and flag below stays OPEN: this tool lists discrepancies and never resolves them; she decides.",
             "", "## Inputs", "", "| file | source | as of |", "|---|---|---|"]
    for path in inputs:
        try:
            data = manifest.read(path)
        except manifest.MissingManifest:
            data = {}
        lines.append(f"| {manifest.to_rel(path)} | {data.get('source', '')} | {data.get('as_of', '')} |")
    lines += ["", "## Providers and proration", "",
              "| record_id | provider | department | plan period | actual period | proration |", "|---|---|---|---|---|---|"]
    lines += [f"| {p['record_id']} | {p['provider']} | {p['department']} | {p['plan']} | {p['actual']} | "
              f"{p['proration']} |" for p in providers]
    lines += ["", "## Epic rows", ""] + ([f"- {n}" for n in epic_notes] or ["- one Epic row per provider"])
    lines += ["", f"## Findings ({len(findings)}, all OPEN, never resolved by this tool)", ""]
    if findings:
        lines += ["| kind | record_id | line | detail | left | right | status |", "|---|---|---|---|---|---|---|"]
        lines += [f"| {f.kind} | {f.record_id} | {f.line} | {f.detail} | {f.left} | {f.right} | {f.status} |"
                  for f in findings]
    else:
        lines.append("None found.")
    return "\n".join(lines) + "\n"


# ---------------------------------------------------------------- B11: workbook


def _read_findings(root: Path, cohort: str) -> list[Finding]:
    path = _staging(root, cohort) / FINDINGS_FILE
    if not path.is_file():
        raise LookbackError(f"{path} not found; run `python -m cpa lookback assemble --cohort {cohort}` first")
    return [Finding(**f) for f in json.loads(path.read_text(encoding="utf-8"))]


def _read_joined(root: Path, cohort: str) -> tuple[Path, list[dict[str, str]]]:
    path = _staging(root, cohort) / JOINED_FILE
    if not path.is_file():
        raise LookbackError(f"{path} not found; run `python -m cpa lookback assemble --cohort {cohort}` first")
    return path, _read_csv(path)


def _joined_ref(row_index: int, column: str) -> str:
    from openpyxl.utils import get_column_letter

    return f"{get_column_letter(JOINED_COLUMNS.index(column) + 1)}{row_index + 2}"


def workbook(cohort: str, *, root: Path | str | None = None) -> Path:
    """B11: `outbox/lookback/<cohort>/Lookback_<cohort>.xlsx` with live variance formulas, yellow placeholders
    with notes, an M2 block per P&L tab and the Verification tab. Returns the workbook path."""
    import openpyxl
    from openpyxl.comments import Comment
    from openpyxl.styles import Font, PatternFill

    from cpa import activity_block, fsutil, manifest, verify
    from cpa.pptx import brand

    cohort = _check_cohort_name(cohort)
    root = _root(root)
    joined_path, rows = _read_joined(root, cohort)
    findings = _read_findings(root, cohort)
    yellow = brand.color("flag_yellow").lstrip("#")  # MissingAssumption before anything is written
    fill = PatternFill(fill_type="solid", fgColor=yellow, bgColor=yellow)
    bold = Font(name="Calibri", bold=True)
    joined_rel = manifest.to_rel(joined_path)

    def placeholder(cell, what: str, supplier: str) -> None:
        cell.value = f"{verify.M2_MISSING_MARKER}: {what}"
        cell.fill = fill
        cell.comment = Comment(f"Missing: {what}. Source: {supplier}. Type the value here; the P&L formulas "
                               "pick it up.", "cpa-lookback", width=320, height=120)

    wb = openpyxl.Workbook()
    summary = wb.active
    summary.title = "Summary"
    order = list(dict.fromkeys(r["record_id"] for r in rows))
    by_key = {(r["record_id"], r["line"]): (i, r) for i, r in enumerate(rows)}
    figures: dict[str, dict[str, Any]] = {}

    def figure(fid: str, value: Any, cell: str, ref_row: int, column: str, status: str, period: str) -> None:
        figures[fid] = {"value": value, "source_file": joined_rel, "ref": _joined_ref(ref_row, column),
                        "cell": cell, "status": status, "period": period}

    # Plan / Actuals / Offer tabs: one row per (record_id, line), constants or placeholders
    tab_rows: dict[str, dict[tuple[str, str], int]] = {"Plan": {}, "Actuals": {}, "Offer": {}}
    heads = {"Plan": "plan_value", "Actuals": "actual_value", "Offer": "offer_value"}
    sheets = {name: wb.create_sheet(name) for name in ("Plan", "Actuals", "Offer")}
    for name, ws in sheets.items():
        ws.append(["record_id", "line", "label", heads[name], "source_file", "source_ref", "period", "status"])
        for c in ws[1]:
            c.font = bold
    for i, r in enumerate(rows):
        kind = r["kind"]
        targets = []
        if kind in (FLOW, RATE):
            targets = [("Plan", "plan"), ("Actuals", "actual")]
        elif kind == OFFER:
            targets = [("Offer", "plan")]
        for name, side in targets:
            ws = sheets[name]
            status = STATUS["offer" if name == "Offer" else side]
            period = r[f"{side}_period"]
            ws.append([r["record_id"], r["line"], r["label"], None, r[f"{side}_file"], r[f"{side}_ref"], period, status])
            n = ws.max_row
            tab_rows[name][(r["record_id"], r["line"])] = n
            value = _num(r[f"{side}_value"])
            cell = ws.cell(row=n, column=4)
            if value is None:
                placeholder(cell, r[f"{side}_missing"] or f"{r['label']} {side}", r[f"{side}_supplier"] or "see join_report.md")
            else:
                cell.value = value
                cell.number_format = "#,##0.00"
                fid_side = "offer" if name == "Offer" else side
                figure(f"{r['record_id']}.{r['line']}.{fid_side}", value, f"'{name}'!D{n}", i, f"{side}_value", status,
                       period)

    # one P&L tab per provider
    pnl_rows: dict[tuple[str, str], int] = {}
    per_provider_findings: dict[str, list[Finding]] = {}
    for f in findings:
        per_provider_findings.setdefault(f.record_id, []).append(f)
    for rid in order:
        first = by_key[(rid, LINES[0].key)][1]
        ws = wb.create_sheet(f"{PNL_PREFIX}{rid}")
        ws["A1"] = f"Lookback {cohort}: {first['provider_label']} ({rid})"
        ws["A1"].font = bold
        ws["A2"] = RETENTION_FRAMING
        flow = next(r for r in rows if r["record_id"] == rid and r["kind"] == FLOW)
        for n, (label, value) in enumerate((("Department", first["department"]), ("Plan period", first["plan_period"]),
                                            ("Actual period", first["actual_period"]),
                                            ("Proration", flow["proration_label"])), start=3):
            ws.cell(row=n, column=1, value=label).font = bold
            ws.cell(row=n, column=2, value=value)
        ws.cell(row=7, column=1, value="Proration factor").font = bold
        factor = _num(flow["factor"])
        if factor is None:
            ws["B7"] = flow["proration_label"]
            ws["B7"].fill = fill
            reasons = "; ".join(f.detail for f in per_provider_findings.get(rid, []) if f.kind in ("PERIOD_BOUNDS", "LOGIC"))
            ws["B7"].comment = Comment(f"Missing: a usable proration factor. {reasons}. Source: cohort.csv period "
                                       "bounds (plan_start, plan_end, actual_start, actual_end).", "cpa-lookback",
                                       width=320, height=120)
        else:
            ws["B7"] = factor
            ws["B7"].number_format = "0.0000"
        heads_row = ("Measure", "Plan", "Proration", "Prorated plan", "Actual", "Variance", "Variance %", "Flags",
                     "Line id")
        for c, text in enumerate(heads_row, start=1):
            ws.cell(row=9, column=c, value=text).font = bold
        r_no = 10
        line_row: dict[str, int] = {}
        for ln in LINES:
            line_row[ln.key] = r_no
            r_no += 1
        for ln in LINES:
            i, row = by_key[(rid, ln.key)]
            r = line_row[ln.key]
            pnl_rows[(rid, ln.key)] = r
            ws.cell(row=r, column=1, value=ln.label)
            ws.cell(row=r, column=3, value=row["proration_label"])
            ws.cell(row=r, column=9, value=ln.key)
            if ln.kind in (FLOW, RATE):
                ws.cell(row=r, column=2, value=f"=Plan!D{tab_rows['Plan'][(rid, ln.key)]}")
                ws.cell(row=r, column=5, value=f"=Actuals!D{tab_rows['Actuals'][(rid, ln.key)]}")
            elif ln.kind == OFFER:
                ws.cell(row=r, column=2, value=f"=Offer!D{tab_rows['Offer'][(rid, ln.key)]}")
                ws.cell(row=r, column=5, value=f"=Actuals!D{tab_rows['Actuals'][(rid, ln.counterpart)]}")
            else:  # TCC = base + supplements
                b, s = line_row["base_salary"], line_row["supplements"]
                ws.cell(row=r, column=2, value=f'=IF(AND(ISNUMBER(B{b}),ISNUMBER(B{s})),B{b}+B{s},"n/a")')
                ws.cell(row=r, column=5, value=f'=IF(AND(ISNUMBER(E{b}),ISNUMBER(E{s})),E{b}+E{s},"n/a")')
            if ln.kind == FLOW:
                ws.cell(row=r, column=4, value=f'=IF(AND(ISNUMBER(B{r}),ISNUMBER($B$7)),B{r}*$B$7,"n/a")')
            else:
                ws.cell(row=r, column=4, value=f'=IF(ISNUMBER(B{r}),B{r},"n/a")')
            ws.cell(row=r, column=6, value=f'=IF(AND(ISNUMBER(D{r}),ISNUMBER(E{r})),E{r}-D{r},"n/a")')
            ws.cell(row=r, column=7, value=f'=IF(AND(ISNUMBER(F{r}),ISNUMBER(D{r})),IF(D{r}=0,"n/a",F{r}/D{r}),"n/a")')
            for c in (2, 4, 5, 6):
                ws.cell(row=r, column=c).number_format = "#,##0"
            ws.cell(row=r, column=7).number_format = "0.0%"
            flags = [f"OPEN {f.kind}" for f in per_provider_findings.get(rid, []) if f.line == ln.key]
            if row["plan_missing"] or row["actual_missing"]:
                flags.append("MISSING input (yellow cell on the linked tab)")
            ws.cell(row=r, column=8, value="; ".join(flags) or None)
            pnl = f"'{ws.title}'"
            if ln.kind == DERIVED:
                if _num(row["plan_value"]) is not None:
                    figure(f"{rid}.tcc.plan", _num(row["plan_value"]), f"{pnl}!B{r}", i, "plan_value",
                           STATUS["plan"], row["plan_period"])
                if _num(row["actual_value"]) is not None:
                    figure(f"{rid}.tcc.actual", _num(row["actual_value"]), f"{pnl}!E{r}", i, "actual_value",
                           STATUS["actual"], row["actual_period"])
            if _num(row["prorated_plan"]) is not None:
                figure(f"{rid}.{ln.key}.prorated_plan", _num(row["prorated_plan"]), f"{pnl}!D{r}", i, "prorated_plan",
                       STATUS["plan"], row["actual_period"])
            if _num(row["variance"]) is not None:
                figure(f"{rid}.{ln.key}.variance", _num(row["variance"]), f"{pnl}!F{r}", i, "variance",
                       STATUS["actual"], row["actual_period"])
        ws.column_dimensions["A"].width = 42
        ws.column_dimensions["C"].width = 40
        ws.column_dimensions["H"].width = 48

    # Summary: framing, one row per provider, links to the P&L variance cells
    summary["A1"] = f"Lookback {cohort}"
    summary["A1"].font = bold
    summary["A2"] = RETENTION_FRAMING
    head = ["record_id", "Provider", "Department", "Plan period", "Actual period", "Proration"]
    head += [f"{LINE_BY_KEY[k].label} variance" for k in SUMMARY_LINES] + ["Open findings"]
    for c, text in enumerate(head, start=1):
        summary.cell(row=4, column=c, value=text).font = bold
    for n, rid in enumerate(order, start=5):
        first = by_key[(rid, LINES[0].key)][1]
        flow = by_key[(rid, "wrvu")][1]
        for c, value in enumerate((rid, first["provider_label"], first["department"], first["plan_period"],
                                   first["actual_period"], flow["proration_label"]), start=1):
            summary.cell(row=n, column=c, value=value)
        for c, key in enumerate(SUMMARY_LINES, start=7):
            r = pnl_rows[(rid, key)]
            summary.cell(row=n, column=c, value=f"='{PNL_PREFIX}{rid}'!F{r}").number_format = "#,##0"
            i, row = by_key[(rid, key)]
            if _num(row["variance"]) is not None:
                figure(f"{rid}.{key}.variance.summary", _num(row["variance"]), f"'Summary'!{summary.cell(row=n, column=c).coordinate}",
                       i, "variance", STATUS["actual"], row["actual_period"])
        open_n = sum(1 for f in findings if f.record_id in (rid, "*"))
        summary.cell(row=n, column=7 + len(SUMMARY_LINES), value=f"{open_n} open")

    disc = wb.create_sheet("Discrepancies")
    disc.append(["kind", "record_id", "line", "detail", "left", "right", "status"])
    for c in disc[1]:
        c.font = bold
    for f in findings:
        disc.append([f.kind, f.record_id, f.line, f.detail, f.left, f.right, f.status])
    disc.column_dimensions["D"].width = 90

    out = _outbox(root, cohort) / f"Lookback_{cohort}.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    fsutil.atomic_write(out, lambda tmp: wb.save(str(tmp)))
    _, as_of = _as_of(joined_path)
    manifest.write(out, source="lookback", report="B11 lookback variance workbook", filters=f"cohort={cohort}",
                   as_of=as_of, inputs=[joined_path])
    manifest.update(out, figures=figures)

    for rid in order:
        metrics: dict[str, Any] = {"cFTE": {"why": "wRVUs per clinical FTE in the retention discussion",
                                            "source": "QGenda cFTE (cpa cfte); not a lookback input yet"}}
        for metric, key in (("wRVUs", "wrvu"), ("Collections (net)", "net_collections"), ("Charges", "charges")):
            row = by_key[(rid, key)][1]
            value = _num(row["actual_value"])
            if value is not None and row["actual_as_of"]:
                metrics[metric] = {"value": value, "period": row["actual_period"], "status": "actual",
                                   "source": "epic", "as_of": row["actual_as_of"]}
            else:
                metrics[metric] = {"why": "the provider's first-year activity against plan",
                                   "source": row["actual_supplier"] or "Epic aggregate extract"}
        activity_block.build(out, f"{PNL_PREFIX}{rid}", metrics, set(activity_block.METRICS))
    verify.build_verification_tab(out)
    return out


# ---------------------------------------------------------------- B12: deck


def _link_value(wb, formula: Any) -> float | None:
    """The constant a `=Sheet!A1` link points at (None for a placeholder, a blank or anything else)."""
    if not isinstance(formula, str) or not formula.startswith("="):
        return _num(formula) if not isinstance(formula, str) else None
    target = formula[1:]
    sheet, _, coord = target.rpartition("!")
    sheet = sheet.strip("'")
    if not sheet or sheet not in wb.sheetnames:
        return None
    value = wb[sheet][coord].value
    return None if isinstance(value, str) else _num(value)


def _mirror(wb, ws, line_rows: dict[str, int]) -> dict[str, dict[str, float | None]]:
    """What each P&L formula evaluates to, computed from the workbook's own constant cells."""
    factor = ws["B7"].value
    factor = None if isinstance(factor, str) else _num(factor)
    out: dict[str, dict[str, float | None]] = {}
    for key, r in line_rows.items():
        ln = LINE_BY_KEY[key]
        if ln.kind == DERIVED:
            b, s = out.get("base_salary", {}), out.get("supplements", {})
            plan = b["B"] + s["B"] if b.get("B") is not None and s.get("B") is not None else None
            actual = b["E"] + s["E"] if b.get("E") is not None and s.get("E") is not None else None
        else:
            plan = _link_value(wb, ws.cell(row=r, column=2).value)
            actual = _link_value(wb, ws.cell(row=r, column=5).value)
        if ln.kind == FLOW:
            prorated = plan * factor if plan is not None and factor is not None else None
        else:
            prorated = plan
        variance = actual - prorated if actual is not None and prorated is not None else None
        out[key] = {"B": plan, "D": prorated, "E": actual, "F": variance}
    return out


def _check_cached(cached_ws, key: str, r: int, values: dict[str, float | None]) -> None:
    for col, mirror in values.items():
        cached = cached_ws[f"{col}{r}"].value
        if cached is None:
            continue  # not recalculated yet: the mirror is the formula's value
        c = _num(cached) if not isinstance(cached, str) else None
        if (c is None) != (mirror is None) or (c is not None and abs(c - mirror) > 1e-6 * max(1.0, abs(c))):
            raise DeckNumberMismatch(f"{cached_ws.title}!{col}{r} ({key}) holds {cached!r} but its inputs give "
                                     f"{mirror!r}; recalculate or rebuild the workbook before the deck")


def _framing_check(text: str, where: str) -> None:
    low = text.casefold()
    for phrase in FRAMING_BANNED:
        if phrase in low:
            raise FramingViolation(f"{where} reads as a performance review ({phrase!r}); the lookback is framed as "
                                   "retention support. Rewrite that text; it is never rewritten automatically")


def deck(cohort: str, *, root: Path | str | None = None) -> Path:
    """B12: `outbox/lookback/<cohort>/Lookback_<cohort>.pptx`: a title slide, a summary slide and one support-plan
    slide per workbook provider, every number read from its workbook cell. Returns the deck path."""
    import openpyxl
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    from cpa import fsutil, manifest, verify
    from cpa.pptx import brand, lint, notes

    cohort = _check_cohort_name(cohort)
    root = _root(root)
    xlsx = _outbox(root, cohort) / f"Lookback_{cohort}.xlsx"
    if not xlsx.is_file():
        raise LookbackError(f"{xlsx} not found; run `python -m cpa lookback workbook --cohort {cohort}` first")
    findings = _read_findings(root, cohort)
    navy, yellow = brand.color("navy"), brand.color("flag_yellow")
    wb = openpyxl.load_workbook(str(xlsx))
    cached = openpyxl.load_workbook(str(xlsx), data_only=True)
    providers = [n[len(PNL_PREFIX):] for n in wb.sheetnames if n.startswith(PNL_PREFIX)]
    xlsx_rel = manifest.to_rel(xlsx)

    # read everything first; framing and number checks stop before a slide is drawn
    data = []
    for rid in providers:
        ws = wb[f"{PNL_PREFIX}{rid}"]
        line_rows = {}
        for r in range(10, ws.max_row + 1):
            key = ws.cell(row=r, column=9).value
            if isinstance(key, str) and key in LINE_BY_KEY:
                line_rows[key] = r
        values = _mirror(wb, ws, line_rows)
        for key, r in line_rows.items():
            _check_cached(cached[ws.title], key, r, values[key])
        title = str(ws["A1"].value or "").split(": ", 1)[-1]
        narrative_path = _staging(root, cohort) / NARRATIVE_DIR / f"{rid}.md"
        narrative = narrative_path.read_text(encoding="utf-8-sig").strip() if narrative_path.is_file() else ""
        if narrative:
            _framing_check(narrative, narrative_path.name)
        data.append({"rid": rid, "sheet": ws.title, "who": title, "values": values, "rows": line_rows,
                     "plan_period": ws["B4"].value, "actual_period": ws["B5"].value, "proration": ws["B6"].value,
                     "narrative": narrative,
                     "findings": [f for f in findings if f.record_id in (rid, "*")]})

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    figures: dict[str, dict[str, Any]] = {}
    note_slides: list = []
    all_text: list[tuple[str, str]] = []

    def textbox(slide, name, x, y, w, h, text, size=14, bold=False, color=None, fill=None):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        box.name = name
        tf = box.text_frame
        tf.word_wrap = True
        for i, para in enumerate(str(text).split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            run = p.add_run()
            run.text = para
            run.font.size, run.font.bold, run.font.name = Pt(size), bold, "Calibri"
            if color:
                run.font.color.rgb = RGBColor.from_string(color.lstrip("#"))
        if fill:
            box.fill.solid()
            box.fill.fore_color.rgb = RGBColor.from_string(fill.lstrip("#"))
        all_text.append((name, str(text)))
        return box

    def table(slide, x, y, w, rows_data, widths):
        shape = slide.shapes.add_table(len(rows_data), len(rows_data[0]), Inches(x), Inches(y), Inches(w),
                                       Inches(0.4 * len(rows_data)))
        tbl = shape.table
        for c, width in enumerate(widths):
            tbl.columns[c].width = Inches(width)
        for r, row in enumerate(rows_data):
            for c, (text, flagged) in enumerate(row):
                cell = tbl.cell(r, c)
                cell.text = str(text)
                for p in cell.text_frame.paragraphs:
                    for run in p.runs:
                        run.font.size, run.font.name, run.font.bold = Pt(12), "Calibri", r == 0
                if flagged:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor.from_string(yellow.lstrip("#"))
                all_text.append(("table", str(text)))
        return shape

    def record(fid, value, sheet, col, r, text, slide_no, status, period):
        figures[fid] = {"value": value, "cell": f"'{sheet}'!{col}{r}", "source_file": xlsx_rel,
                        "ref": f"'{sheet}'!{col}{r}", "text": text, "slide": slide_no, "status": status,
                        "period": period}

    # slide 1: title and purpose
    s1 = prs.slides.add_slide(blank)
    textbox(s1, "title", 0.6, 1.6, 12.1, 1.0, f"Lookback {cohort}: provider support review", 32, True, navy)
    textbox(s1, "framing", 0.6, 2.9, 12.1, 1.6, RETENTION_FRAMING, 18)
    textbox(s1, "source", 0.6, 5.2, 12.1, 0.6, f"Figures from {xlsx.name}; draft for review.", 12)
    note_slides.append(notes.Slide(title=f"Lookback {cohort}", points=[RETENTION_FRAMING]))

    # slide 2: summary
    s2 = prs.slides.add_slide(blank)
    textbox(s2, "title", 0.5, 0.3, 12.3, 0.8, f"Cohort summary: where support can help", 28, True, navy)
    head = [("Record ID", False), ("Provider", False), ("Proration", False)]
    head += [(f"{LINE_BY_KEY[k].label} vs plan", False) for k in SUMMARY_LINES] + [("Open findings", False)]
    body = [head]
    summary_figs = []
    for d in data:
        row = [(d["rid"], False), (d["who"].split(" (")[0], False), (d["proration"], False)]
        for key in SUMMARY_LINES:
            v = d["values"][key]["F"]
            text = _fmt_amount(v, key, signed=True) if v is not None else "n/a"
            row.append((text, v is None))
            if v is not None:
                record(f"{d['rid']}.{key}.variance.summary", v, d["sheet"], "F", d["rows"][key], text, 2, "actual",
                       str(d["actual_period"]))
                summary_figs.append(notes.Figure(f"{d['rid']} {LINE_BY_KEY[key].label} vs plan", text,
                                                 str(d["actual_period"]), "actual"))
        row.append((str(len(d["findings"])), bool(d["findings"])))
        body.append(row)
    table(s2, 0.5, 1.4, 12.3, body, [1.2, 1.6, 3.1, 1.6, 1.6, 1.7, 1.5])
    textbox(s2, "note", 0.5, 6.3, 12.3, 0.8, "Variance = actual minus the plan prorated to the actual period "
                                             "(label per provider). " + RETENTION_SHORT, 12)
    note_slides.append(notes.Slide(title="Cohort summary", points=[
        "Variances compare actuals with the plan prorated to each provider's actual period.", RETENTION_SHORT],
        figures=summary_figs))

    # one support-plan slide per provider in the workbook
    for d in data:
        slide_no = len(prs.slides) + 1
        s = prs.slides.add_slide(blank)
        textbox(s, "title", 0.5, 0.3, 12.3, 0.8, f"{d['who']}: support plan", 28, True, navy)
        textbox(s, "framing", 0.5, 1.05, 12.3, 0.5, RETENTION_SHORT, 14)
        textbox(s, "periods", 0.5, 1.5, 12.3, 0.6, f"Plan period {d['plan_period']}. Actual period "
                                                   f"{d['actual_period']}. Proration: {d['proration']}.", 12)
        body = [[("Measure", False), ("Plan (prorated)", False), ("Actual", False), ("Variance", False)]]
        figs, flagged = [], []
        for key in DECK_LINES:
            v, r = d["values"][key], d["rows"][key]
            cells = [(LINE_BY_KEY[key].label, False)]
            for col, status in (("D", "projected"), ("E", "actual"), ("F", "actual")):
                value = v[col]
                if value is None:
                    text = ("n/a" if col == "F" else "withheld" if col == "D" and v["B"] is not None
                            else verify.M2_MISSING_MARKER)
                    cells.append((text, col != "F"))
                    if col != "F" and LINE_BY_KEY[key].label not in flagged:
                        flagged.append(LINE_BY_KEY[key].label)
                    continue
                text = _fmt_amount(value, key, signed=col == "F")
                cells.append((text, False))
                suffix = DECK_COLUMNS[col]
                record(f"{d['rid']}.{key}.{suffix}", value, d["sheet"], col, r, text, slide_no, status,
                       str(d["actual_period"]))
                figs.append(notes.Figure(f"{LINE_BY_KEY[key].label} {suffix.replace('_', ' ')}", text,
                                         str(d["actual_period"]), status))
            body.append(cells)
        table(s, 0.5, 2.2, 7.4, body, [2.6, 1.6, 1.6, 1.6])
        if d["narrative"]:
            textbox(s, "narrative", 8.2, 2.2, 4.6, 2.4,
                    "Support context (from the manager recommendation letter):\n" + d["narrative"], 12)
        else:
            textbox(s, "narrative", 8.2, 2.2, 4.6, 2.4, brand.flag_text(
                "Recommendation letter narrative", why="the retention support narrative",
                source="manager recommendation letter; the cpa-lookback skill writes "
                       f"narrative/{d['rid']}.md"), 12, fill=yellow)
            flagged.append("Recommendation letter narrative")
        if d["findings"]:
            text = "Open items (listed, not resolved by this tool):\n" + "\n".join(
                f"{f.kind} {f.line}: {f.detail}" for f in d["findings"])
            textbox(s, "open_items", 8.2, 4.8, 4.6, 2.4, text, 11, fill=yellow)
            flagged.append("open items")
        else:
            textbox(s, "open_items", 8.2, 4.8, 4.6, 1.0, "Open items: none found in assembly.", 12)
        note_slides.append(notes.Slide(title=f"{d['who']}: support plan",
                                       points=[RETENTION_SHORT, f"Proration: {d['proration']}",
                                               f"Open findings: {len(d['findings'])} (on the slide)"],
                                       figures=figs, flagged_metrics=flagged))

    for name, text in all_text:
        _framing_check(text, f"slide text ({name})")

    out = _outbox(root, cohort) / f"Lookback_{cohort}.pptx"
    fsutil.atomic_write(out, lambda tmp: prs.save(str(tmp)))
    notes.write_notes(out, note_slides, "committee")
    issues = lint.lint_deck(out, audience="committee", deck_kind="companion")
    if issues:
        raise LintFailed("deck failed format lint: " + "; ".join(f"{i.code} {i.where}" for i in issues))
    _, as_of = _as_of(xlsx)
    manifest.write(out, source="lookback", report="B12 lookback deck", filters=f"cohort={cohort}", as_of=as_of,
                   row_count=len(prs.slides), inputs=[xlsx])
    manifest.update(out, figures=figures)
    return out


# ---------------------------------------------------------------- readiness (D19)


def _ready_extra(ws: Path, key: str) -> list[Path] | None:
    """Every provider on cohort.csv has exactly one bp_<record_id> file; None otherwise. Never raises."""
    try:
        from cpa import manifest

        rows = _read_csv(_staging(ws, key) / COHORT_FILE)
        files: list[Path] = []
        for row in rows:
            rid = row.get("record_id", "")
            hits = [p for p in (ws / "inbox" / "medvitals").glob(f"bp_{rid}.*")
                    if p.is_file() and not manifest.is_sidecar(p)] if rid else []
            if len(hits) != 1:
                return None
            files += hits
        return files or None
    except Exception:  # noqa: BLE001 - a readiness probe must not raise (state.Rule contract)
        return None


READINESS_RULES = (
    state.Rule(WORKFLOW, ("staging/lookback/{key}/" + COHORT_FILE, "inbox/sap/salary_{key}.*",
                          "inbox/epic/*_{key}.*"), "folder", extra=_ready_extra),
)


# ---------------------------------------------------------------- CLI


def _guard(action) -> int:
    from cpa import activity_block, config, crosswalk, periods
    from cpa.pptx import brand
    from cpa.sources.base import SourceError

    try:
        return action()
    except config.MissingAssumption as exc:
        print(f"lookback: {exc}", file=sys.stderr)
        return 2
    except (LookbackError, crosswalk.CrosswalkError, SourceError, config.ConfigError, ValueError,
            activity_block.ActivityBlockError, brand.BrandError, periods.PartialPeriod, OSError) as exc:
        print(f"lookback: {exc}", file=sys.stderr)
        return 3


def _emit(args: argparse.Namespace, data: dict, line: str) -> None:
    print(json.dumps(data, indent=2, ensure_ascii=False) if args.json else line)


def _cmd_assemble(args: argparse.Namespace) -> int:
    def action() -> int:
        res = assemble_result(args.cohort, root=args.root, prior=args.prior)
        _emit(args, res.to_json(), f"{res.joined}\n{len(res.findings)} open finding(s); see {res.report}")
        return 0

    return _guard(action)


def _cmd_workbook(args: argparse.Namespace) -> int:
    def action() -> int:
        out = workbook(args.cohort, root=args.root)
        _emit(args, {"workbook": str(out)}, str(out))
        return 0

    return _guard(action)


def _cmd_deck(args: argparse.Namespace) -> int:
    def action() -> int:
        out = deck(args.cohort, root=args.root)
        _emit(args, {"deck": str(out)}, str(out))
        return 0

    return _guard(action)


def _cmd_run(args: argparse.Namespace) -> int:
    def action() -> int:
        from cpa import manifest

        res = assemble_result(args.cohort, root=args.root, prior=args.prior)
        xlsx = workbook(args.cohort, root=args.root)
        pptx = deck(args.cohort, root=args.root)
        verification = (manifest.read(xlsx).get("verification") or {}).get("summary", "")
        data = {"joined": str(res.joined), "report": str(res.report), "workbook": str(xlsx), "deck": str(pptx),
                "findings": len(res.findings), "verification": verification}
        _emit(args, data, f"{xlsx}\n{pptx}\n{len(res.findings)} open finding(s); verification {verification}")
        return 0

    return _guard(action)


def register(subparsers: argparse._SubParsersAction) -> None:
    """Register `lookback assemble|workbook|deck|run`. Import-cheap (D03)."""
    top = subparsers.add_parser("lookback", help="Lookback: cohort assembly, variance workbook, provider deck "
                                                 "(B10-B12).")
    sub = top.add_subparsers(dest="command", required=True, title="commands")
    leaves = (
        ("assemble", "B10: join the cohort on the MedVitals record ID; write joined.csv and join_report.md.",
         _cmd_assemble),
        ("workbook", "B11: build Lookback_<cohort>.xlsx (live formulas, placeholders, M2, Verification).",
         _cmd_workbook),
        ("deck", "B12: build Lookback_<cohort>.pptx (retention framing, one slide per provider).", _cmd_deck),
        ("run", "B10, B11 and B12 in order.", _cmd_run),
    )
    for name, help_text, func in leaves:
        p = sub.add_parser(name, help=help_text, description=help_text + " Exit 2 when a needed assumption is "
                                                                         "null, 3 on a data stop.")
        p.add_argument("--cohort", required=True, help="Cohort name: the folder under staging/lookback.")
        p.add_argument("--root", type=Path, default=None, help="Workspace root (default: config.workspace()).")
        p.add_argument("--json", action="store_true", help="Print the result as JSON.")
        if name in ("assemble", "run"):
            p.add_argument("--prior", type=Path, default=None,
                           help="Prior lookback workbook (default: staging/lookback/<cohort>/prior_lookback.xlsx "
                                "when present).")
        p.set_defaults(func=func)
