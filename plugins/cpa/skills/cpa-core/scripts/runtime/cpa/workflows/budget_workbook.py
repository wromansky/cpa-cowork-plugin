"""Budget workbook rebuild from SAP CO detail (B17, U20 "budget" part, Build List :213-219).

Streams the A11 CO line-item export (often above 15 MB) through `cpa.sources.get("sap.co_lineitems")`
(hard rule 12: SAP access goes through the adapter only) and produces two flat, pivot-ready layouts:
`Monthly Line Items` (source columns verbatim, plus one column per fiscal period) and `Year Line Items`
(source columns verbatim, plus one column per fiscal year). Every source row is kept, including a row
whose values are entirely blank (hard rule "keep every row"), and all ten cost-structure columns are
kept -- the acceptance line names a past incident where People vs Thing, Active vs Inactive, GAAP Lvl 2
and GAAP Lvl 4 were dropped in a rebuild; `structure_columns()` names them from the adapter's own schema
(no literal duplication) and a missing one raises `DroppedColumns` before anything is written.

Hard rules enforced: 11 (streamed via `TabularAdapter.read`, which is `cpa.bigxlsx.iter_rows` under the
hood -- D10 -- and the output workbook is `openpyxl.Workbook(write_only=True)`; neither leg loads the
whole file), 12 (SAP reached only through the adapter), 10 (an unrecognised fiscal-period value fails
loudly as `UnmappedPeriods`, naming every offending value, rather than being silently dropped).

The C5 prior-month check sheet and `--status`/STATUS_MISSING handling are not implemented here.
"""
from __future__ import annotations

from cpa import office

import argparse
import json
import sys
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import Any

MONTHLY_SHEET = "Monthly Line Items"
YEARLY_SHEET = "Year Line Items"
OUTPUT_STEM = "Budget_Workbook_"
REPORT = "Budget Workbook"
SOURCE_SYSTEM = "sap"  # FIXTURE -- confirm against her file when Workday replaces SAP (hard rule 12)
CO_ADAPTER_NAME = "sap.co_lineitems"

EXIT_OK = 0
EXIT_STOPPED = 1

__all__ = [
    "BudgetWorkbookError", "DroppedColumns", "UnmappedPeriods", "BudgetResult",
    "structure_columns", "run", "register",
]


class BudgetWorkbookError(RuntimeError):
    """The build stopped before anything was written; the message names why."""


class DroppedColumns(BudgetWorkbookError):
    """One or more required columns are absent from the export. `.missing` names every one, `.where`
    says where they were found absent ("export")."""

    def __init__(self, missing: tuple[str, ...], where: str) -> None:
        self.missing = tuple(missing)
        self.where = where
        super().__init__(f"{where}: missing required column(s) {', '.join(self.missing)}")


class UnmappedPeriods(BudgetWorkbookError):
    """A fiscal_period cell was neither a bare FYMM token nor a valid SAP period label. `.values` maps
    each bad text to how many rows carried it."""

    def __init__(self, values: dict[str, int]) -> None:
        self.values = dict(values)
        named = ", ".join(f"{v!r} ({n}x)" for v, n in sorted(values.items()))
        super().__init__(f"unmapped fiscal period value(s): {named}")


@dataclass(frozen=True)
class BudgetResult:
    """What `run` did. `workbook` and `sources` are workspace-relative POSIX strings when inside one."""

    workbook: Path
    fymm: str
    row_count: int
    periods: tuple[str, ...]
    fiscal_years: tuple[int, ...]
    verify_summary: str
    sources: tuple[str, ...] = field(default_factory=tuple)

    def to_json(self) -> dict[str, Any]:
        return {
            "workbook": str(self.workbook), "fymm": self.fymm, "row_count": self.row_count,
            "periods": list(self.periods), "fiscal_years": list(self.fiscal_years),
            "verify_summary": self.verify_summary, "sources": list(self.sources),
        }


def _co_adapter():
    from cpa import sources

    return sources.get(CO_ADAPTER_NAME)


def structure_columns() -> tuple[str, ...]:
    """The ten required canonical CO columns, read from the adapter's own schema (R108); no literal here."""
    adapter = _co_adapter()
    return tuple(c.name for c in adapter.COLUMNS if c.required)


def _period_key(value: Any) -> str | None:
    """A fiscal_period cell -> its FYMM key, or None for a blank cell. Raises ValueError (caller wraps it
    per-row into UnmappedPeriods) for anything that is neither a bare FYMM token nor a valid SAP label."""
    from cpa import periods

    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    try:
        periods.parse_fymm(text)
        return text
    except ValueError:
        pass
    return periods.parse_sap_label(text)  # raises ValueError for anything else; caller wraps it


def _check_required(column_map: dict[str, str], missing: list[str]) -> None:
    if missing:  # adapter.match()'s missing is exactly the required=True columns absent (structure_columns())
        raise DroppedColumns(tuple(missing), where="export")
    layout_missing = [c for c in ("fiscal_period", "amount") if c not in column_map]
    if layout_missing:
        raise DroppedColumns(tuple(layout_missing), where="export")


@dataclass(frozen=True)
class _Scan:
    header: list[str]
    column_map: dict[str, str]
    row_count: int
    fymm_keys: tuple[str, ...]  # sorted
    fiscal_years: tuple[int, ...]  # sorted


def _scan(export: Path) -> _Scan:
    """Pass 1: stream the export once to validate the header and discover the row count and the set of
    fiscal periods present. Raises DroppedColumns / UnmappedPeriods before any row is counted wrong."""
    from cpa import periods

    adapter = _co_adapter()
    header, body = adapter.read(export)
    column_map, missing = adapter.match(header)
    _check_required(column_map, missing)
    period_idx = header.index(column_map["fiscal_period"])
    rows = 0
    keys: set[str] = set()
    bad: dict[str, int] = {}
    try:
        for row in body:
            rows += 1
            raw = row[period_idx] if period_idx < len(row) else None
            try:
                key = _period_key(raw)
            except ValueError:
                text = str(raw).strip()
                bad[text] = bad.get(text, 0) + 1
                continue
            if key is not None:
                keys.add(key)
    finally:
        body.close()
    if bad:
        raise UnmappedPeriods(bad)
    fymm_sorted = tuple(sorted(keys))
    fy_sorted = tuple(sorted({periods.parse_fymm(k)[0] for k in fymm_sorted}))
    return _Scan(header=header, column_map=column_map, row_count=rows, fymm_keys=fymm_sorted,
                fiscal_years=fy_sorted)


def _blank_row(width: int) -> list[None]:
    return [None] * width


def _build_workbook(export: Path, scan: _Scan):
    """Pass 2: a fresh streamed read, writing both layout sheets into a write_only workbook."""
    import openpyxl

    from cpa import periods

    adapter = _co_adapter()
    header, body = adapter.read(export)
    period_idx = header.index(scan.column_map["fiscal_period"])
    amount_idx = header.index(scan.column_map["amount"])
    monthly_headers = [periods.sap_label(k) for k in scan.fymm_keys]
    year_headers = [f"FY{fy}" for fy in scan.fiscal_years]
    fy_of = {k: periods.parse_fymm(k)[0] for k in scan.fymm_keys}

    wb = openpyxl.Workbook(write_only=True)
    ws_m = wb.create_sheet(title=MONTHLY_SHEET)
    from cpa.pptx import brand

    brand.append_generated_row(ws_m, list(header) + monthly_headers, row_number=1, header=True)
    ws_y = wb.create_sheet(title=YEARLY_SHEET)
    brand.append_generated_row(ws_y, list(header) + year_headers, row_number=1, header=True)

    written = 0
    try:
        for row in body:
            row = list(row)
            amount = row[amount_idx] if amount_idx < len(row) else None
            raw_period = row[period_idx] if period_idx < len(row) else None
            key = _period_key(raw_period)  # already proven mappable in pass 1

            m_row = row + _blank_row(len(monthly_headers))
            if key is not None:
                m_row[len(header) + scan.fymm_keys.index(key)] = amount
            brand.append_generated_row(ws_m, m_row, row_number=written + 2)

            y_row = row + _blank_row(len(year_headers))
            if key is not None:
                y_row[len(header) + scan.fiscal_years.index(fy_of[key])] = amount
            brand.append_generated_row(ws_y, y_row, row_number=written + 2)
            written += 1
    finally:
        body.close()
    if written != scan.row_count:  # pass 1 and pass 2 must see the same stream (never a silent mismatch)
        raise BudgetWorkbookError(
            f"{export.name}: row count changed between passes ({scan.row_count} then {written}); "
            "the export may have been edited mid-run"
        )
    return wb


def run(export: Path | str, *, fymm: str | None = None, root: Path | str | None = None) -> BudgetResult:
    """Build `<root|workspace>/outbox/budget/<fymm>/Budget_Workbook_<fymm>.xlsx` from `export` and return
    what was built. `fymm` defaults to the FYMM token in the export's filename. Raises DroppedColumns,
    UnmappedPeriods or a cpa.config error before anything is written."""
    from cpa import config, fsutil, manifest, periods, verify

    export = Path(export)
    if not export.is_file():
        raise FileNotFoundError(f"CO line-item export not found: {export}")
    resolved_fymm = fymm if fymm is not None else periods.fymm_from_filename(export.name)
    periods.parse_fymm(resolved_fymm)  # ValueError on a bad token, before anything is read

    scan = _scan(export)
    wb = _build_workbook(export, scan)

    ws_root = Path(root) if root is not None else config.workspace()
    out_dir = ws_root / "outbox" / "budget" / resolved_fymm
    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / f"{OUTPUT_STEM}{resolved_fymm}.xlsx"
    office.save_workbook(wb, out)

    manifest.write(out, SOURCE_SYSTEM, REPORT, f"fiscal period {resolved_fymm}", date.today().isoformat(),
                   row_count=scan.row_count, inputs=[export], period=resolved_fymm, status="budget")
    # Only attach figure provenance when the input itself carries its genuine source/as-of manifest.
    # Otherwise verification must continue to expose SOURCE_UNKNOWN rather than self-tie-out.
    if manifest.exists(export):
        import openpyxl
        from openpyxl.utils import get_column_letter

        source_manifest = manifest.read(export)
        if source_manifest.get("source") and source_manifest.get("as_of"):
            figures: dict[str, dict[str, Any]] = {}
            source_sheet = None
            if export.suffix.lower() in (".xlsx", ".xlsm", ".xlsb"):
                from cpa import bigxlsx

                source_sheet = bigxlsx.sheet_names(export)[0]
            adapter = _co_adapter()
            header, body = adapter.read(export)
            period_idx = header.index(scan.column_map["fiscal_period"])
            amount_idx = header.index(scan.column_map["amount"])
            try:
                for row_index, raw_row in enumerate(body, start=2):
                    raw_period = raw_row[period_idx] if period_idx < len(raw_row) else None
                    key = _period_key(raw_period)
                    if key is None:
                        continue
                    amount = raw_row[amount_idx] if amount_idx < len(raw_row) else None
                    if not isinstance(amount, (int, float)) or isinstance(amount, bool):
                        continue
                    for out_sheet in (MONTHLY_SHEET, YEARLY_SHEET):
                        if out_sheet == MONTHLY_SHEET:
                            period_offset = scan.fymm_keys.index(key)
                            period = periods.sap_label(key)
                        else:
                            fy = periods.parse_fymm(key)[0]
                            period_offset = scan.fiscal_years.index(fy)
                            period = f"FY{fy}"
                        output_cell = (f"{out_sheet}!{get_column_letter(len(scan.header) + period_offset + 1)}"
                                       f"{row_index}")
                        source_cell = f"{get_column_letter(amount_idx + 1)}{row_index}"
                        if source_sheet is not None:
                            source_cell = f"{source_sheet}!{source_cell}"
                        figures[output_cell] = {
                            "value": amount, "source_file": manifest.to_rel(export), "ref": source_cell,
                            "cell": output_cell, "status": "budget", "period": period,
                        }
            finally:
                body.close()
            manifest.update(out, figures=figures)
    vres = verify.build_verification_tab(out)

    return BudgetResult(
        workbook=out, fymm=resolved_fymm, row_count=scan.row_count, periods=scan.fymm_keys,
        fiscal_years=scan.fiscal_years, verify_summary=vres.summary, sources=(manifest.to_rel(export),),
    )


def _cmd_build(args: argparse.Namespace) -> int:
    from cpa import config

    try:
        result = run(args.export, fymm=args.fymm, root=args.root)
    except (BudgetWorkbookError, FileNotFoundError, ValueError, config.ConfigError) as exc:
        print(f"stopped: {exc}", file=sys.stderr)
        return EXIT_STOPPED
    if args.json:
        print(json.dumps(result.to_json(), ensure_ascii=False))
    else:
        print(f"{result.workbook}: {result.row_count} row(s), {len(result.periods)} period(s), "
              f"verify {result.verify_summary}")
    return EXIT_OK


def register(subparsers: argparse._SubParsersAction) -> None:
    """Register `budget_workbook build`. Import-cheap: no file access at module import time (D03)."""
    top = subparsers.add_parser("budget_workbook", help=__doc__.splitlines()[0])
    sub = top.add_subparsers(dest="command", required=True, title="commands")

    b = sub.add_parser(
        "build", help="Rebuild the budget workbook from an SAP CO line-item export.",
        description="Stream the CO line-item export into outbox/budget/<fymm>/Budget_Workbook_<fymm>.xlsx "
                    "with Monthly Line Items and Year Line Items layouts, every row and all ten cost "
                    "structure columns kept, plus a Verification tab. Exit 0 built, 1 stopped.")
    b.add_argument("--export", required=True, type=Path, help="Path to the SAP CO line-item export.")
    b.add_argument("--fymm", default=None, help="Fiscal period, yymm (default: parsed from --export's name).")
    b.add_argument("--root", type=Path, default=None, help="Workspace root (default: cpa.config.workspace()).")
    b.add_argument("--json", action="store_true", help="Print the result as one JSON object.")
    b.set_defaults(func=_cmd_build)
